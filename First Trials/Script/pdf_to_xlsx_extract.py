import re
import openpyxl
from openpyxl.styles import Alignment
import pandas as pd
import pdfplumber


def fix_bconnect_cell(text):
    """Cleans and reverses extracted RTL text from B-Connect PDFs:

    - Normalizes bracket orientations.
    - Reverses character sequences for Arabic tokens while preserving numbers and
    English.
    - Restores correct word sequence across mixed Arabic/English descriptions.
    """
    if not text:
        return ""

    # 1. Swap parentheses/brackets direction (RTL extraction inversion fix)
    bracket_map = str.maketrans("()[]{}", ")(][}{")
    text = text.translate(bracket_map)

    # Normalize spaces around punctuation marks so tokens split cleanly
    text = re.sub(r"([()\[\]{}])", r" \1 ", text)
    tokens = text.split()

    fixed_tokens = []
    # Fixed Unicode range: added \u before FB50
    arabic_pattern = re.compile(
        r"[\u0600-\u06FF\u0750-\u077F\u08A0-\u08FF\uFB50-\uFDFF\uFE70-\uFEFF]"
    )

    for token in tokens:
        if arabic_pattern.search(token):
            # Split numbers/latin tokens from Arabic characters inside the same word
            sub_tokens = re.split(r"(\d+|[A-Za-z]+|[^\w\s])", token)
            sub_fixed = []
            for st in sub_tokens:
                if arabic_pattern.search(st):
                    sub_fixed.append(st[::-1])  # Reverse Arabic letters
                else:
                    sub_fixed.append(st)  # Preserve digits / Latin text
            fixed_tokens.append("".join(sub_fixed))
        else:
            fixed_tokens.append(token)

    # Group contiguous Arabic words and punctuation, then reverse phrase sequence
    result = []
    arabic_group = []

    for token in fixed_tokens:
        is_arabic_or_symbol = bool(arabic_pattern.search(token)) or token in (
            "()+*-/[]"
        )
        if is_arabic_or_symbol:
            arabic_group.append(token)
        else:
            if arabic_group:
                result.extend(arabic_group[::-1])
                arabic_group = []
            result.append(token)

    if arabic_group:
        result.extend(arabic_group[::-1])

    res_str = " ".join(result)

    # Clean up excess spaces added around brackets
    res_str = re.sub(r"\s+([()\-[\]])", r"\1", res_str)
    res_str = re.sub(r"([()\-[\]])\s+", r"\1", res_str)

    return res_str.strip()


def sanitize_code(code_text):
    """Sanitizes Code column to prevent cross-column overlap errors."""
    if not code_text:
        return ""
    cleaned = re.sub(r"[^\w\-]", "", code_text)
    return cleaned


# File paths
pdf_path = "Items_Catalogue/item_cat_short.pdf"
excel_path = "catalog_final_fixed.xlsx"

extracted_data = []

print("Extracting and parsing PDF catalog...")

with pdfplumber.open(pdf_path) as pdf:
    for page in pdf.pages:
        tables = page.extract_tables()
        for table in tables:
            for row in table:
                if row and len(row) >= 2:
                    raw_code = row[-1].replace("\n", " ").strip() if row[-1] else ""
                    raw_name = row[-2].replace("\n", " ").strip() if row[-2] else ""

                    fixed_code = sanitize_code(raw_code)
                    fixed_name = fix_bconnect_cell(raw_name)

                    if fixed_code != "الكود" and (fixed_code or fixed_name):
                        extracted_data.append([fixed_code, fixed_name])

# Write DataFrame to Excel
df = pd.DataFrame(extracted_data, columns=["الكود", "الاسم"])
df.to_excel(excel_path, index=False)

# Apply explicit formatting via openpyxl
wb = openpyxl.load_workbook(excel_path)
ws = wb.active

# Enable Right-to-Left sheet view
ws.sheet_view.rightToLeft = True

# Format columns and alignments
ws.column_dimensions["A"].width = 18
ws.column_dimensions["B"].width = 55

align_center = Alignment(horizontal="center", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")

for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=2):
    row[0].alignment = align_center  # Code Column (A)
    row[1].alignment = align_right  # Name Column (B)

wb.save(excel_path)
print(f"Extraction complete! Saved to '{excel_path}'.")

#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
====================================================================================================
GENIUS (e-Plus) PHARMACY ERP - MASTER DATA CONSOLIDATION & INTELLIGENCE ENGINE (ENTERPRISE EDITION)
====================================================================================================
Consolidated Production Architecture:
  1. Inventory Master & Batch Control Engine  -> 01_Genius_Inventory_Master.xlsx          [13 Sheets]
  2. Sales Operations & Revenue Engine        -> 02_Genius_Sales_Operations.xlsx          [14 Sheets]
  3. Procurement & Vendor Supply Chain        -> 03_Genius_Purchasing_Operations.xlsx      [12 Sheets]
  4. Treasury & General Financial Ledger      -> 04_Genius_Financial_Ledger.xlsx          [12 Sheets]
  5. POS Cashier & Shift Audit Trail          -> 05_Genius_POS_Shift_Operations.xlsx      [10 Sheets]
  6. Executive Business Intelligence Cockpit  -> 06_Genius_Executive_Intelligence_Cockpit.xlsx [11 Sheets]
====================================================================================================
Total Production Scope: 6 Workbooks | 72 Analytical Sheets | 100% Native Excel Tables (No Totals Clutter)
====================================================================================================
Key Architectural Standards:
  - Universal Historical Continuity: Full inclusion of all years (2023-2026+) with complete ledger integrity.
  - In-Memory OpenPyXL Styling: Direct single-pass worksheet formatting (eliminates 2-pass disk overhead).
  - Exact Net Revenue & COGS Alignment: Header revenue netted against sales returns to prevent profit inflation.
  - Multi-Unit Pack Normalization: Blisters/strips/tablets normalized to standard boxes with zero-division guards.
  - Varbinary Decryption Engine: Bidi buffer reversal + windows-1256 trade name restoration.
  - Corrected Accounting Conventions: AP (ven_current_credit > 0), AR (cust_current_credit > 0).
  - True Weighted Average Cost: SUM(qty * cost) / NULLIF(SUM(qty), 0).
  - Safe Type-Coercion: TRY_CAST on all employee/user auditing keys.
====================================================================================================
"""

import os
import sys
import time
import re
import logging
import urllib.parse
from datetime import datetime
import pandas as pd
from sqlalchemy import create_engine, text
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# --------------------------------------------------------------------------------------------------
# 0. LOGGING CONFIGURATION
# --------------------------------------------------------------------------------------------------
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[logging.StreamHandler(sys.stdout)],
)
logger = logging.getLogger("GeniusMasterEngine")

# --------------------------------------------------------------------------------------------------
# 1. DATABASE CONNECTION SETUP
# --------------------------------------------------------------------------------------------------
SERVER = "localhost,1433"
DATABASE = "Genius"
USERNAME = "sa"
PASSWORD = "Emma@2015"

conn_str = (
    f"DRIVER={{ODBC Driver 18 for SQL Server}};"
    f"SERVER={SERVER};"
    f"DATABASE={DATABASE};"
    f"UID={USERNAME};"
    f"PWD={PASSWORD};"
    f"TrustServerCertificate=yes;"
)
params = urllib.parse.quote_plus(conn_str)
ENGINE = create_engine(f"mssql+pyodbc:///?odbc_connect={params}", pool_pre_ping=True)


# --------------------------------------------------------------------------------------------------
# 2. VARBINARY ITEM DECRYPTION ENGINE (GENIUS ARCHITECTURAL CORE)
# --------------------------------------------------------------------------------------------------
def decode_legacy_bidi_bytes(raw_bytes):
    """
    Decodes obfuscated varbinary medicine names in Item_Catalog.
    Splits binary in half (N // 2), reverses each chunk independently,
    concatenates back together, and decodes with windows-1256.
    """
    if pd.isna(raw_bytes) or not raw_bytes:
        return ""
    try:
        raw_b = bytes(raw_bytes)
        N = len(raw_b)
        L = N // 2
        corrected_bytes = raw_b[:L][::-1] + raw_b[L:][::-1]
        return corrected_bytes.decode("windows-1256", errors="replace").strip()
    except Exception:
        return "Decoding Error"


ITEM_NAME_AR_CACHE = None
ITEM_NAME_EN_CACHE = None


def get_item_name_cache(engine):
    """
    Initializes and caches decrypted trade names across all items in Item_Catalog.
    Provides fallback keys (float, int, str) for reliable joins.
    """
    global ITEM_NAME_AR_CACHE, ITEM_NAME_EN_CACHE
    if ITEM_NAME_AR_CACHE is not None and ITEM_NAME_EN_CACHE is not None:
        return ITEM_NAME_AR_CACHE, ITEM_NAME_EN_CACHE

    logger.info("Initializing Item Decryption Cache from Item_Catalog...")
    q = """
    SELECT itm_id, itm_name_ar, itm_name_en, itm_name_ar_encrypt, itm_name_en_encrypt 
    FROM Item_Catalog WITH (NOLOCK)
    """
    try:
        df_cat = pd.read_sql(q, engine)
        cache_ar = {}
        cache_en = {}

        for _, row in df_cat.iterrows():
            iid = row["itm_id"]
            if pd.isna(iid):
                continue

            name_ar = row["itm_name_ar"]
            if pd.isna(name_ar) or not str(name_ar).strip():
                name_ar = decode_legacy_bidi_bytes(row["itm_name_ar_encrypt"])
            else:
                name_ar = str(name_ar).strip()

            name_en = row["itm_name_en"]
            if pd.isna(name_en) or not str(name_en).strip():
                name_en = decode_legacy_bidi_bytes(row["itm_name_en_encrypt"])
            else:
                name_en = str(name_en).strip()

            if not name_en and name_ar:
                name_en = name_ar
            if not name_ar and name_en:
                name_ar = name_en

            try:
                iid_f = float(iid)
                iid_i = int(iid_f)
                cache_ar[iid_f] = name_ar
                cache_ar[iid_i] = name_ar
                cache_ar[str(iid_i)] = name_ar
                cache_ar[str(iid_f)] = name_ar

                cache_en[iid_f] = name_en
                cache_en[iid_i] = name_en
                cache_en[str(iid_i)] = name_en
                cache_en[str(iid_f)] = name_en
            except ValueError, TypeError:
                cache_ar[iid] = name_ar
                cache_ar[str(iid)] = name_ar
                cache_en[iid] = name_en
                cache_en[str(iid)] = name_en

        ITEM_NAME_AR_CACHE = cache_ar
        ITEM_NAME_EN_CACHE = cache_en
        logger.info(f"Item Decryption Cache built: {len(df_cat):,} items mapped.")
    except Exception as e:
        logger.warning(f"Could not build item decryption cache: {e}")
        ITEM_NAME_AR_CACHE = {}
        ITEM_NAME_EN_CACHE = {}

    return ITEM_NAME_AR_CACHE, ITEM_NAME_EN_CACHE


def read_sql_and_enrich(sql, engine):
    """
    Executes a SQL query and enriches trade names using the decrypted in-memory catalog.
    Handles primary items, substitutes, and original referral medicines.
    """
    df = pd.read_sql(sql, engine)
    cache_ar, cache_en = get_item_name_cache(engine)

    id_col = None
    for candidate in [
        "Item_ID",
        "itm_id",
        "st_itm_id",
        "invd_itm_id",
        "iqc_itm_id",
        "ptd_itm_id",
        "Shortage_Item_ID",
        "Original_Item_ID",
    ]:
        if candidate in df.columns:
            id_col = candidate
            break

    if id_col is not None:
        for ar_col in [
            "Item_Name",
            "Item_Name_Ar",
            "Item_Name_ar",
            "Original_Item_Name",
        ]:
            if ar_col in df.columns:
                mapped = df[id_col].map(cache_ar)
                df[ar_col] = (
                    df[ar_col].replace(r"^\s*$", None, regex=True).fillna(mapped)
                )

        for en_col in ["Item_Name_En", "Item_Name_en", "Original_Item_Name_En"]:
            if en_col in df.columns:
                mapped = df[id_col].map(cache_en)
                df[en_col] = (
                    df[en_col].replace(r"^\s*$", None, regex=True).fillna(mapped)
                )

    if "Substitute_Item_ID" in df.columns:
        if "Substitute_Item_Name" in df.columns:
            mapped_sub_ar = df["Substitute_Item_ID"].map(cache_ar)
            df["Substitute_Item_Name"] = (
                df["Substitute_Item_Name"]
                .replace(r"^\s*$", None, regex=True)
                .fillna(mapped_sub_ar)
            )
        if "Substitute_Item_Name_En" in df.columns:
            mapped_sub_en = df["Substitute_Item_ID"].map(cache_en)
            df["Substitute_Item_Name_En"] = (
                df["Substitute_Item_Name_En"]
                .replace(r"^\s*$", None, regex=True)
                .fillna(mapped_sub_en)
            )

    return df


# --------------------------------------------------------------------------------------------------
# 3. OPENPYXL HIGH-PERFORMANCE STYLING ENGINE (DIRECT IN-MEMORY TABLE FORMATTING)
# --------------------------------------------------------------------------------------------------
THIN_BORDER = Border(
    left=Side(style="thin", color="D3D3D3"),
    right=Side(style="thin", color="D3D3D3"),
    top=Side(style="thin", color="D3D3D3"),
    bottom=Side(style="thin", color="D3D3D3"),
)

COLOR_TO_TABLE_STYLE = {
    "1B4D3E": "TableStyleMedium7",  # Emerald Green (Module 1: Inventory)
    "1F4E79": "TableStyleMedium9",  # Sapphire Blue (Module 2: Sales)
    "4A154B": "TableStyleMedium12",  # Deep Burgundy (Module 3: Purchases)
    "7D5A00": "TableStyleMedium13",  # Bronze Gold   (Module 4: Financials)
    "1A5276": "TableStyleMedium10",  # Slate Teal    (Module 5: POS Shifts)
    "0B132B": "TableStyleMedium1",  # Dark Navy     (Module 6: Executive BI)
}


def get_excel_format_and_alignment(header_name: str):
    """
    Maps column names to appropriate Excel number formats and alignments.
    Distinguishes true decimal ratios (0-1) from raw percentage values (0-100).
    """
    h = header_name.lower().strip()
    is_wage_or_factor = any(
        term in h
        for term in [
            "hourly",
            "wage",
            "points_cash",
            "conversion",
            "per_unit",
            "spread",
            "delta",
        ]
    )

    # True ratios (stored as 0.0 - 1.0) formatted with native percentage: 0.00%
    is_true_ratio = not is_wage_or_factor and (
        h
        in [
            "gross_margin_perc",
            "gross_margin_rate",
            "share_of_total_sales",
            "cumulative_revenue_share",
            "potential_profit_margin_perc",
            "credit_limit_utilization",
        ]
        or any(h.endswith(sfx) for sfx in ["_share", "margin_perc", "margin_rate"])
    )
    if is_true_ratio:
        return "0.00%", Alignment(horizontal="right", vertical="center")

    # Raw database percentage integers/floats (stored as 0.0 - 100.0) formatted as #,##0.00"%"
    is_raw_perc = (
        not is_wage_or_factor
        and (
            any(
                h.endswith(sfx)
                for sfx in ["_perc", "_percentage", "_rate", "_ratio", "_percent", "%"]
            )
            or any(
                h.startswith(pfx) for pfx in ["perc_", "rate_", "ratio_", "percentage_"]
            )
            or any(
                term in h
                for term in [
                    "discount_perc",
                    "copay_perc",
                    "tax_rate",
                    "fee_perc",
                    "discount_rate",
                ]
            )
        )
        and not any(
            term in h
            for term in [
                "revenue",
                "sales",
                "cogs",
                "profit",
                "amount",
                "total",
                "val",
                "cost",
                "price",
                "fee",
            ]
        )
    )

    if is_raw_perc:
        return '#,##0.00"%"', Alignment(horizontal="right", vertical="center")

    # Financial currencies, monetary values, and unit cost/prices
    if (
        any(
            term in h
            for term in [
                "price",
                "cost",
                "val",
                "total",
                "net",
                "credit",
                "debit",
                "balance",
                "revenue",
                "cogs",
                "profit",
                "cash",
                "amount",
                "paid",
                "refund",
                "fee",
                "disbursed",
                "salary",
                "bonus",
                "bouns",
                "pen",
                "gross",
                "outlay",
                "yield",
                "liability",
                "float",
                "deduction",
                "penalty",
                "variance",
                "inflow",
                "outflow",
                "drop",
            ]
        )
        or is_wage_or_factor
    ):
        return "#,##0.00", Alignment(horizontal="right", vertical="center")

    # Quantities, unit counts, lines, inventory units
    if any(
        term in h
        for term in [
            "qty",
            "quantity",
            "count",
            "items",
            "no_",
            "units",
            "lines",
            "days",
            "hours",
            "rank",
        ]
    ):
        return "#,##0.00", Alignment(horizontal="right", vertical="center")

    return None, Alignment(horizontal="center", vertical="center")


def style_worksheet(
    ws, header_fill_hex="1F4E79", is_kpi=False, existing_table_names=None
):
    """
    Applies high-performance formatting directly to an active openpyxl worksheet.
    Enforces native Excel Table definitions, dynamic column auto-fitting, Segoe UI fonts,
    thin cell borders, frozen top headers, and no total rows.
    """
    if existing_table_names is None:
        existing_table_names = set()

    max_row = ws.max_row
    max_col = ws.max_column

    if max_row < 2:
        ws.append([None] * max_col)
        max_row = 2

    data_font = Font(name="Segoe UI", size=9)
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 26

    headers = []
    seen_headers = {}
    for c in range(1, max_col + 1):
        raw_val = ws.cell(row=1, column=c).value
        cell_val = str(raw_val if raw_val is not None else f"Col_{c}").strip()
        if not cell_val:
            cell_val = f"Col_{c}"

        if cell_val in seen_headers:
            seen_headers[cell_val] += 1
            unique_h = f"{cell_val}_{seen_headers[cell_val]}"
        else:
            seen_headers[cell_val] = 0
            unique_h = cell_val

        ws.cell(row=1, column=c, value=unique_h)
        headers.append(unique_h)

    col_widths = {}
    for col_idx in range(1, max_col + 1):
        header_val = headers[col_idx - 1]
        max_len = len(header_val)

        for row_idx in range(2, max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            cell.font = data_font
            cell.border = THIN_BORDER

            if val is not None:
                str_val = str(val)
                if len(str_val) > max_len:
                    max_len = len(str_val)

                if isinstance(val, (int, float)):
                    fmt, align = get_excel_format_and_alignment(header_val)
                    if fmt:
                        cell.number_format = fmt
                    cell.alignment = align
                elif isinstance(val, (datetime, pd.Timestamp)):
                    cell.number_format = "yyyy-mm-dd hh:mm"
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        col_letter = get_column_letter(col_idx)
        col_widths[col_letter] = min(max(max_len + 4, 12), 50)

    clean_title = re.sub(r"[^a-zA-Z0-9_]", "_", ws.title).strip("_")
    if not clean_title or clean_title[0].isdigit():
        clean_title = f"T_{clean_title}"
    base_name = f"Tab_{clean_title}"[:200]
    candidate = base_name
    counter = 1
    while candidate.lower() in existing_table_names:
        candidate = f"{base_name}_{counter}"
        counter += 1
    existing_table_names.add(candidate.lower())

    tab_ref = f"A1:{get_column_letter(max_col)}{max_row}"
    tab = Table(displayName=candidate, ref=tab_ref)
    table_style = COLOR_TO_TABLE_STYLE.get(header_fill_hex, "TableStyleMedium9")
    tab.tableStyleInfo = TableStyleInfo(
        name=table_style,
        showRowStripes=True,
        showFirstColumn=False,
        showLastColumn=False,
        showColumnStripes=False,
    )
    ws.add_table(tab)

    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width


def write_and_style_sheet(
    writer, sheet_name, df, header_fill_hex, is_kpi=False, table_names=None
):
    """
    Helper function that writes a DataFrame directly to Excel and styles it in-memory
    within the active writer session, eliminating the need to re-read files from disk.
    """
    if table_names is None:
        table_names = set()
    df.to_excel(writer, sheet_name=sheet_name, index=False)
    ws = writer.sheets[sheet_name]
    style_worksheet(
        ws,
        header_fill_hex=header_fill_hex,
        is_kpi=is_kpi,
        existing_table_names=table_names,
    )


# ==================================================================================================
# MODULE 1: INVENTORY MASTER & BATCH CONTROL ENGINE
# ==================================================================================================
def generate_inventory_workbook(output_dir):
    filename = os.path.join(output_dir, "01_Genius_Inventory_Master.xlsx")
    logger.info(f"Generating Module 1: {filename}...")

    # 1. Active Store Batches with Barcode CTE & Vendor Fallback
    q_batches = """
    WITH PrimaryBarcodes AS (
        SELECT io_itm_id, MIN(io_itm_int) AS Primary_Barcode
        FROM Item_Objects WITH (NOLOCK)
        GROUP BY io_itm_id
    )
    SELECT 
        ics.sto_id AS Store_ID,
        st.sto_name_ar AS Store_Name,
        CASE WHEN ics.sto_id = 47 THEN 'Dispensary Retail (صيدلية)' ELSE 'Consignment / Quarantine (تحت التصريف)' END AS Store_Type,
        ics.itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        pb.Primary_Barcode AS International_Barcode,
        ic.itm_name_ar AS Item_Name_Ar,
        ic.itm_name_en AS Item_Name_En,
        ics.c_id AS Batch_Class_ID,
        icl.itm_batch_no AS Manufacturer_Batch_No,
        ics.itm_qty AS Batch_Quantity,
        u.u_name_ar AS Base_Unit,
        ics.pharm_price AS Unit_Cost_Price,
        ics.sell_price AS Unit_Sell_Price,
        ics.sell_tax AS Unit_Tax,
        CAST(ics.itm_qty * ics.pharm_price AS DECIMAL(18,2)) AS Total_Batch_Cost_Valuation,
        CAST(ics.itm_qty * ics.sell_price AS DECIMAL(18,2)) AS Total_Batch_Retail_Valuation,
        CONVERT(VARCHAR(10), ics.itm_expiry_date, 120) AS Expiry_Date,
        COALESCE(v.ven_name_ar, v_cl.ven_name_ar, 'Unassigned / Opening Stock') AS Preferred_Vendor,
        ics.itm_location AS Rack_Location
    FROM Item_Class_Store ics WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = ics.itm_id
    JOIN Store st WITH (NOLOCK) ON st.sto_id = ics.sto_id
    LEFT JOIN PrimaryBarcodes pb ON pb.io_itm_id = ics.itm_id
    LEFT JOIN Item_Class icl WITH (NOLOCK) ON icl.c_id = ics.c_id AND icl.itm_id = ics.itm_id AND icl.stock_id = ics.sto_id
    LEFT JOIN Units u WITH (NOLOCK) ON u.u_id = ic.itm_unit1
    LEFT JOIN Vendor v WITH (NOLOCK) ON v.ven_id = ics.ven_id
    LEFT JOIN Vendor v_cl WITH (NOLOCK) ON v_cl.ven_id = icl.ven_id
    WHERE ics.itm_qty > 0
    ORDER BY ics.sto_id, ics.itm_id;
    """

    # 2. Consolidated Stock Valuation Across Entire Chain
    q_consolidated = """
    WITH PrimaryBarcodes AS (
        SELECT io_itm_id, MIN(io_itm_int) AS Primary_Barcode
        FROM Item_Objects WITH (NOLOCK)
        GROUP BY io_itm_id
    )
    SELECT 
        ic.itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        pb.Primary_Barcode AS International_Barcode,
        ic.itm_name_ar AS Item_Name_Ar,
        c.com_name_ar AS Manufacturer,
        COUNT(DISTINCT ics.c_id) AS Active_Batch_Count,
        SUM(ics.itm_qty) AS Total_Stock_Units,
        CAST(SUM(ics.itm_qty * ics.pharm_price) / NULLIF(SUM(ics.itm_qty), 0) AS DECIMAL(18,4)) AS True_Weighted_Avg_Cost,
        MAX(ics.sell_price) AS Current_Retail_Price,
        SUM(CAST(ics.itm_qty * ics.pharm_price AS DECIMAL(18,2))) AS Consolidated_Cost_Value,
        SUM(CAST(ics.itm_qty * ics.sell_price AS DECIMAL(18,2))) AS Consolidated_Retail_Value,
        MIN(ics.itm_expiry_date) AS Nearest_Expiry
    FROM Item_Class_Store ics WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = ics.itm_id
    LEFT JOIN PrimaryBarcodes pb ON pb.io_itm_id = ic.itm_id
    LEFT JOIN Company c WITH (NOLOCK) ON c.com_id = ic.com_id
    WHERE ics.itm_qty > 0
    GROUP BY ic.itm_id, ic.itm_code, pb.Primary_Barcode, ic.itm_name_ar, c.com_name_ar
    ORDER BY Consolidated_Cost_Value DESC;
    """

    # 3. Inter-Store Transfers (Between Dispensary 47 and Quarantine 48)
    q_transfers = """
    SELECT 
        sth.stnh_id AS Transfer_ID,
        CASE sth.stnh_flag 
            WHEN 'R' THEN 'Received (مستلم)' 
            WHEN 'S' THEN 'In Transit (مرسل)' 
            WHEN 'A' THEN 'Approved (معتمد)' 
            WHEN 'P' THEN 'Pending (معلق)' 
            ELSE sth.stnh_flag 
        END AS Transfer_Status,
        fs.sto_name_ar AS From_Store,
        ts.sto_name_ar AS To_Store,
        st.st_itm_id AS Item_ID,
        ic.itm_name_ar AS Item_Name,
        st.st_itm_quantity AS Transferred_Qty,
        st.st_itm_cost AS Line_Unit_Cost,
        st.st_itm_sell AS Line_Unit_Sell,
        CAST(st.st_itm_quantity * st.st_itm_cost AS DECIMAL(18,2)) AS Total_Line_Cost,
        CAST(st.st_itm_quantity * st.st_itm_sell AS DECIMAL(18,2)) AS Total_Line_Retail,
        sth.sec_insert_date AS Transfer_Date,
        COALESCE(emp.e_Name, usr.usr_name, CAST(sth.sec_insert_uid AS VARCHAR(50))) AS Operator
    FROM Store_Trans_h sth WITH (NOLOCK)
    JOIN Store_Trans st WITH (NOLOCK) ON st.stnh_id = sth.stnh_id AND st.st_from_store = sth.stnh_f_Sto_id
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = st.st_itm_id
    JOIN Store fs WITH (NOLOCK) ON fs.sto_id = sth.stnh_f_Sto_id
    JOIN Store ts WITH (NOLOCK) ON ts.sto_id = sth.stnh_t_Sto_id
    LEFT JOIN Users usr WITH (NOLOCK) ON usr.usr_e_id = TRY_CAST(sth.sec_insert_uid AS DECIMAL(18,0))
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(sth.sec_insert_uid AS DECIMAL(18,0))
    ORDER BY sth.stnh_id DESC;
    """

    # 4. Physical Stocktaking & Audits (Inventory_h & Inventory_d)
    q_audits = """
    SELECT 
        ih.inv_id AS Audit_Doc_ID,
        st.sto_name_ar AS Store_Name,
        ih.inv_begin_time AS Audit_Timestamp,
        ih.inv_flag AS Doc_State,
        id.invd_itm_id AS Item_ID,
        ic.itm_name_ar AS Item_Name,
        id.invd_c_id AS Batch_ID,
        id.invd_old_qty AS System_Expected_Qty,
        id.invd_new_qty AS Physical_Counted_Qty,
        (id.invd_new_qty - id.invd_old_qty) AS Variance_Qty,
        id.invd_pharm_price AS Unit_Cost,
        CAST((id.invd_new_qty - id.invd_old_qty) * id.invd_pharm_price AS DECIMAL(18,2)) AS Financial_Cost_Impact,
        COALESCE(emp.e_Name, CAST(ih.sec_insert_uid AS VARCHAR(50))) AS Audited_By
    FROM Inventory_h ih WITH (NOLOCK)
    JOIN Inventory_d id WITH (NOLOCK) ON id.invd_inv_id = ih.inv_id
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = id.invd_itm_id
    JOIN Store st WITH (NOLOCK) ON st.sto_id = ih.inv_sto_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(ih.sec_insert_uid AS DECIMAL(18,0))
    ORDER BY ih.inv_id DESC;
    """

    # 5. Quantity Adjustments & Destructions (Item_Quantity_Changes)
    q_adjustments = """
    SELECT 
        iqc.iqc_id AS Adj_ID,
        st.sto_name_ar AS Store_Name,
        iqc.iqc_itm_id AS Item_ID,
        ic.itm_name_ar AS Item_Name,
        dc.dc_name_ar AS Cause_Description,
        iqc.iqc_old_quantity AS Old_Quantity,
        iqc.iqc_new_quantity AS New_Quantity,
        (iqc.iqc_new_quantity - iqc.iqc_old_quantity) AS Quantity_Delta,
        iqc.iqc_pharm_price AS Cost_Price,
        CAST((iqc.iqc_new_quantity - iqc.iqc_old_quantity) * iqc.iqc_pharm_price AS DECIMAL(18,2)) AS Adjustment_Cost_Value,
        iqc.iqc_notes AS Notes,
        iqc.sec_insert_date AS Action_Timestamp
    FROM Item_Quantity_Changes iqc WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = iqc.iqc_itm_id
    LEFT JOIN Store st WITH (NOLOCK) ON st.sto_id = iqc.iqc_store_id
    LEFT JOIN destroy_causes dc WITH (NOLOCK) ON dc.dc_id = iqc.iqc_dc_id
    ORDER BY iqc.iqc_id DESC;
    """

    # 6. Opening Inventory Baseline (Stock_Open & Stock_Open_d)
    q_opening = """
    SELECT 
        so.so_id AS Opening_Doc_ID,
        st.sto_name_ar AS Store_Name,
        sod.itm_id AS Item_ID,
        ic.itm_name_ar AS Item_Name,
        sod.c_id AS Batch_ID,
        sod.itm_qty AS Opening_Quantity,
        sod.pharm_price AS Cost_Price,
        sod.sell_price AS Retail_Price,
        sod.itm_expiry_date AS Expiry_Date,
        so.sec_insert_date AS Setup_Date
    FROM Stock_Open so WITH (NOLOCK)
    JOIN Stock_Open_d sod WITH (NOLOCK) ON sod.so_id = so.so_id
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = sod.itm_id
    JOIN Store st WITH (NOLOCK) ON st.sto_id = so.so_sto_id
    ORDER BY so.so_id DESC;
    """

    # 7. Real-Time Forensic Stock Audit (Trigger: watch_stock on Item_Class_Store)
    q_forensic_stock = """
    SELECT TOP 10000
        w.id AS Audit_Log_ID,
        st.sto_name_ar AS Store_Name,
        w.itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        ic.itm_name_ar AS Item_Name,
        w.c_id AS Batch_Class_ID,
        w.old_qty AS Previous_Quantity,
        w.new_qty AS Adjusted_Quantity,
        CAST(w.new_qty - ISNULL(w.old_qty, 0) AS DECIMAL(18,2)) AS Quantity_Variance,
        w.oper AS Operation_Type,
        w.prog_nm AS Software_Source,
        w.pc_nm AS Client_Machine_Host,
        CONVERT(VARCHAR(10), w.itm_expiry_date, 120) AS Expiry_Date,
        w.insert_dt AS Log_Timestamp
    FROM watch_qty_chng w WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = w.itm_id
    LEFT JOIN Store st WITH (NOLOCK) ON st.sto_id = w.sto_id
    ORDER BY w.id DESC;
    """

    # 8. Shortages Notebook & Replenishment Demand (Shortage)
    q_shortages = """
    SELECT 
        s.str_id AS Shortage_ID,
        st.sto_name_ar AS Store_Name,
        s.str_itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        ic.itm_name_ar AS Item_Name,
        CASE s.str_type 
            WHEN 'S' THEN 'System Suggested Shortage' 
            WHEN 'B' THEN 'Cashier/POS Request' 
            WHEN 'I' THEN 'Inventory Review' 
            ELSE s.str_type 
        END AS Request_Type,
        COALESCE(NULLIF(s.str_requested_qty, 0), ic.itm_default_limit, 1.0) AS Requested_Units,
        CASE s.str_requested WHEN 1 THEN 'Ordered/Fulfilled' ELSE 'Pending Demand' END AS Fulfillment_Status,
        s.str_notes AS Operational_Notes,
        s.sec_insert_date AS Request_Date,
        COALESCE(emp.e_Name, CAST(s.sec_insert_uid AS VARCHAR(50))) AS Requested_By
    FROM Shortage s WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = s.str_itm_id
    LEFT JOIN Store st WITH (NOLOCK) ON st.sto_id = s.sto_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(s.sec_insert_uid AS DECIMAL(18,0))
    ORDER BY s.str_id DESC;
    """

    # 9. Expiry Date Modifications Audit (Item_Expire_Change)
    q_expiry_changes = """
    SELECT 
        iec.iec_id AS Expiry_Audit_ID,
        st.sto_name_ar AS Store_Name,
        iec.iec_itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        ic.itm_name_ar AS Item_Name,
        iec.iec_c_id AS Batch_Class_ID,
        CONVERT(VARCHAR(10), iec.iec_old_expire, 120) AS Old_Expiry_Date,
        CONVERT(VARCHAR(10), iec.iec_new_expire, 120) AS New_Expiry_Date,
        iec.iec_old_qty AS Stock_Quantity_Affected,
        iec.iec_itm_sell_price AS Retail_Price,
        CAST(iec.iec_old_qty * iec.iec_itm_sell_price AS DECIMAL(18,2)) AS Total_Valuation_Impact,
        iec.sec_insert_date AS Change_Timestamp,
        COALESCE(emp.e_Name, CAST(iec.sec_insert_uid AS VARCHAR(50))) AS Modified_By
    FROM Item_Expire_Change iec WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = iec.iec_itm_id
    LEFT JOIN Store st WITH (NOLOCK) ON st.sto_id = iec.iec_sto_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(iec.sec_insert_uid AS DECIMAL(18,0))
    ORDER BY iec.iec_id DESC;
    """

    # 10. Medicine Official Repricing & Revaluation Gains (Item_Changes)
    q_repricing = """
    SELECT 
        ch.ich_id AS Reprice_Audit_ID,
        ch.ich_itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        ic.itm_name_ar AS Item_Name,
        ch.ich_sell_price_old AS Old_Retail_Price,
        ch.ich_sell_price AS New_Official_Price,
        CAST(ch.ich_sell_price - ch.ich_sell_price_old AS DECIMAL(18,2)) AS Price_Delta_Per_Unit,
        ch.ich_pharm_price_old AS Old_Cost,
        ch.ich_pharm_price AS New_Cost,
        ch.ich_qty AS Stock_Quantity_At_Reprice,
        CAST(ch.ich_qty * (ch.ich_sell_price - ch.ich_sell_price_old) AS DECIMAL(18,2)) AS Revaluation_Gain_Retail,
        CAST(ch.ich_qty * (ch.ich_pharm_price - ch.ich_pharm_price_old) AS DECIMAL(18,2)) AS Revaluation_Gain_Cost,
        ch.sec_insert_date AS Repricing_Timestamp,
        COALESCE(emp.e_Name, CAST(ch.sec_insert_uid AS VARCHAR(50))) AS Authorized_By
    FROM Item_Changes ch WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = ch.ich_itm_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(ch.sec_insert_uid AS DECIMAL(18,0))
    WHERE (ch.ich_sell_price - ch.ich_sell_price_old) <> 0
    ORDER BY ch.sec_insert_date DESC;
    """

    # 11. Clinical Indications & Diagnoses Directory (Item_Usage_Causes & Usage_Causes)
    q_usage_indications = """
    SELECT 
        iuc.iuc_itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        ic.itm_name_ar AS Item_Name,
        uc.uc_name_ar AS Medical_Indication_Ar,
        uc.uc_name_en AS Medical_Indication_En
    FROM Item_Usage_Causes iuc WITH (NOLOCK)
    JOIN Usage_Causes uc WITH (NOLOCK) ON uc.uc_id = iuc.iuc_uc_id
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = iuc.iuc_itm_id
    ORDER BY uc.uc_name_ar, iuc.iuc_itm_id;
    """

    # 12. Generic Drug Substitutions & Therapeutic Alternates (Item_Alternates)
    q_substitutions = """
    SELECT 
        ia.itm_id AS Original_Item_ID,
        ic1.itm_code AS Original_Code,
        ic1.itm_name_ar AS Original_Item_Name,
        c1.com_name_ar AS Original_Manufacturer,
        ic1.itm_def_sell_price AS Original_Retail_Price,
        ia.alt_itm_id AS Substitute_Item_ID,
        ic2.itm_code AS Substitute_Code,
        ic2.itm_name_ar AS Substitute_Item_Name,
        c2.com_name_ar AS Substitute_Manufacturer,
        ic2.itm_def_sell_price AS Substitute_Retail_Price,
        CAST(ic2.itm_def_sell_price - ic1.itm_def_sell_price AS DECIMAL(18,2)) AS Price_Difference,
        ia.sec_insert_date AS Linked_Date
    FROM Item_Alternates ia WITH (NOLOCK)
    JOIN Item_Catalog ic1 WITH (NOLOCK) ON ic1.itm_id = ia.itm_id
    JOIN Item_Catalog ic2 WITH (NOLOCK) ON ic2.itm_id = ia.alt_itm_id
    LEFT JOIN Company c1 WITH (NOLOCK) ON c1.com_id = ic1.com_id
    LEFT JOIN Company c2 WITH (NOLOCK) ON c2.com_id = ic2.com_id
    ORDER BY ia.itm_id, ia.alt_itm_id;
    """

    # 13. System Inventory Governance & Operational Parameters (Sys_setting)
    q_sys_governance = """
    SELECT 
        rec_name AS Pharmacy_Legal_Name,
        rec_tel AS Contact_Phone,
        rec_loc AS Registered_Address,
        ret_days_No AS Policy_Return_Days_Allowed,
        CASE prevent_expired_sales WHEN 1 THEN 'Enabled (بيع المنتهي محظور)' ELSE 'Disabled' END AS Block_Expired_Sales,
        CASE stock_avg_cost WHEN 1 THEN 'Enabled (حساب متوسط التكلفة)' ELSE 'Disabled' END AS Weighted_Average_Costing,
        CASE pur_extra_disc_update_stock WHEN 1 THEN 'Enabled (خصم الشراء الإضافي يحدّث التكلفة)' ELSE 'Disabled' END AS Purchase_Disc_Updates_Cost,
        last_month_close_dt AS Last_Monthly_Inventory_Close,
        rec_footer AS Receipt_Terms_Footer
    FROM Sys_setting WITH (NOLOCK);
    """

    # --- SINGLE-PASS IN-MEMORY EXCEL WRITING & FORMATTING ---
    table_names = set()
    fill_hex = "1B4D3E"  # Emerald Green
    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        write_and_style_sheet(
            writer,
            "Current_Store_Batches",
            read_sql_and_enrich(q_batches, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Consolidated_Inventory",
            read_sql_and_enrich(q_consolidated, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Stock_Transfers",
            read_sql_and_enrich(q_transfers, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Physical_Audits_Variance",
            read_sql_and_enrich(q_audits, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Quantity_Adjustments",
            read_sql_and_enrich(q_adjustments, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Opening_Stock",
            read_sql_and_enrich(q_opening, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Inventory_Forensic_Audit",
            read_sql_and_enrich(q_forensic_stock, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Shortages_Catalog_Registry",
            read_sql_and_enrich(q_shortages, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Expiry_Date_Change_Audit",
            read_sql_and_enrich(q_expiry_changes, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Medicine_Repricing_Gains",
            read_sql_and_enrich(q_repricing, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Clinical_Usage_Indications",
            read_sql_and_enrich(q_usage_indications, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Generic_Substitutions_Mapping",
            read_sql_and_enrich(q_substitutions, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "System_Governance_Policies",
            read_sql_and_enrich(q_sys_governance, ENGINE),
            fill_hex,
            table_names=table_names,
        )

    logger.info(
        "Module 1 (01_Genius_Inventory_Master.xlsx) generated successfully with 13 styled sheets."
    )


# ==================================================================================================
# MODULE 2: SALES OPERATIONS & REVENUE ENGINE
# ==================================================================================================
def generate_sales_workbook(output_dir):
    filename = os.path.join(output_dir, "02_Genius_Sales_Operations.xlsx")
    logger.info(f"Generating Module 2: {filename}...")

    # 1. Master Sales Invoices Header (sales_trans_h)
    q_sales_h = """
    SELECT 
        h.sth_id AS Bill_ID,
        st.sto_name_ar AS Store_Name,
        h.sec_insert_date AS Invoice_Date,
        CASE h.bill_typ 
            WHEN 1 THEN 'Cash (نقدي)' 
            WHEN 2 THEN 'Credit (آجل)' 
            WHEN 3 THEN 'Delivery (توصيل)' 
            WHEN 4 THEN 'Visa (فيزا)' 
            WHEN 5 THEN 'Insurance (تأمين)'
            ELSE 'Other' 
        END AS Bill_Type,
        ISNULL(c.cust_name_ar, 'عميل نقدي (Cash Walk-in)') AS Customer_Name,
        ct.ct_name_ar AS Customer_Type,
        h.no_of_items AS Total_Lines,
        h.total_bill AS Gross_Amount,
        h.total_des_mon AS Cash_Discount,
        h.total_dis_per AS Percentage_Discount,
        h.total_bill_after_disc AS Net_Invoice_Amount,
        h.total_tax AS Tax_Amount,
        h.sth_cash AS Paid_Cash,
        h.sth_rest AS Customer_Remaining_Debt,
        h.sth_extra_expenses AS Delivery_Service_Fee,
        COALESCE(emp.e_Name, CAST(h.emp_id AS VARCHAR(50))) AS Cashier_Name,
        h.sth_pc_name AS Station_Hostname
    FROM sales_trans_h h WITH (NOLOCK)
    JOIN Store st WITH (NOLOCK) ON st.sto_id = h.sto_id
    LEFT JOIN Customer c WITH (NOLOCK) ON c.cust_id = h.cust_id
    LEFT JOIN Customer_Type ct WITH (NOLOCK) ON ct.ct_id = c.cust_type
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(h.emp_id AS DECIMAL(18,0))
    WHERE h.sth_flag = 'C'
    ORDER BY h.sth_id DESC;
    """

    # 2. Detailed Dispensed Sales Lines with Optimized Barcode CTE (sales_trans_d)
    q_sales_d = """
    WITH PrimaryBarcodes AS (
        SELECT io_itm_id, MIN(io_itm_int) AS Primary_Barcode
        FROM Item_Objects WITH (NOLOCK)
        GROUP BY io_itm_id
    )
    SELECT 
        d.sth_id AS Bill_ID,
        d.std_id AS Line_No,
        h.sec_insert_date AS Sale_Date,
        d.itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        pb.Primary_Barcode AS International_Barcode,
        ic.itm_name_ar AS Item_Name,
        d.c_id AS Batch_Class_ID,
        d.exp_date AS Batch_Expiry,
        d.qnty AS Dispensed_Quantity,
        ISNULL(d.itm_back, 0) AS Returned_Quantity,
        (d.qnty - ISNULL(d.itm_back, 0)) AS Net_Sold_Units,
        d.itm_unit AS Unit_Level,
        d.itm_sell AS Unit_Sale_Price,
        d.itm_cost AS Unit_Cost_COGS,
        d.itm_dis_per AS Item_Discount_Perc,
        d.itm_dis_mon AS Item_Discount_Cash,
        CAST(
            CASE WHEN (d.qnty - ISNULL(d.itm_back, 0)) > 0 THEN
                (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell * (1.0 - (ISNULL(d.itm_dis_per, 0) / 100.0))) - ISNULL(d.itm_dis_mon, 0))
                * (1.0 - (ISNULL(h.total_dis_per, 0) / 100.0))
                - (ISNULL(h.total_des_mon, 0) * (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell) / NULLIF(h.total_bill, 0)))
            ELSE 0.00 END
        AS DECIMAL(18,2)) AS Net_Line_Revenue,
        CAST((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_cost AS DECIMAL(18,2)) AS Net_Line_COGS,
        CAST(
            CASE WHEN (d.qnty - ISNULL(d.itm_back, 0)) > 0 THEN
                (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell * (1.0 - (ISNULL(d.itm_dis_per, 0) / 100.0))) - ISNULL(d.itm_dis_mon, 0))
                * (1.0 - (ISNULL(h.total_dis_per, 0) / 100.0))
                - (ISNULL(h.total_des_mon, 0) * (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell) / NULLIF(h.total_bill, 0)))
            ELSE 0.00 END
            - ((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_cost)
        AS DECIMAL(18,2)) AS Realized_Gross_Profit
    FROM sales_trans_d d WITH (NOLOCK)
    JOIN sales_trans_h h WITH (NOLOCK) ON h.sth_id = d.sth_id
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = d.itm_id
    LEFT JOIN PrimaryBarcodes pb ON pb.io_itm_id = d.itm_id
    WHERE h.sth_flag = 'C'
    ORDER BY d.sth_id DESC, d.std_id ASC;
    """

    # 3. Cashier Drawer Tender Splits (sales_trans_payment)
    q_payments = """
    SELECT 
        stp.stp_sth_id AS Bill_ID,
        st.sto_name_ar AS Store_Name,
        fcs.fcs_name_ar AS Cash_Drawer_Name,
        stp.stp_value AS Tendered_Value,
        stp.stp_insert_date AS Payment_Timestamp,
        COALESCE(emp.e_Name, CAST(stp.stp_insert_uid AS VARCHAR(50))) AS Processed_By,
        stp.stp_pc_name AS Station
    FROM sales_trans_payment stp WITH (NOLOCK)
    JOIN Store st WITH (NOLOCK) ON st.sto_id = stp.stp_sto_id
    JOIN F_Cash_Store fcs WITH (NOLOCK) ON fcs.fcs_id = stp.stp_fcs_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(stp.stp_insert_uid AS DECIMAL(18,0))
    ORDER BY stp.stp_sth_id DESC;
    """

    # 4. Sales Return & Refund Auditing (sales_return & sales_return_payment)
    q_returns = """
    SELECT 
        sr.sr_id AS Return_Doc_ID,
        sr.sth_id AS Original_Bill_ID,
        h.sec_insert_date AS Original_Sale_Date,
        sr.returned_items_no AS Items_Count_Returned,
        sr.returned_items_value AS Gross_Refund_Value,
        sr.total_back_tax AS Tax_Refunded,
        srp.srp_value AS Actual_Cash_Refunded,
        fcs.fcs_name_ar AS Refunded_From_Drawer,
        sr.sec_insert_date AS Return_Timestamp,
        COALESCE(emp.e_Name, CAST(sr.sec_insert_uid AS VARCHAR(50))) AS Handled_By
    FROM sales_return sr WITH (NOLOCK)
    JOIN sales_trans_h h WITH (NOLOCK) ON h.sth_id = sr.sth_id
    LEFT JOIN sales_return_payment srp WITH (NOLOCK) ON srp.srp_sr_id = sr.sr_id
    LEFT JOIN F_Cash_Store fcs WITH (NOLOCK) ON fcs.fcs_id = srp.srp_fcs_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(sr.sec_insert_uid AS DECIMAL(18,0))
    ORDER BY sr.sr_id DESC;
    """

    # 5. Electronic Visa / MasterCard Settlements (sales_visa_info)
    q_visa = """
    SELECT 
        svi.sth_id AS Bill_ID,
        b.bnk_name_ar AS Terminal_Bank,
        svi.visa_name AS Cardholder_Name,
        svi.visa_no AS Card_Masked_Number,
        svi.exp_date AS Card_Expiry,
        svi.visa_value AS Transaction_Amount,
        svi.get_fees_perc AS Bank_Fee_Percentage,
        svi.sec_insert_date AS Transaction_Date
    FROM sales_visa_info svi WITH (NOLOCK)
    LEFT JOIN Bank b WITH (NOLOCK) ON b.bnk_id = svi.visa_bank_id
    ORDER BY svi.sth_id DESC;
    """

    # 6. Entire Invoice Voiding & Deletions (sales_trans_h_del)
    q_del_bills = """
    SELECT 
        del.sth_id AS Voided_Bill_ID,
        st.sto_name_ar AS Store_Name,
        del.total_bill_net AS Voided_Amount,
        del.no_of_items AS Voided_Lines,
        del.sec_insert_date AS Creation_Time,
        del.sth_delTime AS Deletion_Time,
        del.sth_del_pc_name AS Deletion_Station,
        COALESCE(emp.e_Name, CAST(del.sec_insert_uid AS VARCHAR(50))) AS Original_Cashier
    FROM sales_trans_h_del del WITH (NOLOCK)
    JOIN Store st WITH (NOLOCK) ON st.sto_id = del.sto_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(del.sec_insert_uid AS DECIMAL(18,0))
    ORDER BY del.sth_delTime DESC;
    """

    # 7. Cashier Post-Save Line Item Voiding (sales_trans_d_del)
    q_deleted_items = """
    SELECT 
        del.sth_id AS Voided_Bill_ID,
        del.sec_insert_date AS Original_Creation_Time,
        del.std_delTime AS Deletion_Timestamp,
        del.std_del_pc_name AS Station_Hostname,
        del.itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        ic.itm_name_ar AS Item_Name,
        del.qnty AS Voided_Quantity,
        del.itm_sell AS Unit_Sale_Price,
        del.itm_cost AS Unit_Cost,
        CAST(del.qnty * del.itm_sell AS DECIMAL(18,2)) AS Total_Voided_Value,
        COALESCE(emp.e_Name, CAST(del.sec_insert_uid AS VARCHAR(50))) AS Voided_By_Cashier
    FROM sales_trans_d_del del WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = del.itm_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(del.sec_insert_uid AS DECIMAL(18,0))
    ORDER BY del.std_delTime DESC;
    """

    # 8. Quantity & Price Modifications on Completed Invoices (sales_trans_d_modified_qnty)
    q_qty_modifications = """
    SELECT 
        m.sth_id AS Bill_ID,
        m.itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        ic.itm_name_ar AS Item_Name,
        m.qnty AS Original_Quantity,
        m.itm_back AS Return_Quantity,
        m.new AS Modified_Quantity,
        m.itm_sell AS Original_Selling_Price,
        m.newprice AS Modified_Selling_Price,
        m.itm_unit AS Unit_Level
    FROM sales_trans_d_modified_qnty m WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = m.itm_id
    JOIN sales_trans_h h WITH (NOLOCK) ON h.sth_id = m.sth_id
    ORDER BY m.sth_id DESC;
    """

    # 9. Medical Insurance & Syndicate Contract Claims (Customer_Contracts & sales_trans_h)
    q_insurance_claims = """
    SELECT 
        h.sth_id AS Bill_ID,
        st.sto_name_ar AS Store_Name,
        h.sec_insert_date AS Prescription_Date,
        ISNULL(c.cust_name_ar, 'Walk-in Beneficiary') AS Patient_Name,
        ISNULL(cc.cc_name, 'Syndicate / Corporate Contract') AS Insurance_Syndicate,
        ISNULL(cc.cc_number, 'N/A') AS Policy_Contract_No,
        h.fh_medins_rec_name AS Beneficiary_Member_Name,
        h.fh_medins_ticket_num AS Claim_Ticket_No,
        h.fh_medins_ins_num AS Insurance_Card_No,
        h.total_bill_after_disc AS Total_Prescription_Value,
        ISNULL(cc.cc_pay_perc, 0.0) AS Patient_Copay_Perc,
        CAST(h.total_bill_after_disc - h.fh_company_part AS DECIMAL(18,2)) AS Patient_Copay_Paid,
        h.fh_company_part AS Syndicate_Claim_Liability,
        COALESCE(emp.e_Name, CAST(h.emp_id AS VARCHAR(50))) AS Dispensing_Pharmacist
    FROM sales_trans_h h WITH (NOLOCK)
    JOIN Store st WITH (NOLOCK) ON st.sto_id = h.sto_id
    LEFT JOIN Customer_Contracts cc WITH (NOLOCK) ON cc.cc_id = h.fh_contract_id
    LEFT JOIN Customer c WITH (NOLOCK) ON c.cust_id = h.cust_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(h.emp_id AS DECIMAL(18,0))
    WHERE h.sth_flag = 'C' AND (h.fh_contract_id > 0 OR h.bill_typ = 5 OR h.fh_company_part > 0)
    ORDER BY h.sth_id DESC;
    """

    # 10. Prescribing Doctor & Medical Clinic Analytics (Doctors_specializations & Clinic_specializations)
    q_doctor_referrals = """
    SELECT 
        h.sth_id AS Bill_ID,
        h.sec_insert_date AS Prescription_Date,
        ISNULL(h.fh_medins_Doc_name, 'Direct Walk-in / Self-Medication') AS Referring_Doctor,
        ds.ds_name_ar AS Doctor_Specialization,
        cs.cs_name_ar AS Clinic_Specialization,
        c.cust_name_ar AS Patient_Name,
        h.no_of_items AS Prescribed_Items_Count,
        h.total_bill_after_disc AS Prescription_Total_Value,
        COALESCE(emp.e_Name, CAST(h.emp_id AS VARCHAR(50))) AS Dispensing_Pharmacist
    FROM sales_trans_h h WITH (NOLOCK)
    LEFT JOIN Customer c WITH (NOLOCK) ON c.cust_id = h.cust_id
    LEFT JOIN Doctors_specializations ds WITH (NOLOCK) ON ds.ds_id = h.fh_doc_spec_id
    LEFT JOIN Clinic_specializations cs WITH (NOLOCK) ON cs.cs_id = h.fh_clinic_spec_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(h.emp_id AS DECIMAL(18,0))
    WHERE h.sth_flag = 'C' AND (h.fh_doc_spec_id > 0 OR h.fh_medins_Doc_name IS NOT NULL)
    ORDER BY h.sth_id DESC;
    """

    # 11. Home Delivery Dispatch, Courier Payouts & Zones (Customer_Delivery & Delivery_cost_region)
    q_delivery_dispatch = """
    SELECT 
        h.sth_id AS Bill_ID,
        st.sto_name_ar AS Store_Name,
        h.sec_insert_date AS Order_Time,
        ISNULL(c.cust_name_ar, 'Walk-in / Phone Client') AS Customer_Name,
        c.cust_mobile AS Mobile_Phone,
        ISNULL(sdi.address, cd.cd_address) AS Delivery_Address,
        dcr.dc_region_name_ar AS Delivery_Zone,
        h.total_bill_after_disc AS Order_Value,
        h.sth_extra_expenses AS Delivery_Fee_Charged,
        deliv_emp.e_Name AS Courier_Assigned,
        CASE h.sth_deliv_closed WHEN 1 THEN 'Delivered & Closed' ELSE 'Pending Dispatch' END AS Dispatch_Status,
        COALESCE(emp.e_Name, CAST(h.emp_id AS VARCHAR(50))) AS Cashier_Name
    FROM sales_trans_h h WITH (NOLOCK)
    JOIN Store st WITH (NOLOCK) ON st.sto_id = h.sto_id
    LEFT JOIN Customer c WITH (NOLOCK) ON c.cust_id = h.cust_id
    LEFT JOIN Customer_Delivery cd WITH (NOLOCK) ON cd.cd_cust_id = c.cust_id
    LEFT JOIN Delivery_cost_region dcr WITH (NOLOCK) ON dcr.dc_id = cd.cd_region_id
    LEFT JOIN sales_deliv_info sdi WITH (NOLOCK) ON sdi.sth_id = h.sth_id
    LEFT JOIN Employee deliv_emp WITH (NOLOCK) ON deliv_emp.e_id = sdi.deliv_emp_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(h.emp_id AS DECIMAL(18,0))
    WHERE h.sth_flag = 'C' AND (h.bill_typ = 3 OR h.sth_extra_expenses > 0 OR sdi.sdi_id IS NOT NULL)
    ORDER BY h.sth_id DESC;
    """

    # 12. Negative Stock Discrepancies / Sold Without Stock (sales_not_exist_modified)
    q_sold_without_stock = """
    SELECT 
        sne.sth_id AS Bill_ID,
        sne.std_id AS Line_No,
        sne.sec_insert_date AS Sale_Timestamp,
        sne.itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        ic.itm_name_ar AS Item_Name,
        sne.qnty AS Overdrawn_Sold_Qty,
        sne.itm_sell AS Selling_Price,
        sne.itm_cost AS Standard_Cost,
        CAST(sne.qnty * sne.itm_sell AS DECIMAL(18,2)) AS Revenue_Realized,
        sne.ClassStockAfter AS Negative_Stock_Position_After_Sale,
        COALESCE(emp.e_Name, CAST(sne.sec_insert_uid AS VARCHAR(50))) AS Cashier_Name
    FROM sales_not_exist_modified sne WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = sne.itm_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(sne.sec_insert_uid AS DECIMAL(18,0))
    ORDER BY sne.sth_id DESC;
    """

    # 13. Customer Loyalty Rewards & Redemptions (Customer & sales_cobons)
    q_loyalty_rewards = """
    SELECT 
        c.cust_id AS Customer_ID,
        c.cust_code AS Customer_Code,
        c.cust_name_ar AS Customer_Name,
        c.cust_mobile AS Mobile_Number,
        c.cust_curr_points AS Accumulated_Loyalty_Points,
        c.pont2mony AS Points_Cash_Conversion_Rate,
        CAST(c.cust_curr_points * c.pont2mony AS DECIMAL(18,2)) AS Available_Reward_Value_L_E,
        ISNULL((
            SELECT CAST(SUM(sc_used_value) AS DECIMAL(18,2)) 
            FROM sales_cobons sc WITH (NOLOCK) 
            JOIN sales_trans_h sh WITH (NOLOCK) ON sh.sth_id = sc.sth_id 
            WHERE sh.cust_id = c.cust_id
        ), 0.00) AS Total_Points_Redeemed_Historical,
        c.sec_update_date AS Last_Loyalty_Activity
    FROM Customer c WITH (NOLOCK)
    WHERE c.cust_active = '1' AND (c.cust_curr_points > 0 OR c.cust_discount_perc > 0)
    ORDER BY c.cust_curr_points DESC;
    """

    # 14. Clinical Indication Sales Cross-Tabulation (Primary Indication De-duplicated - BUG 5 FIXED)
    q_clinical_basket_affinity = """
    WITH PrimaryItemIndication AS (
        SELECT 
            iuc.iuc_itm_id,
            MIN(iuc.iuc_uc_id) AS primary_uc_id
        FROM Item_Usage_Causes iuc WITH (NOLOCK)
        GROUP BY iuc.iuc_itm_id
    )
    SELECT 
        ISNULL(uc.uc_name_ar, 'General Pharmacy / OTC (غير مصنف)') AS Clinical_Indication,
        COUNT(DISTINCT d.sth_id) AS Prescriptions_Dispensed_Count,
        COUNT(d.std_id) AS Item_Lines_Count,
        CAST(SUM((d.qnty - ISNULL(d.itm_back, 0)) * (
            CASE d.itm_unit 
                WHEN 3 THEN (1.0 / ISNULL(NULLIF(ic.itm_unit1_unit3, 0), 1.0)) 
                WHEN 2 THEN (1.0 / ISNULL(NULLIF(ic.itm_unit1_unit2, 0), 1.0)) 
                ELSE 1.0 
            END
        )) AS DECIMAL(18,2)) AS Standard_Packs_Sold,
        CAST(SUM(
            CASE WHEN (d.qnty - ISNULL(d.itm_back, 0)) > 0 THEN
                (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell * (1.0 - (ISNULL(d.itm_dis_per, 0)/100.0))) - ISNULL(d.itm_dis_mon, 0))
                * (1.0 - (ISNULL(h.total_dis_per, 0)/100.0))
                - (ISNULL(h.total_des_mon, 0) * (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell) / NULLIF(h.total_bill, 0)))
            ELSE 0 END
        ) AS DECIMAL(18,2)) AS Net_Indication_Revenue,
        CAST(SUM((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_cost) AS DECIMAL(18,2)) AS Total_COGS,
        CAST(SUM(
            CASE WHEN (d.qnty - ISNULL(d.itm_back, 0)) > 0 THEN
                (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell * (1.0 - (ISNULL(d.itm_dis_per, 0)/100.0))) - ISNULL(d.itm_dis_mon, 0))
                * (1.0 - (ISNULL(h.total_dis_per, 0)/100.0))
                - (ISNULL(h.total_des_mon, 0) * (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell) / NULLIF(h.total_bill, 0)))
            ELSE 0 END
        ) - SUM((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_cost) AS DECIMAL(18,2)) AS Gross_Profit_Yield
    FROM sales_trans_d d WITH (NOLOCK)
    JOIN sales_trans_h h WITH (NOLOCK) ON h.sth_id = d.sth_id
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = d.itm_id
    LEFT JOIN PrimaryItemIndication pii ON pii.iuc_itm_id = ic.itm_id
    LEFT JOIN Usage_Causes uc WITH (NOLOCK) ON uc.uc_id = pii.primary_uc_id
    WHERE h.sth_flag = 'C' 
      AND ISNULL(d.itm_nexist, 0) <> 1 
      AND ISNULL(ic.itm_srvc, 0) = 0
    GROUP BY uc.uc_name_ar
    HAVING SUM(d.qnty - ISNULL(d.itm_back, 0)) > 0
    ORDER BY Net_Indication_Revenue DESC;
    """

    # --- SINGLE-PASS IN-MEMORY EXCEL WRITING & FORMATTING ---
    table_names = set()
    fill_hex = "1F4E79"  # Sapphire Blue
    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        write_and_style_sheet(
            writer,
            "Sales_Invoices",
            read_sql_and_enrich(q_sales_h, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Sales_Line_Items",
            read_sql_and_enrich(q_sales_d, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Drawer_Payments",
            read_sql_and_enrich(q_payments, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Sales_Returns",
            read_sql_and_enrich(q_returns, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Card_Settlements",
            read_sql_and_enrich(q_visa, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Deleted_Bills_Audit",
            read_sql_and_enrich(q_del_bills, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Cashier_Voided_Lines_Audit",
            read_sql_and_enrich(q_deleted_items, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Sales_Quantity_Modifications",
            read_sql_and_enrich(q_qty_modifications, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Insurance_Contract_Claims",
            read_sql_and_enrich(q_insurance_claims, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Doctor_Prescription_Analytics",
            read_sql_and_enrich(q_doctor_referrals, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Home_Delivery_Dispatch",
            read_sql_and_enrich(q_delivery_dispatch, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Sold_Without_Stock_Audit",
            read_sql_and_enrich(q_sold_without_stock, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Customer_Loyalty_Rewards",
            read_sql_and_enrich(q_loyalty_rewards, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Clinical_Indication_Sales",
            read_sql_and_enrich(q_clinical_basket_affinity, ENGINE),
            fill_hex,
            table_names=table_names,
        )

    logger.info(
        "Module 2 (02_Genius_Sales_Operations.xlsx) generated successfully with 14 styled sheets."
    )


# ==================================================================================================
# MODULE 3: PROCUREMENT & VENDOR SUPPLY CHAIN
# ==================================================================================================
def generate_purchases_workbook(output_dir):
    filename = os.path.join(output_dir, "03_Genius_Purchasing_Operations.xlsx")
    logger.info(f"Generating Module 3: {filename}...")

    # 1. Purchase Invoices Header (pur_trans_h)
    q_pur_h = """
    SELECT 
        pth.pth_id AS Purchase_Bill_ID,
        st.sto_name_ar AS Receiving_Store,
        v.ven_name_ar AS Vendor_Name,
        pth.ven_bill_no AS Vendor_Invoice_Ref,
        pth.ven_bill_date AS Vendor_Bill_Date,
        pth.no_of_items AS Total_Lines,
        pth.total_bill AS Gross_Purchase_Total,
        pth.total_dis_per AS Header_Discount_Perc,
        pth.total_des_mon AS Header_Discount_Cash,
        pth.total_after_back AS Net_Purchase_Total,
        pth.pth_paid AS Paid_Amount,
        (pth.total_after_back - pth.pth_paid) AS Remaining_Vendor_Liability,
        pth.pth_source_tax AS Withholding_Tax_Rate,
        COALESCE(emp.e_Name, CAST(pth.sec_insert_uid AS VARCHAR(50))) AS Created_By,
        pth.sec_insert_date AS System_Entry_Date
    FROM pur_trans_h pth WITH (NOLOCK)
    JOIN Store st WITH (NOLOCK) ON st.sto_id = pth.sto_id
    JOIN Vendor v WITH (NOLOCK) ON v.ven_id = pth.ven_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(pth.sec_insert_uid AS DECIMAL(18,0))
    WHERE pth.bill_status <> 1
    ORDER BY pth.pth_id DESC;
    """

    # 2. Detailed Purchase Line Items with Barcode CTE (pur_trans_d)
    q_pur_d = """
    WITH PrimaryBarcodes AS (
        SELECT io_itm_id, MIN(io_itm_int) AS Primary_Barcode
        FROM Item_Objects WITH (NOLOCK)
        GROUP BY io_itm_id
    )
    SELECT 
        ptd.pth_id AS Purchase_Bill_ID,
        ptd.ptd_id AS Line_No,
        pth.sec_insert_date AS Purchase_Date,
        v.ven_name_ar AS Vendor_Name,
        ptd.itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        pb.Primary_Barcode AS International_Barcode,
        ic.itm_name_ar AS Item_Name,
        ptd.c_id AS Assigned_Batch_ID,
        ptd.ptd_batch AS Vendor_Lot_Batch,
        CONVERT(VARCHAR(10), ptd.exp_date, 120) AS Expiry_Date,
        ptd.qnty AS Purchased_Qty,
        ptd.bonus AS Bonus_Free_Qty,
        ptd.itm_pur_price AS Unit_Wholesale_Price,
        ptd.itm_extra_dis AS Line_Discount_Perc,
        ptd.itm_tax_price AS Unit_Tax_Amount,
        ptd.itm_cost AS Effective_Unit_Cost,
        ptd.itm_sell AS Expected_Retail_Price,
        CAST(
            (ptd.qnty * (ptd.itm_pur_price * (1.0 - (ISNULL(ptd.itm_extra_dis,0)/100.0)) + ptd.itm_tax_price))
            * (1.0 - (ISNULL(pth.total_dis_per,0)/100.0))
        AS DECIMAL(18,2)) AS Net_Line_Purchase_Cost,
        ptd.itm_back_qty AS Returned_Qty
    FROM pur_trans_d ptd WITH (NOLOCK)
    JOIN pur_trans_h pth WITH (NOLOCK) ON pth.pth_id = ptd.pth_id
    JOIN Vendor v WITH (NOLOCK) ON v.ven_id = pth.ven_id
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = ptd.itm_id
    LEFT JOIN PrimaryBarcodes pb ON pb.io_itm_id = ptd.itm_id
    WHERE pth.bill_status <> 1
    ORDER BY ptd.pth_id DESC, ptd.ptd_id ASC;
    """

    # 3. General Vendor Return Vouchers (Item_back_h & Item_back_d)
    q_pur_returns = """
    SELECT 
        ibh.ibh_id AS Return_Doc_ID,
        st.sto_name_ar AS Store_Name,
        v.ven_name_ar AS Vendor_Name,
        ibh.ibh_total_value AS Total_Return_Value,
        ibd.ibd_itm_id AS Item_ID,
        ic.itm_name_ar AS Item_Name,
        ibd.ibd_qty AS Returned_Quantity,
        ibd.ibd_pharm_price AS Refund_Unit_Cost,
        ibd.ibd_tax_price AS Refund_Unit_Tax,
        CAST(ibd.ibd_qty * (ibd.ibd_pharm_price + ibd.ibd_tax_price) AS DECIMAL(18,2)) AS Total_Refund_Credit,
        CONVERT(VARCHAR(10), ibd.ibd_exp_date, 120) AS Expiry_Date,
        ibh.sec_insert_date AS Return_Date,
        COALESCE(emp.e_Name, CAST(ibh.sec_insert_uid AS VARCHAR(50))) AS Returned_By
    FROM Item_back_h ibh WITH (NOLOCK)
    JOIN Item_back_d ibd WITH (NOLOCK) ON ibd.ibd_ibh_id = ibh.ibh_id
    JOIN Vendor v WITH (NOLOCK) ON v.ven_id = ibh.ibh_ven_id
    JOIN Store st WITH (NOLOCK) ON st.sto_id = ibh.ibh_sto_id
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = ibd.ibd_itm_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(ibh.sec_insert_uid AS DECIMAL(18,0))
    ORDER BY ibh.ibh_id DESC;
    """

    # 4. Direct Bill-Linked Purchase Returns (pur_trans_d with itm_back > 0)
    q_direct_bill_returns = """
    SELECT 
        ptd.pth_id AS Purchase_Bill_ID,
        ptd.ptd_id AS Line_No,
        pth.ven_bill_no AS Vendor_Invoice_Ref,
        v.ven_name_ar AS Vendor_Name,
        ptd.itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        ic.itm_name_ar AS Item_Name,
        ptd.itm_back_qty AS Returned_Quantity,
        ptd.itm_back_pharm AS Unit_Return_Price,
        ptd.itm_back_tax AS Unit_Return_Tax,
        ptd.itm_back_bonus AS Returned_Bonus_Units,
        CAST((ptd.itm_back_qty * (ptd.itm_back_pharm + ptd.itm_back_tax)) AS DECIMAL(18,2)) AS Total_Invoice_Credit_Claimed,
        ptd.sec_update_date AS Return_Action_Date,
        COALESCE(emp.e_Name, CAST(ptd.sec_update_uid AS VARCHAR(50))) AS Handled_By
    FROM pur_trans_d ptd WITH (NOLOCK)
    JOIN pur_trans_h pth WITH (NOLOCK) ON pth.pth_id = ptd.pth_id
    JOIN Vendor v WITH (NOLOCK) ON v.ven_id = pth.ven_id
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = ptd.itm_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(ptd.sec_update_uid AS DECIMAL(18,0))
    WHERE ptd.itm_back > 0 AND pth.bill_status <> 1
    ORDER BY ptd.pth_id DESC;
    """

    # 5. Historical Vendor Price Quotes & Discount Terms (Item_Vendor_history)
    q_vendor_pricing = """
    SELECT 
        ivh.itm_id AS Item_ID,
        ic.itm_name_ar AS Item_Name,
        v.ven_name_ar AS Vendor_Name,
        ivh.itm_pharm_price AS Quoted_Pharm_Price,
        ivh.itm_sell_price AS Retail_Price,
        ivh.itm_disc1 AS Primary_Discount,
        ivh.itm_disc2 AS Secondary_Discount,
        ivh.itm_disc3 AS Special_Discount,
        ivh.itm_pricing_date AS Price_Recorded_Date
    FROM Item_Vendor_history ivh WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = ivh.itm_id
    JOIN Vendor v WITH (NOLOCK) ON v.ven_id = ivh.ven_id
    ORDER BY ivh.itm_pricing_date DESC;
    """

    # 6. Vendor Ledger Balances & Aging Summary (Vendor)
    q_vendor_ledger = """
    SELECT 
        v.ven_id AS Vendor_ID,
        v.ven_code AS Vendor_Code,
        v.ven_name_ar AS Vendor_Name_Ar,
        v.ven_tel AS Phone_Contact,
        v.ven_max_credit AS Credit_Limit,
        v.ven_start_debit AS Initial_Debit,
        v.ven_start_credit AS Initial_Credit,
        v.ven_current_credit AS Current_Outstanding_Balance,
        CASE 
            WHEN v.ven_current_credit > 0 THEN 'We Owe Vendor - Accounts Payable (دائن - التزام علينا)' 
            WHEN v.ven_current_credit < 0 THEN 'Vendor Owes Pharmacy - Advance/Return (مدين - رصيد لنا)' 
            ELSE 'Settled (خالص)' 
        END AS Balance_Status,
        v.sec_update_date AS Last_Transaction_Date
    FROM Vendor v WITH (NOLOCK)
    WHERE v.ven_active = '1'
    ORDER BY ABS(v.ven_current_credit) DESC;
    """

    # 7. Forensic Vendor Credit Change Audit Trail (vendor_credit_chng)
    q_vendor_credit_audit = """
    SELECT 
        vc.vc_id AS Audit_ID,
        v.ven_name_ar AS Vendor_Name,
        vc.vc_old_credit AS Previous_Balance,
        vc.vc_new_credit AS Updated_Balance,
        CAST(vc.vc_new_credit - vc.vc_old_credit AS DECIMAL(18,2)) AS Balance_Delta,
        vc.vc_operation AS Triggering_Operation,
        vc.vc_prog_name AS Calling_Application,
        vc.vc_pc_name AS Station_Hostname,
        vc.vc_insert_date AS Audit_Timestamp
    FROM vendor_credit_chng vc WITH (NOLOCK)
    JOIN Vendor v WITH (NOLOCK) ON v.ven_id = vc.vc_cust_id
    ORDER BY vc.vc_id DESC;
    """

    # 8. Distributor Price & Discount Arbitrage Optimizer (Item_Vendor)
    q_distributor_arbitrage = """
    SELECT 
        iv.itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        ic.itm_name_ar AS Item_Name,
        v.ven_name_ar AS Distributor_Name,
        iv.itm_pharm_price AS Quoted_Pharm_Price,
        iv.itm_disc1 AS Base_Discount_Perc,
        iv.itm_disc2 AS Extra_Discount_Perc,
        iv.itm_disc3 AS Special_Discount_Perc,
        CAST(iv.itm_disc1 + iv.itm_disc2 + iv.itm_disc3 AS DECIMAL(18,2)) AS Cumulative_Discount_Rate,
        CAST(iv.itm_pharm_price * (1.0 - ((iv.itm_disc1 + iv.itm_disc2 + iv.itm_disc3) / 100.0)) AS DECIMAL(18,2)) AS Net_Procurement_Unit_Cost,
        iv.itm_sell_price AS Retail_Selling_Price,
        CAST(iv.itm_sell_price - (iv.itm_pharm_price * (1.0 - ((iv.itm_disc1 + iv.itm_disc2 + iv.itm_disc3) / 100.0))) AS DECIMAL(18,2)) AS Profit_Spread,
        CASE WHEN iv.avail_flag = '1' THEN 'Available' ELSE 'Check Availability' END AS Distributor_Availability,
        iv.last_price_update AS Quote_Date
    FROM Item_Vendor iv WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = iv.itm_id
    JOIN Vendor v WITH (NOLOCK) ON v.ven_id = iv.ven_id
    WHERE v.ven_active = '1' AND (iv.itm_disc1 + iv.itm_disc2 + iv.itm_disc3) > 0
    ORDER BY Cumulative_Discount_Rate DESC;
    """

    # 9. Master Vendor Directory & B2B Configuration (Vendor & Vendor_Extn)
    q_vendor_directory = """
    SELECT 
        v.ven_id AS Vendor_ID,
        v.ven_code AS Vendor_Code,
        v.ven_name_ar AS Vendor_Name_Ar,
        v.ven_name_en AS Vendor_Name_En,
        v.ven_tel AS Primary_Telephone,
        ve.ve_tel AS Secondary_Telephone,
        ve.ve_contact_name AS Contact_Person,
        v.ven_mgr_name AS Manager_Name,
        v.ven_address AS Registered_Address,
        v.ven_max_credit AS Authorized_Credit_Limit,
        v.ven_b2b_name AS B2B_Platform_Name,
        v.ven_pharm_code AS Pharmacy_Code_At_Vendor,
        CASE v.ven_active WHEN '1' THEN 'Active' ELSE 'Inactive' END AS Account_Status
    FROM Vendor v WITH (NOLOCK)
    LEFT JOIN Vendor_Extn ve WITH (NOLOCK) ON ve.ve_ven_id = v.ven_id
    ORDER BY v.ven_id;
    """

    # 10. Wholesaler Benchmark Catalog Registry (Vendors_Catalogs)
    q_ucp_catalog = """
    SELECT 
        vc.vc_code AS Distributor_Item_Code,
        vc.vc_company AS Wholesaler_Name,
        vc.vc_item_name AS Catalog_Trade_Name,
        vc.vc_pharm_price AS Wholesale_Price,
        vc.vc_sell_price AS Official_Retail_Price,
        vc.vc_tax AS Tax_Included,
        vc.sec_insert_date AS Catalog_Date
    FROM Vendors_Catalogs vc WITH (NOLOCK)
    ORDER BY vc.vc_company, vc.vc_item_name;
    """

    # 11. Supplier Return Early-Warning Radar (BUG 2 FIXED: Fallback to Item_Class.ven_id)
    q_supplier_return_radar = """
    SELECT 
        ics.sto_id AS Store_ID,
        st.sto_name_ar AS Store_Name,
        COALESCE(v.ven_name_ar, v_cl.ven_name_ar, 'Distributor Unassigned / Opening Stock') AS Vendor_Name,
        ic.itm_id AS Item_ID,
        ic.itm_code AS Barcode,
        ic.itm_name_ar AS Item_Name,
        ics.c_id AS Batch_ID,
        ics.itm_qty AS Stock_Quantity,
        ics.pharm_price AS Unit_Cost,
        CAST(ics.itm_qty * ics.pharm_price AS DECIMAL(18,2)) AS Reclaimable_Vendor_Credit,
        CONVERT(VARCHAR(10), ics.itm_expiry_date, 120) AS Expiry_Date,
        DATEDIFF(DAY, GETDATE(), ics.itm_expiry_date) AS Days_Until_Expiration,
        ISNULL(COALESCE(v.ven_return_interval, v_cl.ven_return_interval), 30) AS Return_Window_Days,
        CASE 
            WHEN DATEDIFF(DAY, GETDATE(), ics.itm_expiry_date) <= ISNULL(COALESCE(v.ven_return_interval, v_cl.ven_return_interval), 30) 
            THEN 'CRITICAL: Return to Vendor Immediately'
            WHEN DATEDIFF(DAY, GETDATE(), ics.itm_expiry_date) <= (ISNULL(COALESCE(v.ven_return_interval, v_cl.ven_return_interval), 30) + 30)
            THEN 'WARNING: Approaching Return Deadline'
            ELSE 'Within Normal Shelf Life'
        END AS Return_Action_Priority
    FROM Item_Class_Store ics WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = ics.itm_id
    JOIN Store st WITH (NOLOCK) ON st.sto_id = ics.sto_id
    LEFT JOIN Item_Class icl WITH (NOLOCK) 
        ON icl.c_id = ics.c_id AND icl.itm_id = ics.itm_id AND icl.stock_id = ics.sto_id
    LEFT JOIN Vendor v WITH (NOLOCK) ON v.ven_id = ics.ven_id
    LEFT JOIN Vendor v_cl WITH (NOLOCK) ON v_cl.ven_id = icl.ven_id
    WHERE ics.itm_qty > 0 
      AND ISNULL(ic.itm_nopurreturn, 0) = 0
      AND ics.itm_expiry_date BETWEEN GETDATE() AND DATEADD(DAY, 120, GETDATE())
    ORDER BY Days_Until_Expiration ASC;
    """

    # 12. Replenishment Multi-Distributor Arbitrage Optimizer (Shortage & Item_Vendor)
    q_replenish_arbitrage = """
    SELECT 
        s.str_itm_id AS Shortage_Item_ID,
        ic.itm_code AS Barcode,
        ic.itm_name_ar AS Item_Name,
        COALESCE(NULLIF(s.str_requested_qty, 0), ic.itm_default_limit, 1.0) AS Units_Required,
        v.ven_name_ar AS Competing_Distributor,
        iv.itm_pharm_price AS Official_Pharm_Price,
        iv.itm_disc1 AS Base_Discount_Perc,
        iv.itm_disc2 AS Extra_Discount_Perc,
        iv.itm_disc3 AS Special_Discount_Perc,
        CAST(iv.itm_disc1 + iv.itm_disc2 + iv.itm_disc3 AS DECIMAL(18,2)) AS Total_Offered_Discount_Perc,
        CAST(iv.itm_pharm_price * (1.0 - ((iv.itm_disc1 + iv.itm_disc2 + iv.itm_disc3) / 100.0)) AS DECIMAL(18,2)) AS Net_Procurement_Unit_Cost,
        CAST(COALESCE(NULLIF(s.str_requested_qty, 0), ic.itm_default_limit, 1.0) * (iv.itm_pharm_price * (1.0 - ((iv.itm_disc1 + iv.itm_disc2 + iv.itm_disc3) / 100.0))) AS DECIMAL(18,2)) AS Estimated_PO_Outlay,
        RANK() OVER (PARTITION BY s.str_itm_id ORDER BY (iv.itm_disc1 + iv.itm_disc2 + iv.itm_disc3) DESC) AS Distributor_Margin_Rank
    FROM Shortage s WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = s.str_itm_id
    JOIN Item_Vendor iv WITH (NOLOCK) ON iv.itm_id = s.str_itm_id
    JOIN Vendor v WITH (NOLOCK) ON v.ven_id = iv.ven_id
    WHERE s.str_requested = 0 AND v.ven_active = '1'
    ORDER BY s.str_itm_id, Distributor_Margin_Rank ASC;
    """

    # --- SINGLE-PASS IN-MEMORY EXCEL WRITING & FORMATTING ---
    table_names = set()
    fill_hex = "4A154B"  # Deep Burgundy
    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        write_and_style_sheet(
            writer,
            "Purchase_Bills",
            read_sql_and_enrich(q_pur_h, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Purchase_Line_Items",
            read_sql_and_enrich(q_pur_d, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Vendor_Returns",
            read_sql_and_enrich(q_pur_returns, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Direct_Bill_Purchase_Returns",
            read_sql_and_enrich(q_direct_bill_returns, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Vendor_Pricing_History",
            read_sql_and_enrich(q_vendor_pricing, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Vendor_Ledger_Balances",
            read_sql_and_enrich(q_vendor_ledger, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Vendor_Credit_Audit_Log",
            read_sql_and_enrich(q_vendor_credit_audit, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Distributor_Price_Arbitrage",
            read_sql_and_enrich(q_distributor_arbitrage, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Master_Vendor_Directory",
            read_sql_and_enrich(q_vendor_directory, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Wholesaler_Master_Catalog",
            read_sql_and_enrich(q_ucp_catalog, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Supplier_Return_Early_Warning",
            read_sql_and_enrich(q_supplier_return_radar, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Replenishment_Arbitrage_Engine",
            read_sql_and_enrich(q_replenish_arbitrage, ENGINE),
            fill_hex,
            table_names=table_names,
        )

    logger.info(
        "Module 3 (03_Genius_Purchasing_Operations.xlsx) generated successfully with 12 styled sheets."
    )


# ==================================================================================================
# MODULE 4: TREASURY & GENERAL FINANCIAL LEDGER
# ==================================================================================================
def generate_financial_workbook(output_dir):
    filename = os.path.join(output_dir, "04_Genius_Financial_Ledger.xlsx")
    logger.info(f"Generating Module 4: {filename}...")

    # 1. Cash Safes, Drawers & Commercial Bank Accounts (F_Cash_Store)
    q_cash_stores = """
    SELECT 
        fcs.fcs_id AS Fund_ID,
        fcs.fcs_code AS Code,
        fcs.fcs_name_ar AS Account_Title,
        CASE fcs.fcs_type 
            WHEN '1' THEN 'POS Cashier Drawer (درج كاشير)' 
            WHEN '2' THEN 'Main Treasury/Vault (خزينة رئيسية)' 
            WHEN '3' THEN 'Bank Account (حساب بنكي)' 
            ELSE 'Other Fund' 
        END AS Account_Type,
        fcs.fcs_current_balance AS Current_Live_Balance,
        fcs.fcs_account_number AS Bank_Account_No,
        b.bnk_name_ar AS Associated_Bank,
        fcs.fcs_computer_name AS Bound_Workstation,
        fcs.sec_update_date AS Last_Balance_Update
    FROM F_Cash_Store fcs WITH (NOLOCK)
    LEFT JOIN Bank b WITH (NOLOCK) ON b.bnk_id = fcs.fcs_bank_id
    WHERE fcs.fcs_active = 1
    ORDER BY fcs.fcs_type, fcs.fcs_id;
    """

    # 2. General Treasury Cash Receipts & Disbursements (F_Transaction_Header)
    q_treasury_trans = """
    SELECT 
        fh.fh_id AS Entry_ID,
        fh.fh_actual_date AS Transaction_Date,
        CASE fh.fh_trans_type 
            WHEN '1' THEN 'Cash Movement (نقدية)' 
            WHEN '2' THEN 'Cheque/Commercial Paper (أوراق مالية)' 
            ELSE 'Other' 
        END AS Instrument_Type,
        fh.fh_value AS Amount,
        CASE fh.fh_from_type 
            WHEN '1' THEN 'Customer (عميل)' 
            WHEN '2' THEN 'Vendor (مورد)' 
            WHEN '3' THEN 'Cash Safe/Drawer (خزينة/درج)' 
            ELSE 'Account/Other' 
        END AS Source_Entity_Type,
        CASE 
            WHEN fh.fh_from_type = '1' THEN ISNULL(c_from.cust_name_ar, 'Cash Customer')
            WHEN fh.fh_from_type = '2' THEN v_from.ven_name_ar
            WHEN fh.fh_from_type = '3' THEN fcs_from.fcs_name_ar
            ELSE 'General GL Account'
        END AS Source_Entity_Name,
        CASE fh.fh_to_type 
            WHEN '1' THEN 'Customer (عميل)' 
            WHEN '2' THEN 'Vendor (مورد)' 
            WHEN '3' THEN 'Cash Safe/Drawer (خزينة/درج)' 
            ELSE 'Account/Other' 
        END AS Destination_Entity_Type,
        CASE 
            WHEN fh.fh_to_type = '1' THEN ISNULL(c_to.cust_name_ar, 'Cash Customer')
            WHEN fh.fh_to_type = '2' THEN v_to.ven_name_ar
            WHEN fh.fh_to_type = '3' THEN fcs_to.fcs_name_ar
            ELSE 'General GL Account'
        END AS Destination_Entity_Name,
        fh.fh_notes AS Operational_Narrative,
        fh.fh_computer AS Executing_Machine,
        COALESCE(emp.e_Name, CAST(fh.sec_insert_uid AS VARCHAR(50))) AS Operator_Name
    FROM F_Transaction_Header fh WITH (NOLOCK)
    LEFT JOIN Customer c_from WITH (NOLOCK) ON c_from.cust_id = fh.fh_from_id AND fh.fh_from_type = '1'
    LEFT JOIN Vendor v_from WITH (NOLOCK) ON v_from.ven_id = fh.fh_from_id AND fh.fh_from_type = '2'
    LEFT JOIN F_Cash_Store fcs_from WITH (NOLOCK) ON fcs_from.fcs_id = fh.fh_from_id AND fh.fh_from_type = '3'
    LEFT JOIN Customer c_to WITH (NOLOCK) ON c_to.cust_id = fh.fh_to_id AND fh.fh_to_type = '1'
    LEFT JOIN Vendor v_to WITH (NOLOCK) ON v_to.ven_id = fh.fh_to_id AND fh.fh_to_type = '2'
    LEFT JOIN F_Cash_Store fcs_to WITH (NOLOCK) ON fcs_to.fcs_id = fh.fh_to_id AND fh.fh_to_type = '3'
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(fh.sec_insert_uid AS DECIMAL(18,0))
    ORDER BY fh.fh_id DESC;
    """

    # 3. Commercial Papers & Cheques Portfolio (F_Transaction_Header_Ext)
    q_cheques = """
    SELECT 
        fhe.fhe_id AS Paper_Registry_ID,
        fhe.fh_id AS Treasury_Header_ID,
        fp.fp_name_ar AS Paper_Type,
        fhe.fhe_paper_number AS Cheque_Paper_Number,
        fhe.fhe_paper_create_date AS Issue_Date,
        fhe.fhe_paper_valid_date AS Due_Maturity_Date,
        b.bnk_name_ar AS Deposit_Bank,
        CASE fhe.fhe_paper_status 
            WHEN '0' THEN 'Pending In-Portfolio (تحت التحصيل)' 
            WHEN '1' THEN 'Deposited/Collected (تم التحصيل)' 
            WHEN '2' THEN 'Rejected/Bounced (مرفوض/مرتد)' 
            WHEN '3' THEN 'Cancelled (ملغي)' 
            ELSE 'Unknown' 
        END AS Cheque_Status,
        fh.fh_value AS Face_Value,
        fhe.fhe_expenses AS Bank_Clearance_Expenses
    FROM F_Transaction_Header_Ext fhe WITH (NOLOCK)
    JOIN F_Transaction_Header fh WITH (NOLOCK) ON fh.fh_id = fhe.fh_id
    LEFT JOIN F_Papers fp WITH (NOLOCK) ON fp.fp_id = fhe.fhe_paper_type
    LEFT JOIN Bank b WITH (NOLOCK) ON b.bnk_id = fhe.fhe_paper_deposit_bank_id
    ORDER BY fhe.fhe_id DESC;
    """

    # 4. Automated Double-Entry General Ledger Journals (F_Auto_Doc_h & F_Auto_Doc_d)
    q_gl_engine = """
    SELECT TOP 50000
        fah.fah_id AS Auto_Doc_ID,
        st.sto_name_ar AS Branch_Name,
        fah.fah_type AS System_Trigger,
        fah.fah_form_text AS Originating_Form,
        fah.fah_description AS Journal_Narrative,
        fah.fah_total_value AS Total_Document_Value,
        fad.fad_group_id AS Leg_Group,
        fad.fad_D_Account_name AS Debit_Account,
        fad.fad_D_object AS Debit_Target_Type,
        fad.fad_C_Account_name AS Credit_Account,
        fad.fad_C_object AS Credit_Target_Type,
        fad.fad_value AS Leg_Amount,
        fah.sec_insert_date AS Posting_Timestamp,
        fah.fah_computer_name AS Machine
    FROM F_Auto_Doc_h fah WITH (NOLOCK)
    JOIN F_Auto_Doc_d fad WITH (NOLOCK) ON fad.fad_fah_id = fah.fah_id
    JOIN Store st WITH (NOLOCK) ON st.sto_id = fah.fah_sto_id
    ORDER BY fah.fah_id DESC, fad.fad_group_id ASC;
    """

    # 5. Customer & Vendor Debit / Credit Adjustment Notices (F_Notices)
    q_notices = """
    SELECT 
        fn.fn_id AS Notice_ID,
        CASE fn.fn_type 
            WHEN '1' THEN 'Addition Notice - Credit (إشعار إضافة)' 
            WHEN '2' THEN 'Discount Notice - Debit (إشعار خصم)' 
            ELSE 'Other' 
        END AS Notice_Type,
        CASE fn.fn_who_type 
            WHEN '1' THEN 'Vendor (مورد)' 
            WHEN '2' THEN 'Customer (عميل)' 
            ELSE 'Party' 
        END AS Party_Type,
        fn.fn_who_id AS Party_ID,
        fnc.fnc_name_ar AS Reason_Category,
        fn.fn_value AS Notice_Amount,
        fn.fn_notes AS Explanation,
        fn.sec_insert_date AS Issued_Timestamp,
        COALESCE(emp.e_Name, CAST(fn.sec_insert_uid AS VARCHAR(50))) AS Officer_Name
    FROM F_Notices fn WITH (NOLOCK)
    LEFT JOIN F_Notices_Causes fnc WITH (NOLOCK) ON fnc.fnc_id = fn.fn_cause_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(fn.sec_insert_uid AS DECIMAL(18,0))
    ORDER BY fn.fn_id DESC;
    """

    # 6. Customer Credit Accounts & Receivables Position (Customer)
    q_cust_debt = """
    SELECT 
        c.cust_id AS Customer_ID,
        c.cust_code AS Customer_Code,
        c.cust_name_ar AS Customer_Name,
        c.cust_mobile AS Phone_Number,
        c.cust_max_credit AS Authorized_Credit_Limit,
        c.cust_start_debit AS Initial_Debit,
        c.cust_start_credit AS Initial_Credit,
        c.cust_current_credit AS Current_Balance,
        CASE 
            WHEN c.cust_current_credit > 0 THEN 'Customer Owes Pharmacy - Receivable (مدين)' 
            WHEN c.cust_current_credit < 0 THEN 'Pharmacy Owes Customer - Advance (دائن)' 
            ELSE 'Settled (خالص)' 
        END AS Position_Status,
        c.sec_update_date AS Last_Credit_Change
    FROM Customer c WITH (NOLOCK)
    WHERE c.cust_active = '1'
    ORDER BY c.cust_current_credit DESC;
    """

    # 7. Forensic Customer Credit Change Audit Trail (customer_credit_chng)
    q_customer_credit_audit = """
    SELECT 
        cc.cc_id AS Audit_ID,
        c.cust_name_ar AS Customer_Name,
        cc.cc_old_credit AS Previous_Debt_Balance,
        cc.cc_new_credit AS Updated_Debt_Balance,
        CAST(cc.cc_new_credit - cc.cc_old_credit AS DECIMAL(18,2)) AS Balance_Delta,
        cc.cc_operation AS Operation_Event,
        cc.cc_prog_name AS Originating_Software,
        cc.cc_pc_name AS Workstation_Hostname,
        cc.cc_insert_date AS Action_Timestamp
    FROM customer_credit_chng cc WITH (NOLOCK)
    JOIN Customer c WITH (NOLOCK) ON c.cust_id = cc.cc_cust_id
    ORDER BY cc.cc_id DESC;
    """

    # 8. Hierarchical Chart of Accounts Tree (F_Account_Chart)
    q_chart_of_accounts = """
    SELECT 
        ac.ac_id AS Account_ID,
        ac.ac_code AS Account_Code,
        ac.ac_name_ar AS Account_Name_Ar,
        ac.ac_name_en AS Account_Name_En,
        ac.ac_parent AS Parent_Account_ID,
        parent.ac_name_ar AS Parent_Account_Name,
        ISNULL(ac.ac_type, 'General Ledger') AS Account_Classification,
        ac.ac_start_debit AS Initial_Debit,
        ac.ac_start_credit AS Initial_Credit,
        ac.open_amount AS Opening_Ledger_Balance,
        ac.sec_update_date AS Last_Account_Review
    FROM F_Account_Chart ac WITH (NOLOCK)
    LEFT JOIN F_Account_Chart parent WITH (NOLOCK) ON parent.ac_id = ac.ac_parent
    ORDER BY ac.ac_code ASC;
    """

    # 9. Operational & Indirect Expense Category Registry (Other_Expenses)
    q_operating_expenses = """
    SELECT 
        oe.oe_id AS Expense_Type_ID,
        oe.oe_code AS Expense_Code,
        oe.oe_name_ar AS Expense_Category_Name,
        ac.ac_code AS Linked_GL_Account_Code,
        ac.ac_name_ar AS Linked_GL_Account_Name,
        CASE oe.oe_active WHEN '1' THEN 'Active' ELSE 'Inactive' END AS Category_Status,
        oe.sec_insert_date AS Created_At
    FROM Other_Expenses oe WITH (NOLOCK)
    LEFT JOIN F_Account_Chart ac WITH (NOLOCK) ON ac.ac_id = oe.oe_ac_id
    ORDER BY oe.oe_id ASC;
    """

    # 10. GL Auto-Posting Routing Matrix & Parameters (F_Account_Settings)
    q_gl_rules = """
    SELECT 
        acs.acs_id AS Setting_ID,
        acs.acs_code AS Parameter_Code,
        acs.acs_name_ar AS Transaction_Type_Name,
        acs.acs_ac_id AS Target_Account_ID,
        ac.ac_code AS Target_GL_Account_Code,
        ac.ac_name_ar AS Target_GL_Account_Name
    FROM F_Account_Settings acs WITH (NOLOCK)
    LEFT JOIN F_Account_Chart ac WITH (NOLOCK) ON ac.ac_id = acs.acs_ac_id
    ORDER BY acs.acs_id;
    """

    # 11. Staff Monthly Payroll Disbursements & Deductions (mon_sal)
    q_staff_payroll = """
    SELECT 
        sal.sal_year AS Payout_Year,
        sal.sal_mon AS Payout_Month,
        emp.e_code AS Staff_Code,
        emp.e_Name AS Staff_Name,
        j.j_name_ar AS Job_Role,
        sal.salary AS Base_Salary,
        sal.ch_salary AS Variable_Salary,
        sal.sal_bouns AS Monthly_Bonus,
        sal.hr_late AS Late_Hours_Penalty,
        sal.days_off AS Days_Absent,
        sal.sal_pen AS Disciplinary_Deductions,
        sal.sal_fx_ins_emp + sal.sal_ch_ins_emp AS Employee_Social_Insurance,
        sal.sal_tx AS Income_Tax_Withheld,
        sal.sal_amount AS Net_Salary_Disbursed,
        sal.notes AS Payroll_Notes
    FROM mon_sal sal WITH (NOLOCK)
    JOIN Employee emp WITH (NOLOCK) ON emp.e_id = sal.sal_e_id
    LEFT JOIN Jobs j WITH (NOLOCK) ON j.j_id = emp.e_job_id
    ORDER BY sal.sal_year DESC, sal.sal_mon DESC, sal.sal_amount DESC;
    """

    # 12. ETA E-Invoicing & E-Receipt Compliance Audit Trail (sales_trans_h with ETA Flags)
    q_eta_compliance = """
    SELECT 
        h.sth_id AS Bill_ID,
        st.sto_name_ar AS Store_Name,
        h.sec_insert_date AS Invoice_Timestamp,
        ISNULL(c.cust_name_ar, 'Consumer Walk-in') AS Customer_Name,
        c.taxregno AS Customer_Tax_Registration_No,
        h.total_bill_after_disc AS Net_Invoice_Total,
        h.total_tax AS Sales_Tax_Levied,
        h.UUID AS ETA_Receipt_UUID,
        h.SubmissionId AS ETA_Submission_Ref,
        h.SendDate AS ETA_Transmission_Date,
        CASE 
            WHEN h.UUID IS NOT NULL AND h.SubmissionId IS NOT NULL THEN 'Submitted to Tax Authority (تم الإرسال للضرائب)'
            WHEN h.total_bill_after_disc >= 50000 AND c.taxregno IS NULL THEN 'B2B Compliance Alert: Missing Tax ID (>50k L.E)'
            ELSE 'Standard POS Retail Receipt'
        END AS Tax_Compliance_Status
    FROM sales_trans_h h WITH (NOLOCK)
    JOIN Store st WITH (NOLOCK) ON st.sto_id = h.sto_id
    LEFT JOIN Customer c WITH (NOLOCK) ON c.cust_id = h.cust_id
    WHERE h.sth_flag = 'C' 
      AND (h.UUID IS NOT NULL OR h.SubmissionId IS NOT NULL OR h.total_bill_after_disc >= 50000 OR c.taxregno IS NOT NULL)
    ORDER BY h.sth_id DESC;
    """

    # --- SINGLE-PASS IN-MEMORY EXCEL WRITING & FORMATTING ---
    table_names = set()
    fill_hex = "7D5A00"  # Bronze Gold
    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        write_and_style_sheet(
            writer,
            "Cash_Safes_Banks",
            read_sql_and_enrich(q_cash_stores, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Treasury_Ledger",
            read_sql_and_enrich(q_treasury_trans, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Cheques_Portfolio",
            read_sql_and_enrich(q_cheques, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Automated_GL_Journals",
            read_sql_and_enrich(q_gl_engine, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Debit_Credit_Notices",
            read_sql_and_enrich(q_notices, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Customer_Credit_Balances",
            read_sql_and_enrich(q_cust_debt, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Customer_Credit_Audit_Log",
            read_sql_and_enrich(q_customer_credit_audit, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Chart_Of_Accounts",
            read_sql_and_enrich(q_chart_of_accounts, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Operating_Expenses_Registry",
            read_sql_and_enrich(q_operating_expenses, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "GL_Posting_Rules_Settings",
            read_sql_and_enrich(q_gl_rules, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Staff_Payroll_Disbursements",
            read_sql_and_enrich(q_staff_payroll, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "ETA_EInvoice_Audit_Trail",
            read_sql_and_enrich(q_eta_compliance, ENGINE),
            fill_hex,
            table_names=table_names,
        )

    logger.info(
        "Module 4 (04_Genius_Financial_Ledger.xlsx) generated successfully with 12 styled sheets."
    )


# ==================================================================================================
# MODULE 5: POS CASHIER & SHIFT AUDIT TRAIL
# ==================================================================================================
def generate_shifts_workbook(output_dir):
    filename = os.path.join(output_dir, "05_Genius_POS_Shift_Operations.xlsx")
    logger.info(f"Generating Module 5: {filename}...")

    # 1. Master Shift Reconciliation & Blind Cash Count (Shift_Changes)
    q_shifts = """
    SELECT 
        sc.sc_id AS Shift_Log_ID,
        emp.e_Name AS Cashier_Name,
        fcs_vault.fcs_name_ar AS Transferred_To_Safe,
        sc.sc_shift_start_time AS Shift_Started_At,
        sc.sec_insert_date AS Shift_Closed_At,
        sc.sc_start_cash AS Starting_Float_Cash,
        sc.sc_curr_cash AS System_Calculated_Cash,
        sc.sc_act_cash AS Cashier_Actual_Counted_Cash,
        CAST(sc.sc_act_cash - sc.sc_curr_cash AS DECIMAL(18,2)) AS Cash_Variance_Over_Short,
        CASE 
            WHEN (sc.sc_act_cash - sc.sc_curr_cash) > 0 THEN 'Surplus (زيادة)' 
            WHEN (sc.sc_act_cash - sc.sc_curr_cash) < 0 THEN 'Deficit/Shortage (عجز)' 
            ELSE 'Balanced (مطابق)' 
        END AS Reconciliation_Status,
        sc.sc_trans_value AS Amount_Transferred_To_Vault,
        next_emp.e_Name AS Transferred_To_Employee,
        sc.sc_curr_credit AS Visa_System_Total,
        sc.sc_act_credit AS Visa_Actual_Slip_Total,
        sc.sc_notes AS Shift_Notes
    FROM Shift_Changes sc WITH (NOLOCK)
    JOIN Employee emp WITH (NOLOCK) ON emp.e_id = sc.sc_cashier_id
    LEFT JOIN F_Cash_Store fcs_vault WITH (NOLOCK) ON fcs_vault.fcs_id = sc.sc_fcs_id
    LEFT JOIN Employee next_emp WITH (NOLOCK) ON next_emp.e_id = TRY_CAST(sc.sc_to_employee AS DECIMAL(18,0))
    ORDER BY sc.sc_id DESC;
    """

    # 2. Detailed Shift Ledger Activity Breakdown (Shift_Changes_d)
    q_shifts_d = """
    SELECT 
        scd.scd_sc_id AS Shift_Log_ID,
        st.sto_name_ar AS Store_Name,
        fcs.fcs_name_ar AS Cash_Drawer,
        scd.scd_start AS Shift_Open,
        scd.scd_end AS Shift_Close,
        scd.scd_d_startBalance AS Opening_Balance,
        scd.scd_d_sales AS Cash_Sales_In,
        scd.scd_d_salesfromreturn AS Sales_From_Returns,
        scd.scd_d_visacollect AS Visa_Collections,
        scd.scd_d_cashIn AS Manual_Cash_In,
        scd.scd_c_purchases AS Outbound_Cash_Purchases,
        scd.scd_c_salesReturn AS Outbound_Sales_Refunds,
        scd.scd_c_cashOut AS Manual_Cash_Out_Expenses,
        scd.scd_o_salesVisa AS Visa_Card_Sales,
        scd.scd_o_salesCredit AS Credit_Account_Sales,
        scd.scd_o_salesPending AS Suspended_Bills_Value
    FROM Shift_Changes_d scd WITH (NOLOCK)
    JOIN Store st WITH (NOLOCK) ON st.sto_id = scd.scd_sto_id
    JOIN F_Cash_Store fcs WITH (NOLOCK) ON fcs.fcs_id = scd.scd_fcs_id
    ORDER BY scd.scd_sc_id DESC;
    """

    # 3. Client Workstation Login / Logout Audit Sessions (log_users)
    q_sessions = """
    SELECT 
        lu.lg_id AS Session_ID,
        usr.usr_name AS System_Username,
        emp.e_Name AS Employee_Full_Name,
        fcs.fcs_name_ar AS POS_Drawer_Assigned,
        lu.computer_nm AS Client_Workstation,
        lu.prog_ver AS ePlus_Software_Version,
        lu.lg_begin_tm AS Login_Timestamp,
        lu.lg_close_tm AS Logout_Timestamp,
        DATEDIFF(MINUTE, lu.lg_begin_tm, lu.lg_close_tm) AS Session_Duration_Minutes
    FROM log_users lu WITH (NOLOCK)
    LEFT JOIN Users usr WITH (NOLOCK) ON usr.usr_e_id = TRY_CAST(lu.lg_usr_id AS DECIMAL(18,0))
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = usr.usr_e_id
    LEFT JOIN F_Cash_Store fcs WITH (NOLOCK) ON fcs.fcs_id = lu.f_cash_store_id
    ORDER BY lu.lg_id DESC;
    """

    # 4. Master Cashier Staff Directory & POS Security Privileges (Employee)
    q_staff = """
    SELECT 
        e.e_id AS Employee_ID,
        e.e_code AS Employee_Code,
        e.e_Name AS Staff_Name,
        j.j_name_ar AS Job_Title,
        ews.ews_name AS Assigned_Shift_Scheme,
        ews.ews_start AS Shift_Default_Start,
        ews.ews_end AS Shift_Default_End,
        e.e_bill_disc_perc AS Max_Allowed_Bill_Discount,
        CASE e.e_select_cashier WHEN 1 THEN 'Yes' ELSE 'No' END AS Can_Select_Cashier,
        CASE e.e_change_sales_prices WHEN 1 THEN 'Yes' ELSE 'No' END AS Can_Override_Prices,
        CASE e.e_can_pend WHEN 1 THEN 'Yes' ELSE 'No' END AS Can_Suspend_Bills,
        e.sec_update_date AS Profile_Last_Updated
    FROM Employee e WITH (NOLOCK)
    LEFT JOIN Jobs j WITH (NOLOCK) ON j.j_id = e.e_job_id
    LEFT JOIN Emp_Work_Shifts ews WITH (NOLOCK) ON ews.ews_id = e.e_ews_id
    WHERE e.e_active = 1
    ORDER BY e.e_id ASC;
    """

    # 5. Outbound Drawer Cash Payouts & Expenses (Shift_Changes_d)
    q_drawer_payouts = """
    SELECT 
        scd.scd_sc_id AS Shift_Log_ID,
        st.sto_name_ar AS Store_Name,
        fcs.fcs_name_ar AS Cash_Drawer,
        scd.scd_start AS Shift_Open,
        scd.scd_end AS Shift_Close,
        scd.scd_c_purchases AS Outbound_Local_Purchases_From_Drawer,
        scd.scd_c_cashOut AS Petty_Cash_Operating_Expenses,
        scd.scd_c_salesReturn AS Cash_Refunds_Given_To_Customers,
        scd.scd_c_NoticeOut AS Credit_Debit_Notices_Paid,
        CAST(scd.scd_c_purchases + scd.scd_c_cashOut + scd.scd_c_salesReturn AS DECIMAL(18,2)) AS Total_Cash_Drained_From_Register,
        COALESCE(emp.e_Name, CAST(scd.scd_insert_uid AS VARCHAR(50))) AS Cashier_Name
    FROM Shift_Changes_d scd WITH (NOLOCK)
    JOIN Store st WITH (NOLOCK) ON st.sto_id = scd.scd_sto_id
    JOIN F_Cash_Store fcs WITH (NOLOCK) ON fcs.fcs_id = scd.scd_fcs_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(scd.scd_insert_uid AS DECIMAL(18,0))
    WHERE (scd.scd_c_purchases + scd.scd_c_cashOut + scd.scd_c_salesReturn) > 0
    ORDER BY scd.scd_sc_id DESC;
    """

    # 6. Cashier Workstation & Register Drawer Authorizations (Emp_cash_store)
    q_drawer_authorizations = """
    SELECT 
        ecs.ecs_e_id AS Employee_ID,
        emp.e_Name AS Staff_Name,
        st.sto_name_ar AS Branch_Assigned,
        fcs.fcs_name_ar AS Authorized_POS_Drawer,
        fcs.fcs_computer_name AS Dedicated_Workstation,
        CASE emp.e_select_cashier WHEN 1 THEN 'Yes' ELSE 'No' END AS Can_Switch_Drawers,
        ecs.ecs_insert_date AS Assignment_Date
    FROM Emp_cash_store ecs WITH (NOLOCK)
    JOIN Employee emp WITH (NOLOCK) ON emp.e_id = ecs.ecs_e_id
    JOIN Store st WITH (NOLOCK) ON st.sto_id = ecs.ecs_sto_id
    JOIN F_Cash_Store fcs WITH (NOLOCK) ON fcs.fcs_id = ecs.ecs_fcs_id
    ORDER BY st.sto_id, emp.e_id;
    """

    # 7. Employee Store & Inventory Access Permissions (Emp_Stores)
    q_emp_stores = """
    SELECT 
        es.es_e_id AS Employee_ID,
        emp.e_Name AS Staff_Name,
        st.sto_name_ar AS Authorized_Store,
        CASE es.es_readOnly WHEN 1 THEN 'Read-Only (عرض فقط)' ELSE 'Full Read/Write (كامل)' END AS Permission_Level,
        es.sec_update_date AS Authorization_Timestamp
    FROM Emp_Stores es WITH (NOLOCK)
    JOIN Employee emp WITH (NOLOCK) ON emp.e_id = es.es_e_id
    JOIN Store st WITH (NOLOCK) ON st.sto_id = es.es_sto_id
    ORDER BY emp.e_id, st.sto_id;
    """

    # 8. Employee Contract Compensation Terms (Emp_Salary)
    q_emp_salary_terms = """
    SELECT 
        es.es_id AS Contract_ID,
        emp.e_Name AS Staff_Name,
        j.j_name_ar AS Role_Title,
        es.es_Type AS Salary_Periodicity,
        es.es_value AS Base_Contracted_Salary,
        es.es_over_h AS Overtime_Hourly_Rate,
        es.es_insurance_disc AS Contracted_Insurance_Deduction,
        es.es_tax_disc AS Contracted_Tax_Deduction,
        es.es_date AS Contract_Effective_Date
    FROM Emp_Salary es WITH (NOLOCK)
    JOIN Employee emp WITH (NOLOCK) ON emp.e_id = es.es_e_id
    LEFT JOIN Jobs j WITH (NOLOCK) ON j.j_id = emp.e_job_id
    ORDER BY emp.e_id;
    """

    # 9. Intercepted Duplicate Sales Vouchers (F_Transaction_Header_SaveDeleteRecords)
    q_duplicate_vouchers = """
    SELECT 
        fh_id AS Intercepted_Record_ID,
        fh_code AS Voucher_Code,
        fh_value AS Voucher_Value,
        fh_notes AS Operational_Narrative,
        fh_actual_date AS Voucher_Date,
        fh_computer AS Workstation_Hostname,
        sec_insert_date AS Purged_Timestamp
    FROM F_Transaction_Header_SaveDeleteRecords WITH (NOLOCK)
    ORDER BY fh_actual_date DESC;
    """

    # 10. Three-Way Shift Cash Flow Statement (BUG 1 FIXED: Decoupled Drawer & Vault Joins)
    q_shift_cash_flow_matrix = """
    SELECT 
        sc.sc_id AS Shift_Log_ID,
        emp.e_Name AS Cashier_Name,
        fcs_drawer.fcs_name_ar AS Cash_Drawer,
        sc.sc_shift_start_time AS Shift_Started,
        sc.sec_insert_date AS Shift_Closed,
        sc.sc_start_cash AS Starting_Float,
        ISNULL(scd.Cash_Sales, 0.00) AS Inflow_Cash_Sales,
        ISNULL(scd.Visa_Collect, 0.00) AS Inflow_Card_Collections,
        ISNULL(scd.Debt_Paid, 0.00) AS Inflow_Customer_Debt_Paid,
        ISNULL(scd.Pur_Refunds, 0.00) AS Inflow_Supplier_Refunds,
        CAST(ISNULL(scd.Cash_Sales, 0) + ISNULL(scd.Debt_Paid, 0) + ISNULL(scd.Pur_Refunds, 0) AS DECIMAL(18,2)) AS Total_Gross_Inflows,
        ISNULL(scd.Local_Purchases, 0.00) AS Outflow_Local_Purchases,
        ISNULL(scd.Sales_Refunds, 0.00) AS Outflow_Sales_Refunds,
        ISNULL(scd.Petty_Expenses, 0.00) AS Outflow_Operating_Expenses,
        CAST(ISNULL(scd.Local_Purchases, 0) + ISNULL(scd.Sales_Refunds, 0) + ISNULL(scd.Petty_Expenses, 0) AS DECIMAL(18,2)) AS Total_Gross_Outflows,
        sc.sc_curr_cash AS System_Calculated_Closing_Cash,
        sc.sc_act_cash AS Cashier_Counted_Closing_Cash,
        CAST(sc.sc_act_cash - sc.sc_curr_cash AS DECIMAL(18,2)) AS Register_Discrepancy_Over_Short,
        sc.sc_trans_value AS Vault_Drop_Transferred,
        fcs_vault.fcs_name_ar AS Transferred_To_Safe,
        next_emp.e_Name AS Vault_Receiver
    FROM Shift_Changes sc WITH (NOLOCK)
    JOIN Employee emp WITH (NOLOCK) ON emp.e_id = sc.sc_cashier_id
    LEFT JOIN (
        SELECT 
            scd_sc_id,
            MIN(scd_fcs_id) AS primary_drawer_id,
            SUM(scd_d_sales) AS Cash_Sales,
            SUM(scd_d_visacollect) AS Visa_Collect,
            SUM(scd_d_cashIn) AS Debt_Paid,
            SUM(scd_d_purReturn) AS Pur_Refunds,
            SUM(scd_c_purchases) AS Local_Purchases,
            SUM(scd_c_salesReturn) AS Sales_Refunds,
            SUM(scd_c_cashOut) AS Petty_Expenses
        FROM Shift_Changes_d WITH (NOLOCK)
        GROUP BY scd_sc_id
    ) scd ON scd.scd_sc_id = sc.sc_id
    LEFT JOIN F_Cash_Store fcs_drawer WITH (NOLOCK) ON fcs_drawer.fcs_id = scd.primary_drawer_id
    LEFT JOIN F_Cash_Store fcs_vault WITH (NOLOCK) ON fcs_vault.fcs_id = sc.sc_fcs_id
    LEFT JOIN Employee next_emp WITH (NOLOCK) ON next_emp.e_id = TRY_CAST(sc.sc_to_employee AS DECIMAL(18,0))
    ORDER BY sc.sc_id DESC;
    """

    # --- SINGLE-PASS IN-MEMORY EXCEL WRITING & FORMATTING ---
    table_names = set()
    fill_hex = "1A5276"  # Slate Teal
    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        write_and_style_sheet(
            writer,
            "Shift_Reconciliation",
            read_sql_and_enrich(q_shifts, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Shift_Operations_Breakdown",
            read_sql_and_enrich(q_shifts_d, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Station_Login_Sessions",
            read_sql_and_enrich(q_sessions, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Cashier_Staff_Roster",
            read_sql_and_enrich(q_staff, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Drawer_Expense_Payouts",
            read_sql_and_enrich(q_drawer_payouts, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Cashier_Drawer_Authorizations",
            read_sql_and_enrich(q_drawer_authorizations, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Staff_Store_Permissions",
            read_sql_and_enrich(q_emp_stores, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Employee_Contract_Terms",
            read_sql_and_enrich(q_emp_salary_terms, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Duplicate_Voucher_Interceptions",
            read_sql_and_enrich(q_duplicate_vouchers, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Shift_Cash_Flow_Statement",
            read_sql_and_enrich(q_shift_cash_flow_matrix, ENGINE),
            fill_hex,
            table_names=table_names,
        )

    logger.info(
        "Module 5 (05_Genius_POS_Shift_Operations.xlsx) generated successfully with 10 styled sheets."
    )


# ==================================================================================================
# MODULE 6: EXECUTIVE BUSINESS INTELLIGENCE & ANALYTICAL COCKPIT
# ==================================================================================================
def generate_executive_bi_workbook(output_dir):
    filename = os.path.join(output_dir, "06_Genius_Executive_Intelligence_Cockpit.xlsx")
    logger.info(f"Generating Module 6: {filename}...")

    # 1. Executive C-Suite KPI Balance Sheet & Scorecard (BUG 4 & Safe 8 Trap Resolved)
    q_kpi_scorecard = """
    SELECT '1. Total Gross Sales Revenue' AS Metric, 
           CAST(ISNULL(SUM(total_bill), 0) AS DECIMAL(18,2)) AS Value_L_E, 
           'Total customer sales billed before discounts' AS Notes 
    FROM sales_trans_h WITH (NOLOCK) WHERE sth_flag='C'

    UNION ALL
    SELECT '2. Total Customer Discounts Conceded', 
           CAST(ISNULL(SUM(total_bill - total_bill_after_disc), 0) AS DECIMAL(18,2)), 
           'Discounts granted to retail customers at invoice level' 
    FROM sales_trans_h WITH (NOLOCK) WHERE sth_flag='C'

    UNION ALL
    SELECT '3. Total Customer Returns Refunded', 
           CAST(ISNULL(SUM(returned_items_value), 0) AS DECIMAL(18,2)), 
           'Total customer merchandise returned and refunded' 
    FROM sales_return WITH (NOLOCK)

    UNION ALL
    SELECT '4. Total Net Sales Realized', 
           CAST(ISNULL(
               (SELECT SUM(total_bill_after_disc) FROM sales_trans_h WITH (NOLOCK) WHERE sth_flag='C')
               - (SELECT SUM(returned_items_value) FROM sales_return WITH (NOLOCK))
           , 0) AS DECIMAL(18,2)), 
           'Actual realized top-line revenue (Netted of customer returns)' 

    UNION ALL
    SELECT '5. Total Cost of Goods Sold (COGS)', 
           CAST(ISNULL(SUM((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_cost), 0) AS DECIMAL(18,2)), 
           'Direct acquisition cost of net dispensed merchandise' 
    FROM sales_trans_d d WITH (NOLOCK) 
    JOIN sales_trans_h h WITH (NOLOCK) ON h.sth_id=d.sth_id 
    WHERE h.sth_flag='C' AND ISNULL(d.itm_nexist, 0) <> 1

    UNION ALL
    SELECT '6. Cumulative Realized Gross Profit', 
           CAST(ISNULL(
               (SELECT SUM(total_bill_after_disc) FROM sales_trans_h WITH (NOLOCK) WHERE sth_flag='C')
               - (SELECT SUM(returned_items_value) FROM sales_return WITH (NOLOCK))
               - (SELECT SUM((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_cost) 
                  FROM sales_trans_d d WITH (NOLOCK) 
                  JOIN sales_trans_h h WITH (NOLOCK) ON h.sth_id=d.sth_id 
                  WHERE h.sth_flag='C' AND ISNULL(d.itm_nexist, 0) <> 1)
           , 0) AS DECIMAL(18,2)), 
           'True Net Revenue minus Total Product Cost (Exact Reconciled Gross Profit)'

    UNION ALL
    SELECT '7. Total Retail Dispensary Stock (At Cost)', 
           CAST(ISNULL(SUM(itm_qty * pharm_price), 0) AS DECIMAL(18,2)), 
           'Current capital tied up in Main Pharmacy Dispensary (Store 47)' 
    FROM Item_Class_Store WITH (NOLOCK) WHERE itm_qty > 0 AND sto_id = 47

    UNION ALL
    SELECT '8. Total Consignment / Quarantine Stock (At Cost)', 
           CAST(ISNULL(SUM(itm_qty * pharm_price), 0) AS DECIMAL(18,2)), 
           'Current stock under quarantine or consignment (Store 48: تحت التصريف)' 
    FROM Item_Class_Store WITH (NOLOCK) WHERE itm_qty > 0 AND sto_id = 48

    UNION ALL
    SELECT '9. Total Combined Stock (At Retail)', 
           CAST(ISNULL(SUM(itm_qty * sell_price), 0) AS DECIMAL(18,2)), 
           'Potential retail revenue of on-shelf inventory across all branches' 
    FROM Item_Class_Store WITH (NOLOCK) WHERE itm_qty > 0

    UNION ALL
    SELECT '10. Total Customer Outstanding Debt (Accounts Receivable)', 
           CAST(ISNULL(SUM(cust_current_credit), 0) AS DECIMAL(18,2)), 
           'Total money owed to pharmacy by credit customers (Receivables)' 
    FROM Customer WITH (NOLOCK) WHERE cust_current_credit > 0 AND cust_active = '1'

    UNION ALL
    SELECT '11. Total Vendor Liabilities (Accounts Payable)', 
           CAST(ISNULL(SUM(ven_current_credit), 0) AS DECIMAL(18,2)), 
           'Total unpaid invoices owed to pharmaceutical distributors (We owe them)' 
    FROM Vendor WITH (NOLOCK) WHERE ven_current_credit > 0 AND ven_active = '1'

    UNION ALL
    SELECT '12. Total Vendor Advances & Return Balances', 
           CAST(ISNULL(ABS(SUM(ven_current_credit)), 0) AS DECIMAL(18,2)), 
           'Distributor credit balances due back to pharmacy (They owe us)' 
    FROM Vendor WITH (NOLOCK) WHERE ven_current_credit < 0 AND ven_active = '1'

    UNION ALL
    SELECT '13. Live Physical POS Register Cash (Drawers)', 
           CAST(ISNULL(SUM(fcs_current_balance), 0) AS DECIMAL(18,2)), 
           'Verified live cash currently resting in active cashier drawers' 
    FROM F_Cash_Store WITH (NOLOCK) WHERE fcs_active = 1 AND fcs_type = '1'

    UNION ALL
    SELECT '14. Total Verified Bank Balances', 
           CAST(ISNULL(SUM(fcs_current_balance), 0) AS DECIMAL(18,2)), 
           'Commercial bank accounts balance' 
    FROM F_Cash_Store WITH (NOLOCK) WHERE fcs_active = 1 AND fcs_type = '3'

    UNION ALL
    SELECT '15. Main Safe Book Ledger (Subject to Owner Drawings Audit)', 
           CAST(ISNULL(SUM(fcs_current_balance), 0) AS DECIMAL(18,2)), 
           'Cumulative book balance of central vault (fcs_id=8: requires reconciliation)' 
    FROM F_Cash_Store WITH (NOLOCK) WHERE fcs_active = 1 AND fcs_type = '2'

    UNION ALL
    SELECT '16. Net POS Cash Drawer Discrepancies (Shortage/Surplus)', 
           CAST(ISNULL(SUM(sc_act_cash - sc_curr_cash), 0) AS DECIMAL(18,2)), 
           'Cumulative historical cashier cash drawer variance' 
    FROM Shift_Changes WITH (NOLOCK);
    """

    # 2. Reconciled Monthly P&L Statement (Sales, Returns, Net Revenue, COGS, Expenses, Payroll)
    q_monthly_pnl = """
    WITH MonthlySales AS (
        SELECT 
            YEAR(sec_insert_date) AS Fin_Year,
            MONTH(sec_insert_date) AS Fin_Month,
            CAST(SUM(total_bill) AS DECIMAL(18,2)) AS Gross_Sales,
            CAST(SUM(total_bill - total_bill_after_disc) AS DECIMAL(18,2)) AS Discounts_Granted,
            CAST(SUM(total_bill_after_disc) AS DECIMAL(18,2)) AS Billed_Revenue
        FROM sales_trans_h WITH (NOLOCK)
        WHERE sth_flag = 'C'
        GROUP BY YEAR(sec_insert_date), MONTH(sec_insert_date)
    ),
    MonthlyReturns AS (
        SELECT 
            YEAR(sec_insert_date) AS Fin_Year,
            MONTH(sec_insert_date) AS Fin_Month,
            CAST(SUM(returned_items_value) AS DECIMAL(18,2)) AS Total_Refunds
        FROM sales_return WITH (NOLOCK)
        GROUP BY YEAR(sec_insert_date), MONTH(sec_insert_date)
    ),
    MonthlyCOGS AS (
        SELECT 
            YEAR(h.sec_insert_date) AS Fin_Year,
            MONTH(h.sec_insert_date) AS Fin_Month,
            CAST(SUM((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_cost) AS DECIMAL(18,2)) AS COGS
        FROM sales_trans_d d WITH (NOLOCK)
        JOIN sales_trans_h h WITH (NOLOCK) ON h.sth_id = d.sth_id
        WHERE h.sth_flag = 'C' AND ISNULL(d.itm_nexist, 0) <> 1
        GROUP BY YEAR(h.sec_insert_date), MONTH(h.sec_insert_date)
    ),
    MonthlyDrawerExpenses AS (
        SELECT 
            YEAR(scd_start) AS Fin_Year,
            MONTH(scd_start) AS Fin_Month,
            CAST(SUM(scd_c_cashOut) AS DECIMAL(18,2)) AS Drawer_Petty_Expenses
        FROM Shift_Changes_d WITH (NOLOCK)
        WHERE scd_start IS NOT NULL
        GROUP BY YEAR(scd_start), MONTH(scd_start)
    ),
    MonthlyPayroll AS (
        SELECT 
            sal_year AS Fin_Year,
            sal_mon AS Fin_Month,
            CAST(SUM(sal_amount) AS DECIMAL(18,2)) AS Total_Payroll_Disbursed
        FROM mon_sal WITH (NOLOCK)
        GROUP BY sal_year, sal_mon
    )
    SELECT 
        s.Fin_Year,
        s.Fin_Month,
        s.Gross_Sales,
        s.Discounts_Granted,
        ISNULL(r.Total_Refunds, 0.00) AS Customer_Returns_Refunded,
        CAST(s.Billed_Revenue - ISNULL(r.Total_Refunds, 0.00) AS DECIMAL(18,2)) AS Net_Realized_Revenue,
        ISNULL(c.COGS, 0.00) AS COGS,
        CAST((s.Billed_Revenue - ISNULL(r.Total_Refunds, 0.00)) - ISNULL(c.COGS, 0.00) AS DECIMAL(18,2)) AS Gross_Profit,
        CAST(((s.Billed_Revenue - ISNULL(r.Total_Refunds, 0.00)) - ISNULL(c.COGS, 0.00)) / NULLIF(s.Billed_Revenue - ISNULL(r.Total_Refunds, 0.00), 0) AS DECIMAL(18,4)) AS Gross_Margin_Perc,
        ISNULL(e.Drawer_Petty_Expenses, 0.00) AS POS_Operating_Expenses,
        ISNULL(p.Total_Payroll_Disbursed, 0.00) AS Payroll_Personnel_Expenses,
        CAST(((s.Billed_Revenue - ISNULL(r.Total_Refunds, 0.00)) - ISNULL(c.COGS, 0.00)) - ISNULL(e.Drawer_Petty_Expenses, 0.00) - ISNULL(p.Total_Payroll_Disbursed, 0.00) AS DECIMAL(18,2)) AS Estimated_Net_Operating_Profit
    FROM MonthlySales s
    LEFT JOIN MonthlyReturns r ON r.Fin_Year = s.Fin_Year AND r.Fin_Month = s.Fin_Month
    LEFT JOIN MonthlyCOGS c ON c.Fin_Year = s.Fin_Year AND c.Fin_Month = s.Fin_Month
    LEFT JOIN MonthlyDrawerExpenses e ON e.Fin_Year = s.Fin_Year AND e.Fin_Month = s.Fin_Month
    LEFT JOIN MonthlyPayroll p ON p.Fin_Year = s.Fin_Year AND p.Fin_Month = s.Fin_Month
    ORDER BY s.Fin_Year DESC, s.Fin_Month DESC;
    """

    # 3. Item-Level ABC Revenue Pareto Analysis
    q_abc_analysis = """
    WITH ItemSales AS (
        SELECT 
            ic.itm_id,
            ic.itm_code AS Barcode,
            ic.itm_name_ar AS Item_Name,
            c.com_name_ar AS Manufacturer,
            CAST(SUM((d.qnty - ISNULL(d.itm_back, 0)) * (
                CASE d.itm_unit 
                    WHEN 3 THEN (1.0 / ISNULL(NULLIF(ic.itm_unit1_unit3, 0), 1.0)) 
                    WHEN 2 THEN (1.0 / ISNULL(NULLIF(ic.itm_unit1_unit2, 0), 1.0)) 
                    ELSE 1.0 
                END
            )) AS DECIMAL(18,2)) AS Standard_Packs_Sold,
            CAST(SUM(
                (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell * (1.0 - (ISNULL(d.itm_dis_per, 0) / 100.0))) - ISNULL(d.itm_dis_mon, 0))
                * (1.0 - (ISNULL(h.total_dis_per, 0) / 100.0))
                - (ISNULL(h.total_des_mon, 0) * (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell) / NULLIF(h.total_bill, 0)))
            ) AS DECIMAL(18,2)) AS Net_Revenue,
            CAST(SUM((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_cost) AS DECIMAL(18,2)) AS Total_COGS,
            CAST(SUM(
                (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell * (1.0 - (ISNULL(d.itm_dis_per, 0) / 100.0))) - ISNULL(d.itm_dis_mon, 0))
                * (1.0 - (ISNULL(h.total_dis_per, 0) / 100.0))
                - (ISNULL(h.total_des_mon, 0) * (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell) / NULLIF(h.total_bill, 0)))
            ) - SUM((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_cost) AS DECIMAL(18,2)) AS Gross_Profit
        FROM sales_trans_d d WITH (NOLOCK)
        JOIN sales_trans_h h WITH (NOLOCK) ON h.sth_id = d.sth_id
        JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = d.itm_id
        LEFT JOIN Company c WITH (NOLOCK) ON c.com_id = ic.com_id
        WHERE h.sth_flag = 'C'
          AND ISNULL(ic.itm_srvc, 0) = 0
          AND ISNULL(d.itm_nexist, 0) <> 1
        GROUP BY ic.itm_id, ic.itm_code, ic.itm_name_ar, c.com_name_ar
        HAVING SUM(d.qnty - ISNULL(d.itm_back, 0)) > 0
    ),
    RankedSales AS (
        SELECT *,
            SUM(Net_Revenue) OVER () AS Total_Market_Revenue,
            SUM(Net_Revenue) OVER (ORDER BY Net_Revenue DESC ROWS BETWEEN UNBOUNDED PRECEDING AND CURRENT ROW) AS Running_Revenue
        FROM ItemSales
    )
    SELECT 
        itm_id AS Item_ID,
        Barcode,
        Item_Name,
        Manufacturer,
        Standard_Packs_Sold,
        Net_Revenue,
        Total_COGS,
        Gross_Profit,
        CAST((Gross_Profit / NULLIF(Net_Revenue, 0)) AS DECIMAL(18,4)) AS Gross_Margin_Perc,
        CAST((Running_Revenue / NULLIF(Total_Market_Revenue, 0)) AS DECIMAL(18,4)) AS Cumulative_Revenue_Share,
        CASE 
            WHEN (Running_Revenue / NULLIF(Total_Market_Revenue, 0)) <= 0.80 THEN 'Class A (Top 80% Rev)'
            WHEN (Running_Revenue / NULLIF(Total_Market_Revenue, 0)) <= 0.95 THEN 'Class B (Next 15% Rev)'
            ELSE 'Class C (Bottom 5% Rev)'
        END AS ABC_Classification
    FROM RankedSales
    ORDER BY Net_Revenue DESC;
    """

    # 4. Expiry & Stagnant Dead Stock Early-Warning Radar
    q_dead_stock_radar = """
    WITH LastItemSale AS (
        SELECT d.itm_id, MAX(h.sec_insert_date) AS Last_Sold_Date
        FROM sales_trans_d d WITH (NOLOCK)
        JOIN sales_trans_h h WITH (NOLOCK) ON h.sth_id = d.sth_id
        WHERE h.sth_flag = 'C'
        GROUP BY d.itm_id
    )
    SELECT 
        ics.sto_id AS Store_ID,
        st.sto_name_ar AS Store_Name,
        ic.itm_id AS Item_ID,
        ic.itm_code AS Barcode,
        ic.itm_name_ar AS Item_Name,
        ics.c_id AS Batch_ID,
        ics.itm_qty AS Stock_On_Hand,
        ics.pharm_price AS Unit_Cost,
        CAST(ics.itm_qty * ics.pharm_price AS DECIMAL(18,2)) AS Capital_At_Risk,
        CONVERT(VARCHAR(10), ics.itm_expiry_date, 120) AS Expiry_Date,
        DATEDIFF(DAY, GETDATE(), ics.itm_expiry_date) AS Days_Until_Expiration,
        CASE 
            WHEN ics.itm_expiry_date < GETDATE() THEN 'EXPIRED (منتهي الصلاحية)'
            WHEN DATEDIFF(DAY, GETDATE(), ics.itm_expiry_date) <= 90 THEN 'CRITICAL EXPIRY (< 90 Days)'
            WHEN DATEDIFF(DAY, GETDATE(), ics.itm_expiry_date) <= 180 THEN 'WARNING EXPIRY (< 180 Days)'
            WHEN lis.Last_Sold_Date < DATEADD(DAY, -180, GETDATE()) OR lis.Last_Sold_Date IS NULL THEN 'STAGNANT DEAD STOCK (>180 Days No Sale)'
            ELSE 'STABLE'
        END AS Stock_Risk_Category,
        ISNULL(lis.Last_Sold_Date, '2000-01-01') AS Last_Sold_Date
    FROM Item_Class_Store ics WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = ics.itm_id
    JOIN Store st WITH (NOLOCK) ON st.sto_id = ics.sto_id
    LEFT JOIN LastItemSale lis ON lis.itm_id = ics.itm_id
    WHERE ics.itm_qty > 0 
      AND (ics.itm_expiry_date <= DATEADD(MONTH, 6, GETDATE()) 
           OR lis.Last_Sold_Date < DATEADD(DAY, -180, GETDATE()) 
           OR lis.Last_Sold_Date IS NULL)
    ORDER BY Days_Until_Expiration ASC;
    """

    # 5. Hourly Footfall, Sales Traffic & Basket Density
    q_hourly_traffic = """
    SELECT 
        DATEPART(HOUR, sec_insert_date) AS Hour_Of_Day_24H,
        COUNT(sth_id) AS Total_Bills_Count,
        CAST(SUM(total_bill_after_disc) AS DECIMAL(18,2)) AS Total_Revenue_Generated,
        CAST(AVG(total_bill_after_disc) AS DECIMAL(18,2)) AS Average_Basket_Size,
        CAST(SUM(total_bill_after_disc) / NULLIF((SELECT SUM(total_bill_after_disc) FROM sales_trans_h WHERE sth_flag='C'), 0) AS DECIMAL(18,4)) AS Share_Of_Total_Sales
    FROM sales_trans_h WITH (NOLOCK)
    WHERE sth_flag = 'C'
    GROUP BY DATEPART(HOUR, sec_insert_date)
    ORDER BY Hour_Of_Day_24H ASC;
    """

    # 6. Strategic Vendor Performance & Supply Concentration Matrix
    q_vendor_performance = """
    SELECT 
        v.ven_id AS Vendor_ID,
        v.ven_name_ar AS Vendor_Name,
        COUNT(DISTINCT pth.pth_id) AS Total_Purchase_Bills,
        CAST(SUM(pth.total_bill) AS DECIMAL(18,2)) AS Total_Purchased_Gross,
        CAST(SUM(pth.total_after_back) AS DECIMAL(18,2)) AS Total_Purchased_Net,
        CAST(AVG(pth.total_dis_per) AS DECIMAL(18,2)) AS Avg_Negotiated_Discount_Perc,
        ISNULL((
            SELECT CAST(SUM(ibh_total_value) AS DECIMAL(18,2)) 
            FROM Item_back_h WITH (NOLOCK) 
            WHERE ibh_ven_id = v.ven_id
        ), 0.00) AS Total_Returned_Value,
        v.ven_current_credit AS Current_Outstanding_Debt
    FROM Vendor v WITH (NOLOCK)
    JOIN pur_trans_h pth WITH (NOLOCK) ON pth.ven_id = v.ven_id
    WHERE pth.bill_status <> 1
    GROUP BY v.ven_id, v.ven_name_ar, v.ven_current_credit
    ORDER BY Total_Purchased_Net DESC;
    """

    # 7. Multi-Branch Operational Benchmark (Dispensary vs Quarantine)
    q_store_matrix = """
    SELECT 
        st.sto_id AS Store_ID,
        st.sto_name_ar AS Store_Name,
        CASE WHEN st.sto_id = 47 THEN 'Dispensary Retail' ELSE 'Consignment / Quarantine' END AS Store_Role,
        ISNULL((SELECT COUNT(*) FROM Item_Class_Store WHERE sto_id=st.sto_id AND itm_qty > 0), 0) AS Active_SKUs_In_Stock,
        ISNULL((SELECT CAST(SUM(itm_qty * pharm_price) AS DECIMAL(18,2)) FROM Item_Class_Store WHERE sto_id=st.sto_id AND itm_qty > 0), 0) AS Total_Stock_Cost_Valuation,
        ISNULL((SELECT CAST(SUM(itm_qty * sell_price) AS DECIMAL(18,2)) FROM Item_Class_Store WHERE sto_id=st.sto_id AND itm_qty > 0), 0) AS Total_Stock_Retail_Valuation,
        ISNULL((SELECT COUNT(*) FROM sales_trans_h WHERE sto_id=st.sto_id AND sth_flag='C'), 0) AS Total_Sales_Transactions,
        ISNULL((SELECT CAST(SUM(total_bill_after_disc) AS DECIMAL(18,2)) FROM sales_trans_h WHERE sto_id=st.sto_id AND sth_flag='C'), 0) AS Total_Sales_Revenue,
        ISNULL((
            SELECT CAST(SUM(sc.sc_act_cash - sc.sc_curr_cash) AS DECIMAL(18,2)) 
            FROM Shift_Changes sc WITH (NOLOCK) 
            WHERE sc.sc_id IN (SELECT scd_sc_id FROM Shift_Changes_d WHERE scd_sto_id=st.sto_id)
        ), 0.00) AS Cumulative_Shift_Variance
    FROM Store st WITH (NOLOCK)
    WHERE st.activated = 1;
    """

    # 8. Monthly Inventory Close Valuation Trend (ICS_Month_Close)
    q_monthly_stock_trend = """
    SELECT 
        YEAR(imc.sec_insert_date) AS Close_Year,
        MONTH(imc.sec_insert_date) AS Close_Month,
        st.sto_name_ar AS Store_Name,
        COUNT(DISTINCT imc.itm_id) AS Distinct_SKUs_Stored,
        CAST(SUM(imc.itm_qty) AS DECIMAL(18,2)) AS Total_Stock_Units,
        CAST(SUM(imc.itm_qty * imc.pharm_price) AS DECIMAL(18,2)) AS Total_Inventory_Cost_Valuation,
        CAST(SUM(imc.itm_qty * imc.sell_price) AS DECIMAL(18,2)) AS Total_Inventory_Retail_Valuation,
        CAST(SUM(imc.itm_qty * (imc.sell_price - imc.pharm_price)) AS DECIMAL(18,2)) AS Unrealized_Gross_Margin,
        CAST(SUM(imc.itm_qty * (imc.sell_price - imc.pharm_price)) / NULLIF(SUM(imc.itm_qty * imc.sell_price), 0) AS DECIMAL(18,4)) AS Potential_Profit_Margin_Perc
    FROM ICS_Month_Close imc WITH (NOLOCK)
    JOIN Store st WITH (NOLOCK) ON st.sto_id = imc.sto_id
    GROUP BY YEAR(imc.sec_insert_date), MONTH(imc.sec_insert_date), st.sto_name_ar
    ORDER BY Close_Year DESC, Close_Month DESC;
    """

    # 9. Therapeutic Class Profitability & Margin Ranking (Scientific_Groups)
    q_therapeutic_profitability = """
    SELECT 
        ISNULL(sg.sg_name_ar, 'General Pharmaceuticals (غير مصنف)') AS Therapeutic_Class,
        COUNT(DISTINCT d.itm_id) AS Active_Medicines_Count,
        CAST(SUM((d.qnty - ISNULL(d.itm_back, 0)) * (
            CASE d.itm_unit 
                WHEN 3 THEN (1.0 / ISNULL(NULLIF(ic.itm_unit1_unit3, 0), 1.0)) 
                WHEN 2 THEN (1.0 / ISNULL(NULLIF(ic.itm_unit1_unit2, 0), 1.0)) 
                ELSE 1.0 
            END
        )) AS DECIMAL(18,2)) AS Standard_Packs_Sold,
        CAST(SUM(
            (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell * (1.0 - (ISNULL(d.itm_dis_per, 0)/100.0))) - ISNULL(d.itm_dis_mon, 0))
            * (1.0 - (ISNULL(h.total_dis_per, 0)/100.0))
            - (ISNULL(h.total_des_mon, 0) * (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell) / NULLIF(h.total_bill, 0)))
        ) AS DECIMAL(18,2)) AS Net_Revenue,
        CAST(SUM((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_cost) AS DECIMAL(18,2)) AS Total_COGS,
        CAST(SUM(
            (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell * (1.0 - (ISNULL(d.itm_dis_per, 0)/100.0))) - ISNULL(d.itm_dis_mon, 0))
            * (1.0 - (ISNULL(h.total_dis_per, 0)/100.0))
            - (ISNULL(h.total_des_mon, 0) * (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell) / NULLIF(h.total_bill, 0)))
        ) - SUM((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_cost) AS DECIMAL(18,2)) AS Gross_Profit,
        CAST((SUM(
            (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell * (1.0 - (ISNULL(d.itm_dis_per, 0)/100.0))) - ISNULL(d.itm_dis_mon, 0))
            * (1.0 - (ISNULL(h.total_dis_per, 0)/100.0))
            - (ISNULL(h.total_des_mon, 0) * (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell) / NULLIF(h.total_bill, 0)))
        ) - SUM((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_cost)) 
             / NULLIF(SUM(
                 (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell * (1.0 - (ISNULL(d.itm_dis_per, 0)/100.0))) - ISNULL(d.itm_dis_mon, 0))
                 * (1.0 - (ISNULL(h.total_dis_per, 0)/100.0))
                 - (ISNULL(h.total_des_mon, 0) * (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell) / NULLIF(h.total_bill, 0)))
             ), 0) AS DECIMAL(18,4)) AS Gross_Margin_Rate
    FROM sales_trans_d d WITH (NOLOCK)
    JOIN sales_trans_h h WITH (NOLOCK) ON h.sth_id = d.sth_id
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = d.itm_id
    LEFT JOIN Scientific_Groups sg WITH (NOLOCK) ON sg.sg_id = ic.itm_scientific_group_id
    WHERE h.sth_flag = 'C'
      AND ISNULL(ic.itm_srvc, 0) = 0
      AND ISNULL(d.itm_nexist, 0) <> 1
    GROUP BY sg.sg_name_ar
    HAVING SUM(d.qnty - ISNULL(d.itm_back, 0)) > 0
    ORDER BY Net_Revenue DESC;
    """

    # 10. Customer Receivables Exposure & Aging Risk Buckets
    q_receivables_aging = """
    SELECT 
        c.cust_id AS Customer_ID,
        c.cust_code AS Customer_Code,
        c.cust_name_ar AS Customer_Name,
        c.cust_mobile AS Phone_Number,
        c.cust_current_credit AS Total_Outstanding_Balance,
        DATEDIFF(DAY, c.sec_update_date, GETDATE()) AS Days_Since_Last_Activity,
        CASE 
            WHEN DATEDIFF(DAY, c.sec_update_date, GETDATE()) <= 30 THEN 'Current (0-30 Days)'
            WHEN DATEDIFF(DAY, c.sec_update_date, GETDATE()) <= 60 THEN 'Aging (31-60 Days)'
            WHEN DATEDIFF(DAY, c.sec_update_date, GETDATE()) <= 90 THEN 'Overdue (61-90 Days)'
            ELSE 'Critical Default (> 90 Days)'
        END AS Aging_Bucket,
        c.cust_max_credit AS Credit_Ceiling_Limit,
        CAST(c.cust_current_credit / NULLIF(c.cust_max_credit, 0) AS DECIMAL(18,4)) AS Credit_Limit_Utilization,
        c.sec_update_date AS Last_Transaction_Date
    FROM Customer c WITH (NOLOCK)
    WHERE c.cust_active = '1' AND c.cust_current_credit > 0
    ORDER BY c.cust_current_credit DESC;
    """

    # 11. Transaction Basket Size Tier Distribution
    q_basket_size_distribution = """
    SELECT 
        CASE 
            WHEN total_bill_after_disc < 50 THEN '1. Micro (< 50 L.E)'
            WHEN total_bill_after_disc <= 150 THEN '2. Small (50 - 150 L.E)'
            WHEN total_bill_after_disc <= 500 THEN '3. Medium (151 - 500 L.E)'
            WHEN total_bill_after_disc <= 1500 THEN '4. Large (501 - 1500 L.E)'
            ELSE '5. VIP (> 1500 L.E)'
        END AS Basket_Tier,
        COUNT(sth_id) AS Transaction_Count,
        CAST(SUM(total_bill_after_disc) AS DECIMAL(18,2)) AS Total_Revenue,
        CAST(AVG(total_bill_after_disc) AS DECIMAL(18,2)) AS Average_Basket_Value,
        CAST(SUM(total_bill_after_disc) / NULLIF((SELECT SUM(total_bill_after_disc) FROM sales_trans_h WHERE sth_flag='C'), 0) AS DECIMAL(18,4)) AS Share_Of_Total_Sales
    FROM sales_trans_h WITH (NOLOCK)
    WHERE sth_flag = 'C'
    GROUP BY 
        CASE 
            WHEN total_bill_after_disc < 50 THEN '1. Micro (< 50 L.E)'
            WHEN total_bill_after_disc <= 150 THEN '2. Small (50 - 150 L.E)'
            WHEN total_bill_after_disc <= 500 THEN '3. Medium (151 - 500 L.E)'
            WHEN total_bill_after_disc <= 1500 THEN '4. Large (501 - 1500 L.E)'
            ELSE '5. VIP (> 1500 L.E)'
        END
    ORDER BY Basket_Tier;
    """

    # --- SINGLE-PASS IN-MEMORY EXCEL WRITING & FORMATTING ---
    table_names = set()
    fill_hex = "0B132B"  # Dark Navy
    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        write_and_style_sheet(
            writer,
            "Executive_KPI_Summary",
            read_sql_and_enrich(q_kpi_scorecard, ENGINE),
            fill_hex,
            is_kpi=True,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Executive_Monthly_P&L",
            read_sql_and_enrich(q_monthly_pnl, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "ABC_Revenue_Pareto",
            read_sql_and_enrich(q_abc_analysis, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Expiry_&_DeadStock_Radar",
            read_sql_and_enrich(q_dead_stock_radar, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Hourly_Traffic_Patterns",
            read_sql_and_enrich(q_hourly_traffic, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Vendor_Supply_Matrix",
            read_sql_and_enrich(q_vendor_performance, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "MultiBranch_Benchmark",
            read_sql_and_enrich(q_store_matrix, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Monthly_Stock_Valuation_Trend",
            read_sql_and_enrich(q_monthly_stock_trend, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Therapeutic_Class_Profitability",
            read_sql_and_enrich(q_therapeutic_profitability, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Customer_Receivables_Aging",
            read_sql_and_enrich(q_receivables_aging, ENGINE),
            fill_hex,
            table_names=table_names,
        )
        write_and_style_sheet(
            writer,
            "Basket_Size_Distribution",
            read_sql_and_enrich(q_basket_size_distribution, ENGINE),
            fill_hex,
            table_names=table_names,
        )

    logger.info(
        "Module 6 (06_Genius_Executive_Intelligence_Cockpit.xlsx) generated successfully with 11 styled sheets."
    )


# ==================================================================================================
# MASTER WORKFLOW ORCHESTRATOR
# ==================================================================================================
def main():
    start_time = time.time()
    logger.info(
        "=========================================================================="
    )
    logger.info(
        "STARTING GENIUS ERP COMPREHENSIVE DATA CONSOLIDATION PIPELINE (MASTER)"
    )
    logger.info(
        "Universal Historical Scope (All Years Included) | In-Memory OpenPyXL Styling"
    )
    logger.info(
        "=========================================================================="
    )

    output_dir = os.path.join(os.getcwd(), "Genius_Consolidated_Reports")
    os.makedirs(output_dir, exist_ok=True)
    logger.info(f"Target Output Directory: {output_dir}")

    # Connectivity Probe
    try:
        with ENGINE.connect() as conn:
            conn.execute(text("SELECT 1"))
        logger.info("Database Connection Test: SUCCESSFUL.")
    except Exception as e:
        logger.error(f"Failed to connect to Microsoft SQL Server: {e}")
        sys.exit(1)

    target_mod = sys.argv[1].strip() if len(sys.argv) > 1 else "all"

    try:
        if target_mod in ("all", "1"):
            generate_inventory_workbook(output_dir)
        if target_mod in ("all", "2"):
            generate_sales_workbook(output_dir)
        if target_mod in ("all", "3"):
            generate_purchases_workbook(output_dir)
        if target_mod in ("all", "4"):
            generate_financial_workbook(output_dir)
        if target_mod in ("all", "5"):
            generate_shifts_workbook(output_dir)
        if target_mod in ("all", "6"):
            generate_executive_bi_workbook(output_dir)
    except Exception as e:
        logger.error(
            f"Extraction halted due to an unexpected exception: {e}", exc_info=True
        )
        sys.exit(1)

    elapsed = round(time.time() - start_time, 2)
    logger.info(
        "=========================================================================="
    )
    logger.info(
        f"ALL 6 MASTER WORKBOOKS (72 SHEETS TOTAL) SUCCESSFULLY GENERATED IN {elapsed}s."
    )
    logger.info(f"Output Location: {output_dir}")
    logger.info(
        "=========================================================================="
    )


if __name__ == "__main__":
    main()

#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
====================================================================================================
GENIUS (e-Plus) PHARMACY ERP - MASTER DATA CONSOLIDATION & INTELLIGENCE ENGINE (DEFINITIVE EDITION)
====================================================================================================
Consolidated Module Architecture:
  1. Inventory Master & Batch Control Engine  -> 01_Genius_Inventory_Master.xlsx          [13 Sheets]
  2. Sales Operations & Revenue Engine        -> 02_Genius_Sales_Operations.xlsx          [13 Sheets]
  3. Procurement & Vendor Supply Chain        -> 03_Genius_Purchasing_Operations.xlsx      [10 Sheets]
  4. Treasury & General Financial Ledger      -> 04_Genius_Financial_Ledger.xlsx          [11 Sheets]
  5. POS Cashier & Shift Audit Trail          -> 05_Genius_POS_Shift_Operations.xlsx      [09 Sheets]
  6. Executive Business Intelligence Cockpit  -> 06_Genius_Executive_Intelligence_Cockpit.xlsx [11 Sheets]
====================================================================================================
Standards Enforced:
  - Universal Year 2024 Exclusion across all transaction logs and analytical models.
  - Multi-Unit Pack Normalization: Blisters/strips/tablets normalized to standard boxes in Pareto.
  - Varbinary Decryption Engine: Bidi buffer reversal + windows-1256 item trade name restoration.
  - Complete Reference Mappings: Includes Clinical Indications, Generic Alternates, and UCP Catalogs.
  - Native Excel Tables: Auto-filters spanning exactly data rows, frozen headers, ZERO total rows.
  - Exact Accounting Sign Conventions: AP (ven_current_credit > 0), AR (cust_current_credit > 0).
  - True Weighted Average Cost: SUM(qty * cost) / NULLIF(SUM(qty), 0).
  - Type-Coercion Safety: All user audit links use TRY_CAST to prevent varchar-to-numeric crashes.
====================================================================================================
"""

import os
import sys
import time
import re
import logging
import urllib.parse
from datetime import datetime
import pandas as pd
from sqlalchemy import create_engine, text
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# --------------------------------------------------------------------------------------------------
# 0. LOGGING CONFIGURATION
# --------------------------------------------------------------------------------------------------
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[logging.StreamHandler(sys.stdout)],
)
logger = logging.getLogger("GeniusMasterEngine")

# --------------------------------------------------------------------------------------------------
# 1. DATABASE CONNECTION SETUP
# --------------------------------------------------------------------------------------------------
SERVER = "localhost,1433"
DATABASE = "Genius"
USERNAME = "sa"
PASSWORD = "Emma@2015"

conn_str = (
    f"DRIVER={{ODBC Driver 18 for SQL Server}};"
    f"SERVER={SERVER};"
    f"DATABASE={DATABASE};"
    f"UID={USERNAME};"
    f"PWD={PASSWORD};"
    f"TrustServerCertificate=yes;"
)
params = urllib.parse.quote_plus(conn_str)
ENGINE = create_engine(f"mssql+pyodbc:///?odbc_connect={params}", pool_pre_ping=True)


# --------------------------------------------------------------------------------------------------
# 1.1 VARBINARY ITEM DECRYPTION ENGINE (GENIUS ARCHITECTURAL CORE)
# --------------------------------------------------------------------------------------------------
def decode_legacy_bidi_bytes(raw_bytes):
    """
    Decodes obfuscated varbinary medicine trade names in Item_Catalog.
    Splits binary buffer in half (N // 2), reverses each half independently,
    concatenates back together, and decodes via windows-1256.
    """
    if pd.isna(raw_bytes) or not raw_bytes:
        return ""
    try:
        raw_b = bytes(raw_bytes)
        N = len(raw_b)
        L = N // 2
        corrected_bytes = raw_b[:L][::-1] + raw_b[L:][::-1]
        return corrected_bytes.decode("windows-1256", errors="replace").strip()
    except Exception:
        return "Decoding Error"


ITEM_NAME_AR_CACHE = None
ITEM_NAME_EN_CACHE = None


def get_item_name_cache(engine):
    """
    Builds a high-performance in-memory cache mapping all possible representations of itm_id
    (float, int, and string) to decrypted Arabic and English trade names.
    """
    global ITEM_NAME_AR_CACHE, ITEM_NAME_EN_CACHE
    if ITEM_NAME_AR_CACHE is not None and ITEM_NAME_EN_CACHE is not None:
        return ITEM_NAME_AR_CACHE, ITEM_NAME_EN_CACHE

    logger.info("Initializing Item Decryption Cache from Item_Catalog...")
    q = """
    SELECT itm_id, itm_name_ar, itm_name_en, itm_name_ar_encrypt, itm_name_en_encrypt 
    FROM Item_Catalog WITH (NOLOCK)
    """
    try:
        df_cat = pd.read_sql(q, engine)
        cache_ar = {}
        cache_en = {}

        for _, row in df_cat.iterrows():
            iid = row["itm_id"]
            if pd.isna(iid):
                continue

            # Arabic resolution
            name_ar = row["itm_name_ar"]
            if pd.isna(name_ar) or not str(name_ar).strip():
                name_ar = decode_legacy_bidi_bytes(row["itm_name_ar_encrypt"])
            else:
                name_ar = str(name_ar).strip()

            # English resolution
            name_en = row["itm_name_en"]
            if pd.isna(name_en) or not str(name_en).strip():
                name_en = decode_legacy_bidi_bytes(row["itm_name_en_encrypt"])
            else:
                name_en = str(name_en).strip()

            # Cross fallbacks
            if not name_en and name_ar:
                name_en = name_ar
            if not name_ar and name_en:
                name_ar = name_en

            # Register numeric and string keys to ensure 100% join hit-rate
            try:
                iid_f = float(iid)
                iid_i = int(iid_f)
                cache_ar[iid_f] = name_ar
                cache_ar[iid_i] = name_ar
                cache_ar[str(iid_i)] = name_ar
                cache_ar[str(iid_f)] = name_ar

                cache_en[iid_f] = name_en
                cache_en[iid_i] = name_en
                cache_en[str(iid_i)] = name_en
                cache_en[str(iid_f)] = name_en
            except ValueError, TypeError:
                cache_ar[iid] = name_ar
                cache_ar[str(iid)] = name_ar
                cache_en[iid] = name_en
                cache_en[str(iid)] = name_en

        ITEM_NAME_AR_CACHE = cache_ar
        ITEM_NAME_EN_CACHE = cache_en
        logger.info(
            f"Item Decryption Cache successfully built: {len(df_cat):,} items mapped."
        )
    except Exception as e:
        logger.warning(f"Could not build item decryption cache: {e}")
        ITEM_NAME_AR_CACHE = {}
        ITEM_NAME_EN_CACHE = {}

    return ITEM_NAME_AR_CACHE, ITEM_NAME_EN_CACHE


def read_sql_and_enrich(sql, engine):
    """
    Executes SQL queries and maps decrypted Arabic/English item names.
    Supports primary item IDs, transfer items, and clinical substitute IDs.
    """
    df = pd.read_sql(sql, engine)
    cache_ar, cache_en = get_item_name_cache(engine)

    # Primary Item Candidate Identification
    id_col = None
    for candidate in [
        "Item_ID",
        "itm_id",
        "st_itm_id",
        "invd_itm_id",
        "iqc_itm_id",
        "ptd_itm_id",
        "Original_Item_ID",
    ]:
        if candidate in df.columns:
            id_col = candidate
            break

    if id_col is not None:
        for ar_col in [
            "Item_Name",
            "Item_Name_Ar",
            "Item_Name_ar",
            "Original_Item_Name",
        ]:
            if ar_col in df.columns:
                mapped = df[id_col].map(cache_ar)
                df[ar_col] = (
                    df[ar_col].replace(r"^\s*$", None, regex=True).fillna(mapped)
                )

        for en_col in ["Item_Name_En", "Item_Name_en", "Original_Item_Name_En"]:
            if en_col in df.columns:
                mapped = df[id_col].map(cache_en)
                df[en_col] = (
                    df[en_col].replace(r"^\s*$", None, regex=True).fillna(mapped)
                )

    # Secondary Item Mapping (Substitutes / Alternates)
    if "Substitute_Item_ID" in df.columns:
        if "Substitute_Item_Name" in df.columns:
            mapped_sub_ar = df["Substitute_Item_ID"].map(cache_ar)
            df["Substitute_Item_Name"] = (
                df["Substitute_Item_Name"]
                .replace(r"^\s*$", None, regex=True)
                .fillna(mapped_sub_ar)
            )
        if "Substitute_Item_Name_En" in df.columns:
            mapped_sub_en = df["Substitute_Item_ID"].map(cache_en)
            df["Substitute_Item_Name_En"] = (
                df["Substitute_Item_Name_En"]
                .replace(r"^\s*$", None, regex=True)
                .fillna(mapped_sub_en)
            )

    return df


# --------------------------------------------------------------------------------------------------
# 2. OPENPYXL ENTERPRISE STYLING HELPERS (ZERO TOTAL ROWS)
# --------------------------------------------------------------------------------------------------
THIN_BORDER = Border(
    left=Side(style="thin", color="D3D3D3"),
    right=Side(style="thin", color="D3D3D3"),
    top=Side(style="thin", color="D3D3D3"),
    bottom=Side(style="thin", color="D3D3D3"),
)

COLOR_TO_TABLE_STYLE = {
    "1B4D3E": "TableStyleMedium7",  # Emerald Green (Inventory)
    "1F4E79": "TableStyleMedium9",  # Sapphire Blue (Sales)
    "4A154B": "TableStyleMedium12",  # Deep Burgundy (Purchasing)
    "7D5A00": "TableStyleMedium13",  # Bronze Gold (Financials)
    "1A5276": "TableStyleMedium10",  # Slate Teal (POS Shifts)
    "0B132B": "TableStyleMedium1",  # Dark Navy (Executive BI)
}


def is_percentage_col(header_name: str) -> bool:
    """Robust percentage column detection avoiding false-positive matches on financial totals."""
    h = header_name.lower().strip()
    if any(
        h.endswith(sfx)
        for sfx in ["_perc", "_rate", "_ratio", "_percent", "_percentage", "%"]
    ):
        return True
    if any(h.startswith(pfx) for pfx in ["perc_", "rate_", "ratio_"]):
        return True
    return bool(
        re.search(r"(?:^|_)(perc|rate|ratio|percent|percentage)(?:_|$)|%", h)
    ) and not any(
        term in h
        for term in [
            "revenue",
            "sales",
            "cogs",
            "profit",
            "amount",
            "total",
            "val",
            "cost",
            "price",
            "fee",
        ]
    )


def style_worksheet(
    ws, header_fill_hex="1F4E79", is_kpi=False, existing_table_names=None
):
    """
    Applies clean corporate styling:
      - Native Excel Table spanning exactly A1:{max_col}{max_row} with auto-filters.
      - Freezes top header row, sets visible gridlines, auto-formats numbers.
      - Auto-calculates column widths (12 to 50 characters).
      - Zero totals row clutter appended at the bottom.
    """
    if existing_table_names is None:
        existing_table_names = set()

    max_row = ws.max_row
    max_col = ws.max_column

    # Excel tables require at least 1 data row beneath headers
    if max_row < 2:
        ws.append([None] * max_col)
        max_row = 2

    data_font = Font(name="Segoe UI", size=9)
    ws.views.sheetView[0].showGridLines = True
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 26

    # Header Uniqueness
    headers = []
    seen_headers = {}
    for c in range(1, max_col + 1):
        raw_val = ws.cell(row=1, column=c).value
        cell_val = str(raw_val if raw_val is not None else f"Col_{c}").strip()
        if not cell_val:
            cell_val = f"Col_{c}"

        if cell_val in seen_headers:
            seen_headers[cell_val] += 1
            unique_h = f"{cell_val}_{seen_headers[cell_val]}"
        else:
            seen_headers[cell_val] = 0
            unique_h = cell_val

        ws.cell(row=1, column=c, value=unique_h)
        headers.append(unique_h)

    col_widths = {}
    for col_idx in range(1, max_col + 1):
        header_val = headers[col_idx - 1]
        max_len = len(header_val)

        for row_idx in range(2, max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            val = cell.value
            cell.font = data_font
            cell.border = THIN_BORDER

            if val is not None:
                str_val = str(val)
                if len(str_val) > max_len:
                    max_len = len(str_val)

                if isinstance(val, (int, float)):
                    if is_percentage_col(header_val):
                        cell.number_format = "0.00%"
                        cell.alignment = Alignment(
                            horizontal="right", vertical="center"
                        )
                    elif any(
                        term in header_val.lower()
                        for term in [
                            "price",
                            "cost",
                            "val",
                            "total",
                            "net",
                            "credit",
                            "debit",
                            "balance",
                            "revenue",
                            "cogs",
                            "profit",
                            "cash",
                            "amount",
                            "paid",
                            "refund",
                            "fee",
                            "disbursed",
                            "salary",
                            "bonus",
                            "bouns",
                            "pen",
                            "gross",
                        ]
                    ):
                        cell.number_format = "#,##0.00"
                        cell.alignment = Alignment(
                            horizontal="right", vertical="center"
                        )
                    elif any(
                        term in header_val.lower()
                        for term in [
                            "qty",
                            "count",
                            "items",
                            "no_",
                            "units",
                            "lines",
                            "days",
                            "hours",
                        ]
                    ):
                        cell.number_format = "#,##0.00"
                        cell.alignment = Alignment(
                            horizontal="right", vertical="center"
                        )
                    else:
                        cell.alignment = Alignment(
                            horizontal="center", vertical="center"
                        )
                elif isinstance(val, (datetime, pd.Timestamp)):
                    cell.number_format = "yyyy-mm-dd hh:mm"
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        col_letter = get_column_letter(col_idx)
        col_widths[col_letter] = min(max(max_len + 4, 12), 50)

    # Native Excel Table Definition
    clean_title = re.sub(r"[^a-zA-Z0-9_]", "_", ws.title).strip("_")
    if not clean_title or clean_title[0].isdigit():
        clean_title = f"T_{clean_title}"
    base_name = f"Tab_{clean_title}"[:200]
    candidate = base_name
    counter = 1
    while candidate.lower() in existing_table_names:
        candidate = f"{base_name}_{counter}"
        counter += 1
    existing_table_names.add(candidate.lower())

    tab_ref = f"A1:{get_column_letter(max_col)}{max_row}"
    tab = Table(displayName=candidate, ref=tab_ref)
    table_style = COLOR_TO_TABLE_STYLE.get(header_fill_hex, "TableStyleMedium9")
    tab.tableStyleInfo = TableStyleInfo(
        name=table_style,
        showRowStripes=True,
        showFirstColumn=False,
        showLastColumn=False,
        showColumnStripes=False,
    )
    ws.add_table(tab)

    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width


# ==================================================================================================
# MODULE 1: INVENTORY MASTER & BATCH CONTROL ENGINE
# ==================================================================================================
def generate_inventory_workbook(output_dir):
    filename = os.path.join(output_dir, "01_Genius_Inventory_Master.xlsx")
    logger.info(f"Generating Module 1: {filename}...")

    # 1. Active Store Batches (EAN-13 Barcode + Manufacturer Lot + Branch Roles)
    q_batches = """
    SELECT 
        ics.sto_id AS Store_ID,
        st.sto_name_ar AS Store_Name,
        CASE WHEN ics.sto_id = 47 THEN 'Dispensary Retail (صيدلية)' ELSE 'Consignment / Quarantine (تحت التصريف)' END AS Store_Type,
        ics.itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        (SELECT TOP 1 io_itm_int FROM Item_Objects WITH (NOLOCK) WHERE io_itm_id = ics.itm_id) AS International_Barcode,
        ic.itm_name_ar AS Item_Name_Ar,
        ic.itm_name_en AS Item_Name_En,
        ics.c_id AS Batch_Class_ID,
        icl.itm_batch_no AS Manufacturer_Batch_No,
        ics.itm_qty AS Batch_Quantity,
        u.u_name_ar AS Base_Unit,
        ics.pharm_price AS Unit_Cost_Price,
        ics.sell_price AS Unit_Sell_Price,
        ics.sell_tax AS Unit_Tax,
        CAST(ics.itm_qty * ics.pharm_price AS DECIMAL(18,2)) AS Total_Batch_Cost_Valuation,
        CAST(ics.itm_qty * ics.sell_price AS DECIMAL(18,2)) AS Total_Batch_Retail_Valuation,
        CONVERT(VARCHAR(10), ics.itm_expiry_date, 120) AS Expiry_Date,
        v.ven_name_ar AS Preferred_Vendor,
        ics.itm_location AS Rack_Location
    FROM Item_Class_Store ics WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = ics.itm_id
    JOIN Store st WITH (NOLOCK) ON st.sto_id = ics.sto_id
    LEFT JOIN Item_Class icl WITH (NOLOCK) ON icl.c_id = ics.c_id AND icl.itm_id = ics.itm_id AND icl.stock_id = ics.sto_id
    LEFT JOIN Units u WITH (NOLOCK) ON u.u_id = ic.itm_unit1
    LEFT JOIN Vendor v WITH (NOLOCK) ON v.ven_id = ics.ven_id
    WHERE ics.itm_qty > 0
    ORDER BY ics.sto_id, ics.itm_id;
    """

    # 2. Consolidated Stock Valuation per Item (True Weighted Average Cost)
    q_consolidated = """
    SELECT 
        ic.itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        (SELECT TOP 1 io_itm_int FROM Item_Objects WITH (NOLOCK) WHERE io_itm_id = ic.itm_id) AS International_Barcode,
        ic.itm_name_ar AS Item_Name_Ar,
        c.com_name_ar AS Manufacturer,
        COUNT(DISTINCT ics.c_id) AS Active_Batch_Count,
        SUM(ics.itm_qty) AS Total_Stock_Units,
        CAST(SUM(ics.itm_qty * ics.pharm_price) / NULLIF(SUM(ics.itm_qty), 0) AS DECIMAL(18,4)) AS True_Weighted_Avg_Cost,
        MAX(ics.sell_price) AS Current_Retail_Price,
        SUM(CAST(ics.itm_qty * ics.pharm_price AS DECIMAL(18,2))) AS Consolidated_Cost_Value,
        SUM(CAST(ics.itm_qty * ics.sell_price AS DECIMAL(18,2))) AS Consolidated_Retail_Value,
        MIN(ics.itm_expiry_date) AS Nearest_Expiry
    FROM Item_Class_Store ics WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = ics.itm_id
    LEFT JOIN Company c WITH (NOLOCK) ON c.com_id = ic.com_id
    WHERE ics.itm_qty > 0
    GROUP BY ic.itm_id, ic.itm_code, ic.itm_name_ar, c.com_name_ar
    ORDER BY Consolidated_Cost_Value DESC;
    """

    # 3. Inter-Store Transfers (Composite Key Join & Safe User Coercion)
    q_transfers = """
    SELECT 
        sth.stnh_id AS Transfer_ID,
        CASE sth.stnh_flag 
            WHEN 'R' THEN 'Received (مستلم)' 
            WHEN 'S' THEN 'In Transit (مرسل)' 
            WHEN 'A' THEN 'Approved (معتمد)' 
            WHEN 'P' THEN 'Pending (معلق)' 
            ELSE sth.stnh_flag 
        END AS Transfer_Status,
        fs.sto_name_ar AS From_Store,
        ts.sto_name_ar AS To_Store,
        st.st_itm_id AS Item_ID,
        ic.itm_name_ar AS Item_Name,
        st.st_itm_quantity AS Transferred_Qty,
        st.st_itm_cost AS Line_Unit_Cost,
        st.st_itm_sell AS Line_Unit_Sell,
        CAST(st.st_itm_quantity * st.st_itm_cost AS DECIMAL(18,2)) AS Total_Line_Cost,
        CAST(st.st_itm_quantity * st.st_itm_sell AS DECIMAL(18,2)) AS Total_Line_Retail,
        sth.sec_insert_date AS Transfer_Date,
        COALESCE(emp.e_Name, usr.usr_name, CAST(sth.sec_insert_uid AS VARCHAR(50))) AS Operator
    FROM Store_Trans_h sth WITH (NOLOCK)
    JOIN Store_Trans st WITH (NOLOCK) ON st.stnh_id = sth.stnh_id AND st.st_from_store = sth.stnh_f_Sto_id
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = st.st_itm_id
    JOIN Store fs WITH (NOLOCK) ON fs.sto_id = sth.stnh_f_Sto_id
    JOIN Store ts WITH (NOLOCK) ON ts.sto_id = sth.stnh_t_Sto_id
    LEFT JOIN Users usr WITH (NOLOCK) ON usr.usr_e_id = TRY_CAST(sth.sec_insert_uid AS DECIMAL(18,0))
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(sth.sec_insert_uid AS DECIMAL(18,0))
    WHERE YEAR(sth.sec_insert_date) <> 2024
    ORDER BY sth.stnh_id DESC;
    """

    # 4. Physical Stock-Take Audits
    q_audits = """
    SELECT 
        ih.inv_id AS Audit_Doc_ID,
        st.sto_name_ar AS Store_Name,
        ih.inv_begin_time AS Audit_Timestamp,
        ih.inv_flag AS Doc_State,
        id.invd_itm_id AS Item_ID,
        ic.itm_name_ar AS Item_Name,
        id.invd_c_id AS Batch_ID,
        id.invd_old_qty AS System_Expected_Qty,
        id.invd_new_qty AS Physical_Counted_Qty,
        (id.invd_new_qty - id.invd_old_qty) AS Variance_Qty,
        id.invd_pharm_price AS Unit_Cost,
        CAST((id.invd_new_qty - id.invd_old_qty) * id.invd_pharm_price AS DECIMAL(18,2)) AS Financial_Cost_Impact,
        COALESCE(emp.e_Name, CAST(ih.sec_insert_uid AS VARCHAR(50))) AS Audited_By
    FROM Inventory_h ih WITH (NOLOCK)
    JOIN Inventory_d id WITH (NOLOCK) ON id.invd_inv_id = ih.inv_id
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = id.invd_itm_id
    JOIN Store st WITH (NOLOCK) ON st.sto_id = ih.inv_sto_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(ih.sec_insert_uid AS DECIMAL(18,0))
    WHERE YEAR(ih.inv_begin_time) <> 2024
    ORDER BY ih.inv_id DESC;
    """

    # 5. Quantity Adjustments & Breakages
    q_adjustments = """
    SELECT 
        iqc.iqc_id AS Adj_ID,
        st.sto_name_ar AS Store_Name,
        iqc.iqc_itm_id AS Item_ID,
        ic.itm_name_ar AS Item_Name,
        dc.dc_name_ar AS Cause_Description,
        iqc.iqc_old_quantity AS Old_Quantity,
        iqc.iqc_new_quantity AS New_Quantity,
        (iqc.iqc_new_quantity - iqc.iqc_old_quantity) AS Quantity_Delta,
        iqc.iqc_pharm_price AS Cost_Price,
        CAST((iqc.iqc_new_quantity - iqc.iqc_old_quantity) * iqc.iqc_pharm_price AS DECIMAL(18,2)) AS Adjustment_Cost_Value,
        iqc.iqc_notes AS Notes,
        iqc.sec_insert_date AS Action_Timestamp
    FROM Item_Quantity_Changes iqc WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = iqc.iqc_itm_id
    LEFT JOIN Store st WITH (NOLOCK) ON st.sto_id = iqc.iqc_store_id
    LEFT JOIN destroy_causes dc WITH (NOLOCK) ON dc.dc_id = iqc.iqc_dc_id
    WHERE YEAR(iqc.sec_insert_date) <> 2024
    ORDER BY iqc.iqc_id DESC;
    """

    # 6. Opening Balances
    q_opening = """
    SELECT 
        so.so_id AS Opening_Doc_ID,
        st.sto_name_ar AS Store_Name,
        sod.itm_id AS Item_ID,
        ic.itm_name_ar AS Item_Name,
        sod.c_id AS Batch_ID,
        sod.itm_qty AS Opening_Quantity,
        sod.pharm_price AS Cost_Price,
        sod.sell_price AS Retail_Price,
        sod.itm_expiry_date AS Expiry_Date,
        so.sec_insert_date AS Setup_Date
    FROM Stock_Open so WITH (NOLOCK)
    JOIN Stock_Open_d sod WITH (NOLOCK) ON sod.so_id = so.so_id
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = sod.itm_id
    JOIN Store st WITH (NOLOCK) ON st.sto_id = so.so_sto_id
    WHERE YEAR(so.sec_insert_date) <> 2024
    ORDER BY so.so_id DESC;
    """

    # 7. Forensic Inventory Audit Trail
    q_forensic_stock = """
    SELECT TOP 10000
        w.id AS Audit_Log_ID,
        st.sto_name_ar AS Store_Name,
        w.itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        ic.itm_name_ar AS Item_Name,
        w.c_id AS Batch_Class_ID,
        w.old_qty AS Previous_Quantity,
        w.new_qty AS Adjusted_Quantity,
        CAST(w.new_qty - ISNULL(w.old_qty, 0) AS DECIMAL(18,2)) AS Quantity_Variance,
        w.oper AS Operation_Type,
        w.prog_nm AS Software_Source,
        w.pc_nm AS Client_Machine_Host,
        CONVERT(VARCHAR(10), w.itm_expiry_date, 120) AS Expiry_Date,
        w.insert_dt AS Log_Timestamp
    FROM watch_qty_chng w WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = w.itm_id
    LEFT JOIN Store st WITH (NOLOCK) ON st.sto_id = w.sto_id
    WHERE YEAR(w.insert_dt) <> 2024
    ORDER BY w.id DESC;
    """

    # 8. Shortages Catalog Registry
    q_shortages = """
    SELECT 
        s.str_id AS Shortage_ID,
        st.sto_name_ar AS Store_Name,
        s.str_itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        ic.itm_name_ar AS Item_Name,
        CASE s.str_type 
            WHEN 'S' THEN 'System Suggested Shortage' 
            WHEN 'B' THEN 'Cashier/POS Request' 
            WHEN 'I' THEN 'Inventory Review' 
            ELSE s.str_type 
        END AS Request_Type,
        s.str_requested_qty AS Requested_Units,
        CASE s.str_requested WHEN 1 THEN 'Ordered/Fulfilled' ELSE 'Pending Demand' END AS Fulfillment_Status,
        s.str_notes AS Operational_Notes,
        s.sec_insert_date AS Request_Date,
        COALESCE(emp.e_Name, CAST(s.sec_insert_uid AS VARCHAR(50))) AS Requested_By
    FROM Shortage s WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = s.str_itm_id
    LEFT JOIN Store st WITH (NOLOCK) ON st.sto_id = s.sto_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(s.sec_insert_uid AS DECIMAL(18,0))
    WHERE YEAR(s.sec_insert_date) <> 2024
    ORDER BY s.str_id DESC;
    """

    # 9. Expiry Date Change Audit
    q_expiry_changes = """
    SELECT 
        iec.iec_id AS Expiry_Audit_ID,
        st.sto_name_ar AS Store_Name,
        iec.iec_itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        ic.itm_name_ar AS Item_Name,
        iec.iec_c_id AS Batch_Class_ID,
        CONVERT(VARCHAR(10), iec.iec_old_expire, 120) AS Old_Expiry_Date,
        CONVERT(VARCHAR(10), iec.iec_new_expire, 120) AS New_Expiry_Date,
        iec.iec_old_qty AS Stock_Quantity_Affected,
        iec.iec_itm_sell_price AS Retail_Price,
        CAST(iec.iec_old_qty * iec.iec_itm_sell_price AS DECIMAL(18,2)) AS Total_Valuation_Impact,
        iec.sec_insert_date AS Change_Timestamp,
        COALESCE(emp.e_Name, CAST(iec.sec_insert_uid AS VARCHAR(50))) AS Modified_By
    FROM Item_Expire_Change iec WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = iec.iec_itm_id
    LEFT JOIN Store st WITH (NOLOCK) ON st.sto_id = iec.iec_sto_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(iec.sec_insert_uid AS DECIMAL(18,0))
    WHERE YEAR(iec.sec_insert_date) <> 2024
    ORDER BY iec.iec_id DESC;
    """

    # 10. Medicine Repricing & Inventory Gain
    q_repricing = """
    SELECT 
        ch.ich_id AS Reprice_Audit_ID,
        ch.ich_itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        ic.itm_name_ar AS Item_Name,
        ch.ich_sell_price_old AS Old_Retail_Price,
        ch.ich_sell_price AS New_Official_Price,
        CAST(ch.ich_sell_price - ch.ich_sell_price_old AS DECIMAL(18,2)) AS Price_Delta_Per_Unit,
        ch.ich_pharm_price_old AS Old_Cost,
        ch.ich_pharm_price AS New_Cost,
        ch.ich_qty AS Stock_Quantity_At_Reprice,
        CAST(ch.ich_qty * (ch.ich_sell_price - ch.ich_sell_price_old) AS DECIMAL(18,2)) AS Revaluation_Gain_Retail,
        CAST(ch.ich_qty * (ch.ich_pharm_price - ch.ich_pharm_price_old) AS DECIMAL(18,2)) AS Revaluation_Gain_Cost,
        ch.sec_insert_date AS Repricing_Timestamp,
        COALESCE(emp.e_Name, CAST(ch.sec_insert_uid AS VARCHAR(50))) AS Authorized_By
    FROM Item_Changes ch WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = ch.ich_itm_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(ch.sec_insert_uid AS DECIMAL(18,0))
    WHERE (ch.ich_sell_price - ch.ich_sell_price_old) <> 0
      AND YEAR(ch.sec_insert_date) <> 2024
    ORDER BY ch.sec_insert_date DESC;
    """

    # 11. Clinical Usage Indications Registry
    q_usage_indications = """
    SELECT 
        iuc.iuc_itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        ic.itm_name_ar AS Item_Name,
        uc.uc_name_ar AS Medical_Indication_Ar,
        uc.uc_name_en AS Medical_Indication_En
    FROM Item_Usage_Causes iuc WITH (NOLOCK)
    JOIN Usage_Causes uc WITH (NOLOCK) ON uc.uc_id = iuc.iuc_uc_id
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = iuc.iuc_itm_id
    ORDER BY uc.uc_name_ar, ic.itm_name_ar;
    """

    # 12. Generic Medicine Substitutions & Clinical Alternates
    q_substitutions = """
    SELECT 
        ia.itm_id AS Original_Item_ID,
        ic1.itm_code AS Original_Code,
        ic1.itm_name_ar AS Original_Item_Name,
        c1.com_name_ar AS Original_Manufacturer,
        ic1.itm_def_sell_price AS Original_Retail_Price,
        ia.alt_itm_id AS Substitute_Item_ID,
        ic2.itm_code AS Substitute_Code,
        ic2.itm_name_ar AS Substitute_Item_Name,
        c2.com_name_ar AS Substitute_Manufacturer,
        ic2.itm_def_sell_price AS Substitute_Retail_Price,
        CAST(ic2.itm_def_sell_price - ic1.itm_def_sell_price AS DECIMAL(18,2)) AS Price_Difference,
        ia.sec_insert_date AS Linked_Date
    FROM Item_Alternates ia WITH (NOLOCK)
    JOIN Item_Catalog ic1 WITH (NOLOCK) ON ic1.itm_id = ia.itm_id
    JOIN Item_Catalog ic2 WITH (NOLOCK) ON ic2.itm_id = ia.alt_itm_id
    LEFT JOIN Company c1 WITH (NOLOCK) ON c1.com_id = ic1.com_id
    LEFT JOIN Company c2 WITH (NOLOCK) ON c2.com_id = ic2.com_id
    ORDER BY ic1.itm_name_ar;
    """

    # 13. System Governance & Inventory Costing Policies
    q_sys_governance = """
    SELECT 
        rec_name AS Pharmacy_Legal_Name,
        rec_tel AS Contact_Phone,
        rec_loc AS Registered_Address,
        ret_days_No AS Policy_Return_Days_Allowed,
        CASE prevent_expired_sales WHEN 1 THEN 'Enabled (بيع المنتهي محظور)' ELSE 'Disabled' END AS Block_Expired_Sales,
        CASE stock_avg_cost WHEN 1 THEN 'Enabled (حساب متوسط التكلفة)' ELSE 'Disabled' END AS Weighted_Average_Costing,
        CASE pur_extra_disc_update_stock WHEN 1 THEN 'Enabled (خصم الشراء الإضافي يحدّث التكلفة)' ELSE 'Disabled' END AS Purchase_Disc_Updates_Cost,
        last_month_close_dt AS Last_Monthly_Inventory_Close,
        rec_footer AS Receipt_Terms_Footer
    FROM Sys_setting WITH (NOLOCK);
    """

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        read_sql_and_enrich(q_batches, ENGINE).to_excel(
            writer, sheet_name="Current_Store_Batches", index=False
        )
        read_sql_and_enrich(q_consolidated, ENGINE).to_excel(
            writer, sheet_name="Consolidated_Inventory", index=False
        )
        read_sql_and_enrich(q_transfers, ENGINE).to_excel(
            writer, sheet_name="Stock_Transfers", index=False
        )
        read_sql_and_enrich(q_audits, ENGINE).to_excel(
            writer, sheet_name="Physical_Audits_Variance", index=False
        )
        read_sql_and_enrich(q_adjustments, ENGINE).to_excel(
            writer, sheet_name="Quantity_Adjustments", index=False
        )
        read_sql_and_enrich(q_opening, ENGINE).to_excel(
            writer, sheet_name="Opening_Stock", index=False
        )
        read_sql_and_enrich(q_forensic_stock, ENGINE).to_excel(
            writer, sheet_name="Inventory_Forensic_Audit", index=False
        )
        read_sql_and_enrich(q_shortages, ENGINE).to_excel(
            writer, sheet_name="Shortages_Catalog_Registry", index=False
        )
        read_sql_and_enrich(q_expiry_changes, ENGINE).to_excel(
            writer, sheet_name="Expiry_Date_Change_Audit", index=False
        )
        read_sql_and_enrich(q_repricing, ENGINE).to_excel(
            writer, sheet_name="Medicine_Repricing_Gains", index=False
        )
        read_sql_and_enrich(q_usage_indications, ENGINE).to_excel(
            writer, sheet_name="Clinical_Usage_Indications", index=False
        )
        read_sql_and_enrich(q_substitutions, ENGINE).to_excel(
            writer, sheet_name="Generic_Substitutions_Mapping", index=False
        )
        read_sql_and_enrich(q_sys_governance, ENGINE).to_excel(
            writer, sheet_name="System_Governance_Policies", index=False
        )

    wb = openpyxl.load_workbook(filename)
    table_names = set()
    for ws_name in wb.sheetnames:
        style_worksheet(
            wb[ws_name],
            header_fill_hex="1B4D3E",
            is_kpi=False,
            existing_table_names=table_names,
        )
    wb.save(filename)
    logger.info("Module 1 Complete.")


# ==================================================================================================
# MODULE 2: SALES OPERATIONS & REVENUE ENGINE
# ==================================================================================================
def generate_sales_workbook(output_dir):
    filename = os.path.join(output_dir, "02_Genius_Sales_Operations.xlsx")
    logger.info(f"Generating Module 2: {filename}...")

    # 1. Sales Headers (Detailed bill types including Insurance Syndicate contract flags)
    q_sales_h = """
    SELECT 
        h.sth_id AS Bill_ID,
        st.sto_name_ar AS Store_Name,
        h.sec_insert_date AS Invoice_Date,
        CASE h.bill_typ 
            WHEN 1 THEN 'Cash (نقدي)' 
            WHEN 2 THEN 'Credit (آجل)' 
            WHEN 3 THEN 'Delivery (توصيل)' 
            WHEN 4 THEN 'Visa (فيزا)' 
            WHEN 5 THEN 'Insurance (تأمين)'
            ELSE 'Other' 
        END AS Bill_Type,
        ISNULL(c.cust_name_ar, 'عميل نقدي (Cash Walk-in)') AS Customer_Name,
        ct.ct_name_ar AS Customer_Type,
        h.no_of_items AS Total_Lines,
        h.total_bill AS Gross_Amount,
        h.total_des_mon AS Cash_Discount,
        h.total_dis_per AS Percentage_Discount,
        h.total_bill_after_disc AS Net_Invoice_Amount,
        h.total_tax AS Tax_Amount,
        h.sth_cash AS Paid_Cash,
        h.sth_rest AS Customer_Remaining_Debt,
        h.sth_extra_expenses AS Delivery_Service_Fee,
        COALESCE(emp.e_Name, CAST(h.emp_id AS VARCHAR(50))) AS Cashier_Name,
        h.sth_pc_name AS Station_Hostname
    FROM sales_trans_h h WITH (NOLOCK)
    JOIN Store st WITH (NOLOCK) ON st.sto_id = h.sto_id
    LEFT JOIN Customer c WITH (NOLOCK) ON c.cust_id = h.cust_id
    LEFT JOIN Customer_Type ct WITH (NOLOCK) ON ct.ct_id = c.cust_type
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(h.emp_id AS DECIMAL(18,0))
    WHERE h.sth_flag = 'C' AND YEAR(h.sec_insert_date) <> 2024
    ORDER BY h.sth_id DESC;
    """

    # 2. Sales Line Items (Accurate Return Deductions & Pro-rated Header Discounts)
    q_sales_d = """
    SELECT 
        d.sth_id AS Bill_ID,
        d.std_id AS Line_No,
        h.sec_insert_date AS Sale_Date,
        d.itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        (SELECT TOP 1 io_itm_int FROM Item_Objects WITH (NOLOCK) WHERE io_itm_id = d.itm_id) AS International_Barcode,
        ic.itm_name_ar AS Item_Name,
        d.c_id AS Batch_Class_ID,
        d.exp_date AS Batch_Expiry,
        d.qnty AS Dispensed_Quantity,
        ISNULL(d.itm_back, 0) AS Returned_Quantity,
        (d.qnty - ISNULL(d.itm_back, 0)) AS Net_Sold_Units,
        d.itm_unit AS Unit_Level,
        d.itm_sell AS Unit_Sale_Price,
        d.itm_cost AS Unit_Cost_COGS,
        d.itm_dis_per AS Item_Discount_Perc,
        d.itm_dis_mon AS Item_Discount_Cash,
        CAST(
            (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell * (1.0 - (ISNULL(d.itm_dis_per, 0) / 100.0))) - ISNULL(d.itm_dis_mon, 0))
            * (1.0 - (ISNULL(h.total_dis_per, 0) / 100.0))
            - ISNULL(ISNULL(h.total_des_mon, 0) / NULLIF(h.no_of_items, 0), 0)
        AS DECIMAL(18,2)) AS Net_Line_Revenue,
        CAST((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_cost AS DECIMAL(18,2)) AS Net_Line_COGS,
        CAST(
            ((((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell * (1.0 - (ISNULL(d.itm_dis_per, 0) / 100.0))) - ISNULL(d.itm_dis_mon, 0))
            * (1.0 - (ISNULL(h.total_dis_per, 0) / 100.0))
            - ISNULL(ISNULL(h.total_des_mon, 0) / NULLIF(h.no_of_items, 0), 0))
            - ((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_cost)
        AS DECIMAL(18,2)) AS Realized_Gross_Profit
    FROM sales_trans_d d WITH (NOLOCK)
    JOIN sales_trans_h h WITH (NOLOCK) ON h.sth_id = d.sth_id
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = d.itm_id
    WHERE h.sth_flag = 'C' AND YEAR(h.sec_insert_date) <> 2024
    ORDER BY d.sth_id DESC, d.std_id ASC;
    """

    # 3. Payment Methods Breakdown
    q_payments = """
    SELECT 
        stp.stp_sth_id AS Bill_ID,
        st.sto_name_ar AS Store_Name,
        fcs.fcs_name_ar AS Cash_Drawer_Name,
        stp.stp_value AS Tendered_Value,
        stp.stp_insert_date AS Payment_Timestamp,
        COALESCE(emp.e_Name, CAST(stp.stp_insert_uid AS VARCHAR(50))) AS Processed_By,
        stp.stp_pc_name AS Station
    FROM sales_trans_payment stp WITH (NOLOCK)
    JOIN Store st WITH (NOLOCK) ON st.sto_id = stp.stp_sto_id
    JOIN F_Cash_Store fcs WITH (NOLOCK) ON fcs.fcs_id = stp.stp_fcs_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(stp.stp_insert_uid AS DECIMAL(18,0))
    WHERE YEAR(stp.stp_insert_date) <> 2024
    ORDER BY stp.stp_sth_id DESC;
    """

    # 4. Sales Returns Operations
    q_returns = """
    SELECT 
        sr.sr_id AS Return_Doc_ID,
        sr.sth_id AS Original_Bill_ID,
        h.sec_insert_date AS Original_Sale_Date,
        sr.returned_items_no AS Items_Count_Returned,
        sr.returned_items_value AS Gross_Refund_Value,
        sr.total_back_tax AS Tax_Refunded,
        srp.srp_value AS Actual_Cash_Refunded,
        fcs.fcs_name_ar AS Refunded_From_Drawer,
        sr.sec_insert_date AS Return_Timestamp,
        COALESCE(emp.e_Name, CAST(sr.sec_insert_uid AS VARCHAR(50))) AS Handled_By
    FROM sales_return sr WITH (NOLOCK)
    JOIN sales_trans_h h WITH (NOLOCK) ON h.sth_id = sr.sth_id
    LEFT JOIN sales_return_payment srp WITH (NOLOCK) ON srp.srp_sr_id = sr.sr_id
    LEFT JOIN F_Cash_Store fcs WITH (NOLOCK) ON fcs.fcs_id = srp.srp_fcs_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(sr.sec_insert_uid AS DECIMAL(18,0))
    WHERE YEAR(sr.sec_insert_date) <> 2024
    ORDER BY sr.sr_id DESC;
    """

    # 5. Electronic Card Settlements
    q_visa = """
    SELECT 
        svi.sth_id AS Bill_ID,
        b.bnk_name_ar AS Terminal_Bank,
        svi.visa_name AS Cardholder_Name,
        svi.visa_no AS Card_Masked_Number,
        svi.exp_date AS Card_Expiry,
        svi.visa_value AS Transaction_Amount,
        svi.get_fees_perc AS Bank_Fee_Percentage,
        svi.sec_insert_date AS Transaction_Date
    FROM sales_visa_info svi WITH (NOLOCK)
    LEFT JOIN Bank b WITH (NOLOCK) ON b.bnk_id = svi.visa_bank_id
    WHERE YEAR(svi.sec_insert_date) <> 2024
    ORDER BY svi.sth_id DESC;
    """

    # 6. Deleted Invoices Audit Log
    q_del_bills = """
    SELECT 
        del.sth_id AS Voided_Bill_ID,
        st.sto_name_ar AS Store_Name,
        del.total_bill_net AS Voided_Amount,
        del.no_of_items AS Voided_Lines,
        del.sec_insert_date AS Creation_Time,
        del.sth_delTime AS Deletion_Time,
        del.sth_del_pc_name AS Deletion_Station,
        COALESCE(emp.e_Name, CAST(del.sec_insert_uid AS VARCHAR(50))) AS Original_Cashier
    FROM sales_trans_h_del del WITH (NOLOCK)
    JOIN Store st WITH (NOLOCK) ON st.sto_id = del.sto_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(del.sec_insert_uid AS DECIMAL(18,0))
    WHERE YEAR(del.sth_delTime) <> 2024
    ORDER BY del.sth_delTime DESC;
    """

    # 7. Cashier Voided Prescription Items (Loss Prevention)
    q_deleted_items = """
    SELECT 
        del.sth_id AS Voided_Bill_ID,
        del.sec_insert_date AS Original_Creation_Time,
        del.std_delTime AS Deletion_Timestamp,
        del.std_del_pc_name AS Station_Hostname,
        del.itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        ic.itm_name_ar AS Item_Name,
        del.qnty AS Voided_Quantity,
        del.itm_sell AS Unit_Sale_Price,
        del.itm_cost AS Unit_Cost,
        CAST(del.qnty * del.itm_sell AS DECIMAL(18,2)) AS Total_Voided_Value,
        COALESCE(emp.e_Name, CAST(del.sec_insert_uid AS VARCHAR(50))) AS Voided_By_Cashier
    FROM sales_trans_d_del del WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = del.itm_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(del.sec_insert_uid AS DECIMAL(18,0))
    WHERE YEAR(del.std_delTime) <> 2024
    ORDER BY del.std_delTime DESC;
    """

    # 8. Post-Ringing Quantity Modifications Audit
    q_qty_modifications = """
    SELECT 
        m.sth_id AS Bill_ID,
        m.itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        ic.itm_name_ar AS Item_Name,
        m.qnty AS Original_Quantity,
        m.itm_back AS Return_Quantity,
        m.new AS Modified_Quantity,
        m.itm_sell AS Original_Selling_Price,
        m.newprice AS Modified_Selling_Price,
        m.itm_unit AS Unit_Level
    FROM sales_trans_d_modified_qnty m WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = m.itm_id
    JOIN sales_trans_h h WITH (NOLOCK) ON h.sth_id = m.sth_id
    WHERE YEAR(h.sec_insert_date) <> 2024
    ORDER BY m.sth_id DESC;
    """

    # 9. Insurance Syndicate Claims
    q_insurance_claims = """
    SELECT 
        h.sth_id AS Bill_ID,
        st.sto_name_ar AS Store_Name,
        h.sec_insert_date AS Prescription_Date,
        c.cust_name_ar AS Patient_Name,
        cc.cc_name AS Insurance_Syndicate,
        cc.cc_number AS Policy_Contract_No,
        h.fh_medins_rec_name AS Beneficiary_Member_Name,
        h.fh_medins_ticket_num AS Claim_Ticket_No,
        h.fh_medins_ins_num AS Insurance_Card_No,
        h.total_bill_after_disc AS Total_Prescription_Value,
        cc.cc_pay_perc AS Patient_Copay_Perc,
        CAST(h.total_bill_after_disc - h.fh_company_part AS DECIMAL(18,2)) AS Patient_Copay_Paid,
        h.fh_company_part AS Syndicate_Claim_Liability,
        COALESCE(emp.e_Name, CAST(h.emp_id AS VARCHAR(50))) AS Dispensing_Pharmacist
    FROM sales_trans_h h WITH (NOLOCK)
    JOIN Customer_Contracts cc WITH (NOLOCK) ON cc.cc_id = h.fh_contract_id
    JOIN Store st WITH (NOLOCK) ON st.sto_id = h.sto_id
    LEFT JOIN Customer c WITH (NOLOCK) ON c.cust_id = h.cust_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(h.emp_id AS DECIMAL(18,0))
    WHERE h.sth_flag = 'C' AND h.fh_contract_id > 0 AND YEAR(h.sec_insert_date) <> 2024
    ORDER BY h.sth_id DESC;
    """

    # 10. Doctor & Clinic Referral Analytics
    q_doctor_referrals = """
    SELECT 
        h.sth_id AS Bill_ID,
        h.sec_insert_date AS Prescription_Date,
        ISNULL(h.fh_medins_Doc_name, 'Direct Walk-in / Self-Medication') AS Referring_Doctor,
        ds.ds_name_ar AS Doctor_Specialization,
        cs.cs_name_ar AS Clinic_Specialization,
        c.cust_name_ar AS Patient_Name,
        h.no_of_items AS Prescribed_Items_Count,
        h.total_bill_after_disc AS Prescription_Total_Value,
        COALESCE(emp.e_Name, CAST(h.emp_id AS VARCHAR(50))) AS Dispensing_Pharmacist
    FROM sales_trans_h h WITH (NOLOCK)
    LEFT JOIN Customer c WITH (NOLOCK) ON c.cust_id = h.cust_id
    LEFT JOIN Doctors_specializations ds WITH (NOLOCK) ON ds.ds_id = h.fh_doc_spec_id
    LEFT JOIN Clinic_specializations cs WITH (NOLOCK) ON cs.cs_id = h.fh_clinic_spec_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(h.emp_id AS DECIMAL(18,0))
    WHERE h.sth_flag = 'C' AND (h.fh_doc_spec_id > 0 OR h.fh_medins_Doc_name IS NOT NULL)
      AND YEAR(h.sec_insert_date) <> 2024
    ORDER BY h.sth_id DESC;
    """

    # 11. Delivery Dispatch & Courier Logistics
    q_delivery_dispatch = """
    SELECT 
        h.sth_id AS Bill_ID,
        st.sto_name_ar AS Store_Name,
        h.sec_insert_date AS Order_Time,
        ISNULL(c.cust_name_ar, 'Walk-in / Phone Client') AS Customer_Name,
        c.cust_mobile AS Mobile_Phone,
        ISNULL(sdi.address, cd.cd_address) AS Delivery_Address,
        dcr.dc_region_name_ar AS Delivery_Zone,
        h.total_bill_after_disc AS Order_Value,
        h.sth_extra_expenses AS Delivery_Fee_Charged,
        deliv_emp.e_Name AS Courier_Assigned,
        CASE h.sth_deliv_closed WHEN 1 THEN 'Delivered & Closed' ELSE 'Pending Dispatch' END AS Dispatch_Status,
        COALESCE(emp.e_Name, CAST(h.emp_id AS VARCHAR(50))) AS Cashier_Name
    FROM sales_trans_h h WITH (NOLOCK)
    JOIN Store st WITH (NOLOCK) ON st.sto_id = h.sto_id
    LEFT JOIN Customer c WITH (NOLOCK) ON c.cust_id = h.cust_id
    LEFT JOIN Customer_Delivery cd WITH (NOLOCK) ON cd.cd_cust_id = c.cust_id
    LEFT JOIN Delivery_cost_region dcr WITH (NOLOCK) ON dcr.dc_id = cd.cd_region_id
    LEFT JOIN sales_deliv_info sdi WITH (NOLOCK) ON sdi.sth_id = h.sth_id
    LEFT JOIN Employee deliv_emp WITH (NOLOCK) ON deliv_emp.e_id = sdi.deliv_emp_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(h.emp_id AS DECIMAL(18,0))
    WHERE h.sth_flag = 'C' AND (h.bill_typ = 3 OR h.sth_extra_expenses > 0 OR sdi.sdi_id IS NOT NULL)
      AND YEAR(h.sec_insert_date) <> 2024
    ORDER BY h.sth_id DESC;
    """

    # 12. Sold Without Stock Negative Stock Overrides
    q_sold_without_stock = """
    SELECT 
        sne.sth_id AS Bill_ID,
        sne.std_id AS Line_No,
        sne.sec_insert_date AS Sale_Timestamp,
        sne.itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        ic.itm_name_ar AS Item_Name,
        sne.qnty AS Overdrawn_Sold_Qty,
        sne.itm_sell AS Selling_Price,
        sne.itm_cost AS Standard_Cost,
        CAST(sne.qnty * sne.itm_sell AS DECIMAL(18,2)) AS Revenue_Realized,
        sne.ClassStockAfter AS Negative_Stock_Position_After_Sale,
        COALESCE(emp.e_Name, CAST(sne.sec_insert_uid AS VARCHAR(50))) AS Cashier_Name
    FROM sales_not_exist_modified sne WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = sne.itm_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(sne.sec_insert_uid AS DECIMAL(18,0))
    WHERE YEAR(sne.sec_insert_date) <> 2024
    ORDER BY sne.sth_id DESC;
    """

    # 13. Customer Loyalty & Rewards Registry
    q_loyalty_rewards = """
    SELECT 
        c.cust_id AS Customer_ID,
        c.cust_code AS Customer_Code,
        c.cust_name_ar AS Customer_Name,
        c.cust_mobile AS Mobile_Number,
        c.cust_curr_points AS Accumulated_Loyalty_Points,
        c.pont2mony AS Points_Cash_Conversion_Rate,
        CAST(c.cust_curr_points * c.pont2mony AS DECIMAL(18,2)) AS Available_Reward_Value_L_E,
        ISNULL((
            SELECT CAST(SUM(sc_used_value) AS DECIMAL(18,2)) 
            FROM sales_cobons sc WITH (NOLOCK) 
            JOIN sales_trans_h sh WITH (NOLOCK) ON sh.sth_id = sc.sth_id 
            WHERE sh.cust_id = c.cust_id AND YEAR(sh.sec_insert_date) <> 2024
        ), 0.00) AS Total_Points_Redeemed_Historical,
        c.sec_update_date AS Last_Loyalty_Activity
    FROM Customer c WITH (NOLOCK)
    WHERE c.cust_active = '1' AND (c.cust_curr_points > 0 OR c.cust_discount_perc > 0)
    ORDER BY c.cust_curr_points DESC;
    """

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        read_sql_and_enrich(q_sales_h, ENGINE).to_excel(
            writer, sheet_name="Sales_Invoices", index=False
        )
        read_sql_and_enrich(q_sales_d, ENGINE).to_excel(
            writer, sheet_name="Sales_Line_Items", index=False
        )
        read_sql_and_enrich(q_payments, ENGINE).to_excel(
            writer, sheet_name="Drawer_Payments", index=False
        )
        read_sql_and_enrich(q_returns, ENGINE).to_excel(
            writer, sheet_name="Sales_Returns", index=False
        )
        read_sql_and_enrich(q_visa, ENGINE).to_excel(
            writer, sheet_name="Card_Settlements", index=False
        )
        read_sql_and_enrich(q_del_bills, ENGINE).to_excel(
            writer, sheet_name="Deleted_Bills_Audit", index=False
        )
        read_sql_and_enrich(q_deleted_items, ENGINE).to_excel(
            writer, sheet_name="Cashier_Voided_Lines_Audit", index=False
        )
        read_sql_and_enrich(q_qty_modifications, ENGINE).to_excel(
            writer, sheet_name="Sales_Quantity_Modifications", index=False
        )
        read_sql_and_enrich(q_insurance_claims, ENGINE).to_excel(
            writer, sheet_name="Insurance_Contract_Claims", index=False
        )
        read_sql_and_enrich(q_doctor_referrals, ENGINE).to_excel(
            writer, sheet_name="Doctor_Prescription_Analytics", index=False
        )
        read_sql_and_enrich(q_delivery_dispatch, ENGINE).to_excel(
            writer, sheet_name="Home_Delivery_Dispatch", index=False
        )
        read_sql_and_enrich(q_sold_without_stock, ENGINE).to_excel(
            writer, sheet_name="Sold_Without_Stock_Audit", index=False
        )
        read_sql_and_enrich(q_loyalty_rewards, ENGINE).to_excel(
            writer, sheet_name="Customer_Loyalty_Rewards", index=False
        )

    wb = openpyxl.load_workbook(filename)
    table_names = set()
    for ws_name in wb.sheetnames:
        style_worksheet(
            wb[ws_name],
            header_fill_hex="1F4E79",
            is_kpi=False,
            existing_table_names=table_names,
        )
    wb.save(filename)
    logger.info("Module 2 Complete.")


# ==================================================================================================
# MODULE 3: PROCUREMENT & VENDOR SUPPLY CHAIN
# ==================================================================================================
def generate_purchases_workbook(output_dir):
    filename = os.path.join(output_dir, "03_Genius_Purchasing_Operations.xlsx")
    logger.info(f"Generating Module 3: {filename}...")

    # 1. Purchase Invoices Header
    q_pur_h = """
    SELECT 
        pth.pth_id AS Purchase_Bill_ID,
        st.sto_name_ar AS Receiving_Store,
        v.ven_name_ar AS Vendor_Name,
        pth.ven_bill_no AS Vendor_Invoice_Ref,
        pth.ven_bill_date AS Vendor_Bill_Date,
        pth.no_of_items AS Total_Lines,
        pth.total_bill AS Gross_Purchase_Total,
        pth.total_dis_per AS Header_Discount_Perc,
        pth.total_des_mon AS Header_Discount_Cash,
        pth.total_after_back AS Net_Purchase_Total,
        pth.pth_paid AS Paid_Amount,
        (pth.total_after_back - pth.pth_paid) AS Remaining_Vendor_Liability,
        pth.pth_source_tax AS Withholding_Tax_Rate,
        COALESCE(emp.e_Name, CAST(pth.sec_insert_uid AS VARCHAR(50))) AS Created_By,
        pth.sec_insert_date AS System_Entry_Date
    FROM pur_trans_h pth WITH (NOLOCK)
    JOIN Store st WITH (NOLOCK) ON st.sto_id = pth.sto_id
    JOIN Vendor v WITH (NOLOCK) ON v.ven_id = pth.ven_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(pth.sec_insert_uid AS DECIMAL(18,0))
    WHERE pth.bill_status <> 1 AND YEAR(pth.sec_insert_date) <> 2024
    ORDER BY pth.pth_id DESC;
    """

    # 2. Purchase Invoices Details (With Lot Batch)
    q_pur_d = """
    SELECT 
        ptd.pth_id AS Purchase_Bill_ID,
        ptd.ptd_id AS Line_No,
        pth.sec_insert_date AS Purchase_Date,
        v.ven_name_ar AS Vendor_Name,
        ptd.itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        (SELECT TOP 1 io_itm_int FROM Item_Objects WITH (NOLOCK) WHERE io_itm_id = ptd.itm_id) AS International_Barcode,
        ic.itm_name_ar AS Item_Name,
        ptd.c_id AS Assigned_Batch_ID,
        ptd.ptd_batch AS Vendor_Lot_Batch,
        CONVERT(VARCHAR(10), ptd.exp_date, 120) AS Expiry_Date,
        ptd.qnty AS Purchased_Qty,
        ptd.bonus AS Bonus_Free_Qty,
        ptd.itm_pur_price AS Unit_Wholesale_Price,
        ptd.itm_extra_dis AS Line_Discount_Perc,
        ptd.itm_tax_price AS Unit_Tax_Amount,
        ptd.itm_cost AS Effective_Unit_Cost,
        ptd.itm_sell AS Expected_Retail_Price,
        CAST(
            (ptd.qnty * (ptd.itm_pur_price * (1.0 - (ISNULL(ptd.itm_extra_dis,0)/100.0)) + ptd.itm_tax_price))
            * (1.0 - (ISNULL(pth.total_dis_per,0)/100.0))
        AS DECIMAL(18,2)) AS Net_Line_Purchase_Cost,
        ptd.itm_back_qty AS Returned_Qty
    FROM pur_trans_d ptd WITH (NOLOCK)
    JOIN pur_trans_h pth WITH (NOLOCK) ON pth.pth_id = ptd.pth_id
    JOIN Vendor v WITH (NOLOCK) ON v.ven_id = pth.ven_id
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = ptd.itm_id
    WHERE pth.bill_status <> 1 AND YEAR(pth.sec_insert_date) <> 2024
    ORDER BY ptd.pth_id DESC, ptd.ptd_id ASC;
    """

    # 3. Vendor General Returns
    q_pur_returns = """
    SELECT 
        ibh.ibh_id AS Return_Doc_ID,
        st.sto_name_ar AS Store_Name,
        v.ven_name_ar AS Vendor_Name,
        ibh.ibh_total_value AS Total_Return_Value,
        ibd.ibd_itm_id AS Item_ID,
        ic.itm_name_ar AS Item_Name,
        ibd.ibd_qty AS Returned_Quantity,
        ibd.ibd_pharm_price AS Refund_Unit_Cost,
        ibd.ibd_tax_price AS Refund_Unit_Tax,
        CAST(ibd.ibd_qty * (ibd.ibd_pharm_price + ibd.ibd_tax_price) AS DECIMAL(18,2)) AS Total_Refund_Credit,
        CONVERT(VARCHAR(10), ibd.ibd_exp_date, 120) AS Expiry_Date,
        ibh.sec_insert_date AS Return_Date,
        COALESCE(emp.e_Name, CAST(ibh.sec_insert_uid AS VARCHAR(50))) AS Returned_By
    FROM Item_back_h ibh WITH (NOLOCK)
    JOIN Item_back_d ibd WITH (NOLOCK) ON ibd.ibd_ibh_id = ibh.ibh_id
    JOIN Vendor v WITH (NOLOCK) ON v.ven_id = ibh.ibh_ven_id
    JOIN Store st WITH (NOLOCK) ON st.sto_id = ibh.ibh_sto_id
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = ibd.ibd_itm_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(ibh.sec_insert_uid AS DECIMAL(18,0))
    WHERE YEAR(ibh.sec_insert_date) <> 2024
    ORDER BY ibh.ibh_id DESC;
    """

    # 4. Direct Bill-Level Purchase Returns
    q_direct_bill_returns = """
    SELECT 
        ptd.pth_id AS Purchase_Bill_ID,
        ptd.ptd_id AS Line_No,
        pth.ven_bill_no AS Vendor_Invoice_Ref,
        v.ven_name_ar AS Vendor_Name,
        ptd.itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        ic.itm_name_ar AS Item_Name,
        ptd.itm_back_qty AS Returned_Quantity,
        ptd.itm_back_pharm AS Unit_Return_Price,
        ptd.itm_back_tax AS Unit_Return_Tax,
        ptd.itm_back_bonus AS Returned_Bonus_Units,
        CAST((ptd.itm_back_qty * (ptd.itm_back_pharm + ptd.itm_back_tax)) AS DECIMAL(18,2)) AS Total_Invoice_Credit_Claimed,
        ptd.sec_update_date AS Return_Action_Date,
        COALESCE(emp.e_Name, CAST(ptd.sec_update_uid AS VARCHAR(50))) AS Handled_By
    FROM pur_trans_d ptd WITH (NOLOCK)
    JOIN pur_trans_h pth WITH (NOLOCK) ON pth.pth_id = ptd.pth_id
    JOIN Vendor v WITH (NOLOCK) ON v.ven_id = pth.ven_id
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = ptd.itm_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(ptd.sec_update_uid AS DECIMAL(18,0))
    WHERE ptd.itm_back > 0 AND pth.bill_status <> 1 AND YEAR(pth.sec_insert_date) <> 2024
    ORDER BY ptd.pth_id DESC;
    """

    # 5. Vendor Pricing History & Quotations
    q_vendor_pricing = """
    SELECT 
        ivh.itm_id AS Item_ID,
        ic.itm_name_ar AS Item_Name,
        v.ven_name_ar AS Vendor_Name,
        ivh.itm_pharm_price AS Quoted_Pharm_Price,
        ivh.itm_sell_price AS Retail_Price,
        ivh.itm_disc1 AS Primary_Discount,
        ivh.itm_disc2 AS Secondary_Discount,
        ivh.itm_disc3 AS Special_Discount,
        ivh.itm_pricing_date AS Price_Recorded_Date
    FROM Item_Vendor_history ivh WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = ivh.itm_id
    JOIN Vendor v WITH (NOLOCK) ON v.ven_id = ivh.ven_id
    WHERE YEAR(ivh.itm_pricing_date) <> 2024
    ORDER BY ivh.itm_pricing_date DESC;
    """

    # 6. Vendor Accounts Payable & Advance Summary Ledger (Strict Signs & Absolute Ordering)
    q_vendor_ledger = """
    SELECT 
        v.ven_id AS Vendor_ID,
        v.ven_code AS Vendor_Code,
        v.ven_name_ar AS Vendor_Name_Ar,
        v.ven_tel AS Phone_Contact,
        v.ven_max_credit AS Credit_Limit,
        v.ven_start_debit AS Initial_Debit,
        v.ven_start_credit AS Initial_Credit,
        v.ven_current_credit AS Current_Outstanding_Balance,
        CASE 
            WHEN v.ven_current_credit > 0 THEN 'We Owe Vendor - Accounts Payable (دائن - التزام علينا)' 
            WHEN v.ven_current_credit < 0 THEN 'Vendor Owes Pharmacy - Advance/Return (مدين - رصيد لنا)' 
            ELSE 'Settled (خالص)' 
        END AS Balance_Status,
        v.sec_update_date AS Last_Transaction_Date
    FROM Vendor v WITH (NOLOCK)
    WHERE v.ven_active = '1'
    ORDER BY ABS(v.ven_current_credit) DESC;
    """

    # 7. Vendor Credit Change Forensic Log
    q_vendor_credit_audit = """
    SELECT 
        vc.vc_id AS Audit_ID,
        v.ven_name_ar AS Vendor_Name,
        vc.vc_old_credit AS Previous_Balance,
        vc.vc_new_credit AS Updated_Balance,
        CAST(vc.vc_new_credit - vc.vc_old_credit AS DECIMAL(18,2)) AS Balance_Delta,
        vc.vc_operation AS Triggering_Operation,
        vc.vc_prog_name AS Calling_Application,
        vc.vc_pc_name AS Station_Hostname,
        vc.vc_insert_date AS Audit_Timestamp
    FROM vendor_credit_chng vc WITH (NOLOCK)
    JOIN Vendor v WITH (NOLOCK) ON v.ven_id = vc.vc_cust_id
    WHERE YEAR(vc.vc_insert_date) <> 2024
    ORDER BY vc.vc_id DESC;
    """

    # 8. Wholesaler Discount Arbitrage Matrix
    q_distributor_arbitrage = """
    SELECT 
        iv.itm_id AS Item_ID,
        ic.itm_code AS Internal_Code,
        ic.itm_name_ar AS Item_Name,
        v.ven_name_ar AS Distributor_Name,
        iv.itm_pharm_price AS Quoted_Pharm_Price,
        iv.itm_disc1 AS Base_Discount_Perc,
        iv.itm_disc2 AS Extra_Discount_Perc,
        iv.itm_disc3 AS Special_Discount_Perc,
        CAST(iv.itm_disc1 + iv.itm_disc2 + iv.itm_disc3 AS DECIMAL(18,2)) AS Cumulative_Discount_Rate,
        CAST(iv.itm_pharm_price * (1.0 - ((iv.itm_disc1 + iv.itm_disc2 + iv.itm_disc3) / 100.0)) AS DECIMAL(18,2)) AS Net_Procurement_Unit_Cost,
        iv.itm_sell_price AS Retail_Selling_Price,
        CAST(iv.itm_sell_price - (iv.itm_pharm_price * (1.0 - ((iv.itm_disc1 + iv.itm_disc2 + iv.itm_disc3) / 100.0))) AS DECIMAL(18,2)) AS Profit_Spread,
        CASE WHEN iv.avail_flag = '1' THEN 'Available' ELSE 'Check Availability' END AS Distributor_Availability,
        iv.last_price_update AS Quote_Date
    FROM Item_Vendor iv WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = iv.itm_id
    JOIN Vendor v WITH (NOLOCK) ON v.ven_id = iv.ven_id
    WHERE v.ven_active = '1' AND (iv.itm_disc1 + iv.itm_disc2 + iv.itm_disc3) > 0
    ORDER BY Cumulative_Discount_Rate DESC;
    """

    # 9. Master Vendor Directory (Vendor + Vendor_Extn)
    q_vendor_directory = """
    SELECT 
        v.ven_id AS Vendor_ID,
        v.ven_code AS Vendor_Code,
        v.ven_name_ar AS Vendor_Name_Ar,
        v.ven_name_en AS Vendor_Name_En,
        v.ven_tel AS Primary_Telephone,
        ve.ve_tel AS Secondary_Telephone,
        ve.ve_contact_name AS Contact_Person,
        v.ven_mgr_name AS Manager_Name,
        v.ven_address AS Registered_Address,
        v.ven_max_credit AS Authorized_Credit_Limit,
        v.ven_b2b_name AS B2B_Platform_Name,
        v.ven_pharm_code AS Pharmacy_Code_At_Vendor,
        CASE v.ven_active WHEN '1' THEN 'Active' ELSE 'Inactive' END AS Account_Status
    FROM Vendor v WITH (NOLOCK)
    LEFT JOIN Vendor_Extn ve WITH (NOLOCK) ON ve.ve_ven_id = v.ven_id
    ORDER BY v.ven_id;
    """

    # 10. Official UCP (United Company for Pharmacists) Wholesaler Catalog
    q_ucp_catalog = """
    SELECT 
        vc.vc_code AS Distributor_Item_Code,
        vc.vc_company AS Wholesaler_Name,
        vc.vc_item_name AS Catalog_Trade_Name,
        vc.vc_pharm_price AS Wholesale_Price,
        vc.vc_sell_price AS Official_Retail_Price,
        vc.vc_tax AS Tax_Included,
        vc.sec_insert_date AS Catalog_Date
    FROM Vendors_Catalogs vc WITH (NOLOCK)
    ORDER BY vc.vc_company, vc.vc_item_name;
    """

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        read_sql_and_enrich(q_pur_h, ENGINE).to_excel(
            writer, sheet_name="Purchase_Bills", index=False
        )
        read_sql_and_enrich(q_pur_d, ENGINE).to_excel(
            writer, sheet_name="Purchase_Line_Items", index=False
        )
        read_sql_and_enrich(q_pur_returns, ENGINE).to_excel(
            writer, sheet_name="Vendor_Returns", index=False
        )
        read_sql_and_enrich(q_direct_bill_returns, ENGINE).to_excel(
            writer, sheet_name="Direct_Bill_Purchase_Returns", index=False
        )
        read_sql_and_enrich(q_vendor_pricing, ENGINE).to_excel(
            writer, sheet_name="Vendor_Pricing_History", index=False
        )
        read_sql_and_enrich(q_vendor_ledger, ENGINE).to_excel(
            writer, sheet_name="Vendor_Ledger_Balances", index=False
        )
        read_sql_and_enrich(q_vendor_credit_audit, ENGINE).to_excel(
            writer, sheet_name="Vendor_Credit_Audit_Log", index=False
        )
        read_sql_and_enrich(q_distributor_arbitrage, ENGINE).to_excel(
            writer, sheet_name="Distributor_Price_Arbitrage", index=False
        )
        read_sql_and_enrich(q_vendor_directory, ENGINE).to_excel(
            writer, sheet_name="Master_Vendor_Directory", index=False
        )
        read_sql_and_enrich(q_ucp_catalog, ENGINE).to_excel(
            writer, sheet_name="Wholesaler_Master_Catalog", index=False
        )

    wb = openpyxl.load_workbook(filename)
    table_names = set()
    for ws_name in wb.sheetnames:
        style_worksheet(
            wb[ws_name],
            header_fill_hex="4A154B",
            is_kpi=False,
            existing_table_names=table_names,
        )
    wb.save(filename)
    logger.info("Module 3 Complete.")


# ==================================================================================================
# MODULE 4: TREASURY & GENERAL FINANCIAL LEDGER
# ==================================================================================================
def generate_financial_workbook(output_dir):
    filename = os.path.join(output_dir, "04_Genius_Financial_Ledger.xlsx")
    logger.info(f"Generating Module 4: {filename}...")

    # 1. Cash Safes, Drawers & Banks
    q_cash_stores = """
    SELECT 
        fcs.fcs_id AS Fund_ID,
        fcs.fcs_code AS Code,
        fcs.fcs_name_ar AS Account_Title,
        CASE fcs.fcs_type 
            WHEN '1' THEN 'POS Cashier Drawer (درج كاشير)' 
            WHEN '2' THEN 'Main Treasury/Vault (خزينة رئيسية)' 
            WHEN '3' THEN 'Bank Account (حساب بنكي)' 
            ELSE 'Other Fund' 
        END AS Account_Type,
        fcs.fcs_current_balance AS Current_Live_Balance,
        fcs.fcs_account_number AS Bank_Account_No,
        b.bnk_name_ar AS Associated_Bank,
        fcs.fcs_computer_name AS Bound_Workstation,
        fcs.sec_update_date AS Last_Balance_Update
    FROM F_Cash_Store fcs WITH (NOLOCK)
    LEFT JOIN Bank b WITH (NOLOCK) ON b.bnk_id = fcs.fcs_bank_id
    WHERE fcs.fcs_active = 1
    ORDER BY fcs.fcs_type, fcs.fcs_id;
    """

    # 2. Daily Treasury Transactions (Full Entity Name Decoding)
    q_treasury_trans = """
    SELECT 
        fh.fh_id AS Entry_ID,
        fh.fh_actual_date AS Transaction_Date,
        CASE fh.fh_trans_type 
            WHEN '1' THEN 'Cash Movement (نقدية)' 
            WHEN '2' THEN 'Cheque/Commercial Paper (أوراق مالية)' 
            ELSE 'Other' 
        END AS Instrument_Type,
        fh.fh_value AS Amount,
        CASE fh.fh_from_type 
            WHEN '1' THEN 'Customer (عميل)' 
            WHEN '2' THEN 'Vendor (مورد)' 
            WHEN '3' THEN 'Cash Safe/Drawer (خزينة/درج)' 
            ELSE 'Account/Other' 
        END AS Source_Entity_Type,
        CASE 
            WHEN fh.fh_from_type = '1' THEN ISNULL(c_from.cust_name_ar, 'Cash Customer')
            WHEN fh.fh_from_type = '2' THEN v_from.ven_name_ar
            WHEN fh.fh_from_type = '3' THEN fcs_from.fcs_name_ar
            ELSE 'General GL Account'
        END AS Source_Entity_Name,
        CASE fh.fh_to_type 
            WHEN '1' THEN 'Customer (عميل)' 
            WHEN '2' THEN 'Vendor (مورد)' 
            WHEN '3' THEN 'Cash Safe/Drawer (خزينة/درج)' 
            ELSE 'Account/Other' 
        END AS Destination_Entity_Type,
        CASE 
            WHEN fh.fh_to_type = '1' THEN ISNULL(c_to.cust_name_ar, 'Cash Customer')
            WHEN fh.fh_to_type = '2' THEN v_to.ven_name_ar
            WHEN fh.fh_to_type = '3' THEN fcs_to.fcs_name_ar
            ELSE 'General GL Account'
        END AS Destination_Entity_Name,
        fh.fh_notes AS Operational_Narrative,
        fh.fh_computer AS Executing_Machine,
        COALESCE(emp.e_Name, CAST(fh.sec_insert_uid AS VARCHAR(50))) AS Operator_Name
    FROM F_Transaction_Header fh WITH (NOLOCK)
    LEFT JOIN Customer c_from WITH (NOLOCK) ON c_from.cust_id = fh.fh_from_id AND fh.fh_from_type = '1'
    LEFT JOIN Vendor v_from WITH (NOLOCK) ON v_from.ven_id = fh.fh_from_id AND fh.fh_from_type = '2'
    LEFT JOIN F_Cash_Store fcs_from WITH (NOLOCK) ON fcs_from.fcs_id = fh.fh_from_id AND fh.fh_from_type = '3'
    LEFT JOIN Customer c_to WITH (NOLOCK) ON c_to.cust_id = fh.fh_to_id AND fh.fh_to_type = '1'
    LEFT JOIN Vendor v_to WITH (NOLOCK) ON v_to.ven_id = fh.fh_to_id AND fh.fh_to_type = '2'
    LEFT JOIN F_Cash_Store fcs_to WITH (NOLOCK) ON fcs_to.fcs_id = fh.fh_to_id AND fh.fh_to_type = '3'
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(fh.sec_insert_uid AS DECIMAL(18,0))
    WHERE YEAR(fh.fh_actual_date) <> 2024
    ORDER BY fh.fh_id DESC;
    """

    # 3. Commercial Papers & Cheques Portfolio
    q_cheques = """
    SELECT 
        fhe.fhe_id AS Paper_Registry_ID,
        fhe.fh_id AS Treasury_Header_ID,
        fp.fp_name_ar AS Paper_Type,
        fhe.fhe_paper_number AS Cheque_Paper_Number,
        fhe.fhe_paper_create_date AS Issue_Date,
        fhe.fhe_paper_valid_date AS Due_Maturity_Date,
        b.bnk_name_ar AS Deposit_Bank,
        CASE fhe.fhe_paper_status 
            WHEN '0' THEN 'Pending In-Portfolio (تحت التحصيل)' 
            WHEN '1' THEN 'Deposited/Collected (تم التحصيل)' 
            WHEN '2' THEN 'Rejected/Bounced (مرفوض/مرتد)' 
            WHEN '3' THEN 'Cancelled (ملغي)' 
            ELSE 'Unknown' 
        END AS Cheque_Status,
        fh.fh_value AS Face_Value,
        fhe.fhe_expenses AS Bank_Clearance_Expenses
    FROM F_Transaction_Header_Ext fhe WITH (NOLOCK)
    JOIN F_Transaction_Header fh WITH (NOLOCK) ON fh.fh_id = fhe.fh_id
    LEFT JOIN F_Papers fp WITH (NOLOCK) ON fp.fp_id = fhe.fhe_paper_type
    LEFT JOIN Bank b WITH (NOLOCK) ON b.bnk_id = fhe.fhe_paper_deposit_bank_id
    WHERE (fhe.fhe_paper_create_date IS NULL OR YEAR(fhe.fhe_paper_create_date) <> 2024)
    ORDER BY fhe.fhe_id DESC;
    """

    # 4. Automated GL Document Engine
    q_gl_engine = """
    SELECT 
        fah.fah_id AS Auto_Doc_ID,
        st.sto_name_ar AS Branch_Name,
        fah.fah_type AS System_Trigger,
        fah.fah_form_text AS Originating_Form,
        fah.fah_description AS Journal_Narrative,
        fah.fah_total_value AS Total_Document_Value,
        fad.fad_group_id AS Leg_Group,
        fad.fad_D_Account_name AS Debit_Account,
        fad.fad_D_object AS Debit_Target_Type,
        fad.fad_C_Account_name AS Credit_Account,
        fad.fad_C_object AS Credit_Target_Type,
        fad.fad_value AS Leg_Amount,
        fah.sec_insert_date AS Posting_Timestamp,
        fah.fah_computer_name AS Machine
    FROM F_Auto_Doc_h fah WITH (NOLOCK)
    JOIN F_Auto_Doc_d fad WITH (NOLOCK) ON fad.fad_fah_id = fah.fah_id
    JOIN Store st WITH (NOLOCK) ON st.sto_id = fah.fah_sto_id
    WHERE YEAR(fah.sec_insert_date) <> 2024
    ORDER BY fah.fah_id DESC, fad.fad_group_id ASC;
    """

    # 5. Debit / Credit Adjustment Notices
    q_notices = """
    SELECT 
        fn.fn_id AS Notice_ID,
        CASE fn.fn_type 
            WHEN '1' THEN 'Addition Notice - Credit (إشعار إضافة)' 
            WHEN '2' THEN 'Discount Notice - Debit (إشعار خصم)' 
            ELSE 'Other' 
        END AS Notice_Type,
        CASE fn.fn_who_type 
            WHEN '1' THEN 'Vendor (مورد)' 
            WHEN '2' THEN 'Customer (عميل)' 
            ELSE 'Party' 
        END AS Party_Type,
        fn.fn_who_id AS Party_ID,
        fnc.fnc_name_ar AS Reason_Category,
        fn.fn_value AS Notice_Amount,
        fn.fn_notes AS Explanation,
        fn.sec_insert_date AS Issued_Timestamp,
        COALESCE(emp.e_Name, CAST(fn.sec_insert_uid AS VARCHAR(50))) AS Officer_Name
    FROM F_Notices fn WITH (NOLOCK)
    LEFT JOIN F_Notices_Causes fnc WITH (NOLOCK) ON fnc.fnc_id = fn.fn_cause_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(fn.sec_insert_uid AS DECIMAL(18,0))
    WHERE YEAR(fn.sec_insert_date) <> 2024
    ORDER BY fn.fn_id DESC;
    """

    # 6. Customer Credit & Debt Roster (Accounts Receivable)
    q_cust_debt = """
    SELECT 
        c.cust_id AS Customer_ID,
        c.cust_code AS Customer_Code,
        c.cust_name_ar AS Customer_Name,
        c.cust_mobile AS Phone_Number,
        c.cust_max_credit AS Authorized_Credit_Limit,
        c.cust_start_debit AS Starting_Debit,
        c.cust_start_credit AS Starting_Credit,
        c.cust_current_credit AS Current_Balance,
        CASE 
            WHEN c.cust_current_credit > 0 THEN 'Customer Owes Pharmacy - Receivable (مدين)' 
            WHEN c.cust_current_credit < 0 THEN 'Pharmacy Owes Customer - Advance (دائن)' 
            ELSE 'Settled (خالص)' 
        END AS Position_Status,
        c.sec_update_date AS Last_Credit_Change
    FROM Customer c WITH (NOLOCK)
    WHERE c.cust_active = '1'
    ORDER BY c.cust_current_credit DESC;
    """

    # 7. Customer Credit Change Forensic Log
    q_customer_credit_audit = """
    SELECT 
        cc.cc_id AS Audit_ID,
        c.cust_name_ar AS Customer_Name,
        cc.cc_old_credit AS Previous_Debt_Balance,
        cc.cc_new_credit AS Updated_Debt_Balance,
        CAST(cc.cc_new_credit - cc.cc_old_credit AS DECIMAL(18,2)) AS Balance_Delta,
        cc.cc_operation AS Operation_Event,
        cc.cc_prog_name AS Originating_Software,
        cc.cc_pc_name AS Workstation_Hostname,
        cc.cc_insert_date AS Action_Timestamp
    FROM customer_credit_chng cc WITH (NOLOCK)
    JOIN Customer c WITH (NOLOCK) ON c.cust_id = cc.cc_cust_id
    WHERE YEAR(cc.cc_insert_date) <> 2024
    ORDER BY cc.cc_id DESC;
    """

    # 8. Hierarchical Chart of Accounts
    q_chart_of_accounts = """
    SELECT 
        ac.ac_id AS Account_ID,
        ac.ac_code AS Account_Code,
        ac.ac_name_ar AS Account_Name_Ar,
        ac.ac_name_en AS Account_Name_En,
        ac.ac_parent AS Parent_Account_ID,
        parent.ac_name_ar AS Parent_Account_Name,
        ISNULL(ac.ac_type, 'General Ledger') AS Account_Classification,
        ac.ac_start_debit AS Initial_Debit,
        ac.ac_start_credit AS Initial_Credit,
        ac.open_amount AS Opening_Ledger_Balance,
        ac.sec_update_date AS Last_Account_Review
    FROM F_Account_Chart ac WITH (NOLOCK)
    LEFT JOIN F_Account_Chart parent WITH (NOLOCK) ON parent.ac_id = ac.ac_parent
    ORDER BY ac.ac_code ASC;
    """

    # 9. Operating Expenses Category Mapping to GL
    q_operating_expenses = """
    SELECT 
        oe.oe_id AS Expense_Type_ID,
        oe.oe_code AS Expense_Code,
        oe.oe_name_ar AS Expense_Category_Name,
        ac.ac_code AS Linked_GL_Account_Code,
        ac.ac_name_ar AS Linked_GL_Account_Name,
        CASE oe.oe_active WHEN '1' THEN 'Active' ELSE 'Inactive' END AS Category_Status,
        oe.sec_insert_date AS Created_At
    FROM Other_Expenses oe WITH (NOLOCK)
    LEFT JOIN F_Account_Chart ac WITH (NOLOCK) ON ac.ac_id = oe.oe_ac_id
    ORDER BY oe.oe_id ASC;
    """

    # 10. Automated GL Parameter Mapping Rules
    q_gl_rules = """
    SELECT 
        acs.acs_id AS Setting_ID,
        acs.acs_code AS Parameter_Code,
        acs.acs_name_ar AS Transaction_Type_Name,
        acs.acs_ac_id AS Target_Account_ID,
        ac.ac_code AS Target_GL_Account_Code,
        ac.ac_name_ar AS Target_GL_Account_Name
    FROM F_Account_Settings acs WITH (NOLOCK)
    LEFT JOIN F_Account_Chart ac WITH (NOLOCK) ON ac.ac_id = acs.acs_ac_id
    ORDER BY acs.acs_id;
    """

    # 11. Staff Monthly Payroll & Statutory Deductions Ledger
    q_staff_payroll = """
    SELECT 
        sal.sal_year AS Payout_Year,
        sal.sal_mon AS Payout_Month,
        emp.e_code AS Staff_Code,
        emp.e_Name AS Staff_Name,
        j.j_name_ar AS Job_Role,
        sal.salary AS Base_Salary,
        sal.ch_salary AS Variable_Salary,
        sal.sal_bouns AS Monthly_Bonus,
        sal.hr_late AS Late_Hours_Penalty,
        sal.days_off AS Days_Absent,
        sal.sal_pen AS Disciplinary_Deductions,
        sal.sal_fx_ins_emp + sal.sal_ch_ins_emp AS Employee_Social_Insurance,
        sal.sal_tx AS Income_Tax_Withheld,
        sal.sal_amount AS Net_Salary_Disbursed,
        sal.notes AS Payroll_Notes
    FROM mon_sal sal WITH (NOLOCK)
    JOIN Employee emp WITH (NOLOCK) ON emp.e_id = sal.sal_e_id
    LEFT JOIN Jobs j WITH (NOLOCK) ON j.j_id = emp.e_job_id
    WHERE sal.sal_year <> 2024
    ORDER BY sal.sal_year DESC, sal.sal_mon DESC, sal.sal_amount DESC;
    """

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        read_sql_and_enrich(q_cash_stores, ENGINE).to_excel(
            writer, sheet_name="Cash_Safes_Banks", index=False
        )
        read_sql_and_enrich(q_treasury_trans, ENGINE).to_excel(
            writer, sheet_name="Treasury_Ledger", index=False
        )
        read_sql_and_enrich(q_cheques, ENGINE).to_excel(
            writer, sheet_name="Cheques_Portfolio", index=False
        )
        read_sql_and_enrich(q_gl_engine, ENGINE).to_excel(
            writer, sheet_name="Automated_GL_Journals", index=False
        )
        read_sql_and_enrich(q_notices, ENGINE).to_excel(
            writer, sheet_name="Debit_Credit_Notices", index=False
        )
        read_sql_and_enrich(q_cust_debt, ENGINE).to_excel(
            writer, sheet_name="Customer_Credit_Balances", index=False
        )
        read_sql_and_enrich(q_customer_credit_audit, ENGINE).to_excel(
            writer, sheet_name="Customer_Credit_Audit_Log", index=False
        )
        read_sql_and_enrich(q_chart_of_accounts, ENGINE).to_excel(
            writer, sheet_name="Chart_Of_Accounts", index=False
        )
        read_sql_and_enrich(q_operating_expenses, ENGINE).to_excel(
            writer, sheet_name="Operating_Expenses_Registry", index=False
        )
        read_sql_and_enrich(q_gl_rules, ENGINE).to_excel(
            writer, sheet_name="GL_Posting_Rules_Settings", index=False
        )
        read_sql_and_enrich(q_staff_payroll, ENGINE).to_excel(
            writer, sheet_name="Staff_Payroll_Disbursements", index=False
        )

    wb = openpyxl.load_workbook(filename)
    table_names = set()
    for ws_name in wb.sheetnames:
        style_worksheet(
            wb[ws_name],
            header_fill_hex="7D5A00",
            is_kpi=False,
            existing_table_names=table_names,
        )
    wb.save(filename)
    logger.info("Module 4 Complete.")


# ==================================================================================================
# MODULE 5: POS CASHIER & SHIFT AUDIT TRAIL
# ==================================================================================================
def generate_shifts_workbook(output_dir):
    filename = os.path.join(output_dir, "05_Genius_POS_Shift_Operations.xlsx")
    logger.info(f"Generating Module 5: {filename}...")

    # 1. POS Cashier Shifts Closure Summary
    q_shifts = """
    SELECT 
        sc.sc_id AS Shift_Log_ID,
        emp.e_Name AS Cashier_Name,
        fcs.fcs_name_ar AS Cash_Drawer,
        sc.sc_shift_start_time AS Shift_Started_At,
        sc.sec_insert_date AS Shift_Closed_At,
        sc.sc_start_cash AS Starting_Float_Cash,
        sc.sc_curr_cash AS System_Calculated_Cash,
        sc.sc_act_cash AS Cashier_Actual_Counted_Cash,
        CAST(sc.sc_act_cash - sc.sc_curr_cash AS DECIMAL(18,2)) AS Cash_Variance_Over_Short,
        CASE 
            WHEN (sc.sc_act_cash - sc.sc_curr_cash) > 0 THEN 'Surplus (زيادة)' 
            WHEN (sc.sc_act_cash - sc.sc_curr_cash) < 0 THEN 'Deficit/Shortage (عجز)' 
            ELSE 'Balanced (مطابق)' 
        END AS Reconciliation_Status,
        sc.sc_trans_value AS Amount_Transferred_To_Vault,
        next_emp.e_Name AS Transferred_To_Employee,
        sc.sc_curr_credit AS Visa_System_Total,
        sc.sc_act_credit AS Visa_Actual_Slip_Total,
        sc.sc_notes AS Shift_Notes
    FROM Shift_Changes sc WITH (NOLOCK)
    JOIN Employee emp WITH (NOLOCK) ON emp.e_id = sc.sc_cashier_id
    JOIN F_Cash_Store fcs WITH (NOLOCK) ON fcs.fcs_id = sc.sc_fcs_id
    LEFT JOIN Employee next_emp WITH (NOLOCK) ON next_emp.e_id = TRY_CAST(sc.sc_to_employee AS DECIMAL(18,0))
    WHERE YEAR(sc.sec_insert_date) <> 2024
    ORDER BY sc.sc_id DESC;
    """

    # 2. Shift Operations Breakdown Details
    q_shifts_d = """
    SELECT 
        scd.scd_sc_id AS Shift_Log_ID,
        st.sto_name_ar AS Store_Name,
        fcs.fcs_name_ar AS Cash_Drawer,
        scd.scd_start AS Shift_Open,
        scd.scd_end AS Shift_Close,
        scd.scd_d_startBalance AS Opening_Balance,
        scd.scd_d_sales AS Cash_Sales_In,
        scd.scd_d_salesfromreturn AS Sales_From_Returns,
        scd.scd_d_visacollect AS Visa_Collections,
        scd.scd_d_cashIn AS Manual_Cash_In,
        scd.scd_c_purchases AS Outbound_Cash_Purchases,
        scd.scd_c_salesReturn AS Outbound_Sales_Refunds,
        scd.scd_c_cashOut AS Manual_Cash_Out_Expenses,
        scd.scd_o_salesVisa AS Visa_Card_Sales,
        scd.scd_o_salesCredit AS Credit_Account_Sales,
        scd.scd_o_salesPending AS Suspended_Bills_Value
    FROM Shift_Changes_d scd WITH (NOLOCK)
    JOIN Store st WITH (NOLOCK) ON st.sto_id = scd.scd_sto_id
    JOIN F_Cash_Store fcs WITH (NOLOCK) ON fcs.fcs_id = scd.scd_fcs_id
    WHERE (scd.scd_start IS NULL OR YEAR(scd.scd_start) <> 2024)
    ORDER BY scd.scd_sc_id DESC;
    """

    # 3. User Login / Logout Station Sessions
    q_sessions = """
    SELECT 
        lu.lg_id AS Session_ID,
        usr.usr_name AS System_Username,
        emp.e_Name AS Employee_Full_Name,
        fcs.fcs_name_ar AS POS_Drawer_Assigned,
        lu.computer_nm AS Client_Workstation,
        lu.prog_ver AS ePlus_Software_Version,
        lu.lg_begin_tm AS Login_Timestamp,
        lu.lg_close_tm AS Logout_Timestamp,
        DATEDIFF(MINUTE, lu.lg_begin_tm, lu.lg_close_tm) AS Session_Duration_Minutes
    FROM log_users lu WITH (NOLOCK)
    LEFT JOIN Users usr WITH (NOLOCK) ON usr.usr_e_id = TRY_CAST(lu.lg_usr_id AS DECIMAL(18,0))
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = usr.usr_e_id
    LEFT JOIN F_Cash_Store fcs WITH (NOLOCK) ON fcs.fcs_id = lu.f_cash_store_id
    WHERE YEAR(lu.lg_begin_tm) <> 2024
    ORDER BY lu.lg_id DESC;
    """

    # 4. Staff Shift Schedule & Privileges
    q_staff = """
    SELECT 
        e.e_id AS Employee_ID,
        e.e_code AS Employee_Code,
        e.e_Name AS Staff_Name,
        j.j_name_ar AS Job_Title,
        ews.ews_name AS Assigned_Shift_Scheme,
        ews.ews_start AS Shift_Default_Start,
        ews.ews_end AS Shift_Default_End,
        e.e_bill_disc_perc AS Max_Allowed_Bill_Discount,
        CASE e.e_select_cashier WHEN 1 THEN 'Yes' ELSE 'No' END AS Can_Select_Cashier,
        CASE e.e_change_sales_prices WHEN 1 THEN 'Yes' ELSE 'No' END AS Can_Override_Prices,
        CASE e.e_can_pend WHEN 1 THEN 'Yes' ELSE 'No' END AS Can_Suspend_Bills,
        e.sec_update_date AS Profile_Last_Updated
    FROM Employee e WITH (NOLOCK)
    LEFT JOIN Jobs j WITH (NOLOCK) ON j.j_id = e.e_job_id
    LEFT JOIN Emp_Work_Shifts ews WITH (NOLOCK) ON ews.ews_id = e.e_ews_id
    WHERE e.e_active = 1
    ORDER BY e.e_id ASC;
    """

    # 5. Shift Cash Drawer Expense Payouts
    q_drawer_payouts = """
    SELECT 
        scd.scd_sc_id AS Shift_Log_ID,
        st.sto_name_ar AS Store_Name,
        fcs.fcs_name_ar AS Cash_Drawer,
        scd.scd_start AS Shift_Open,
        scd.scd_end AS Shift_Close,
        scd.scd_c_purchases AS Outbound_Local_Purchases_From_Drawer,
        scd.scd_c_cashOut AS Petty_Cash_Operating_Expenses,
        scd.scd_c_salesReturn AS Cash_Refunds_Given_To_Customers,
        scd.scd_c_NoticeOut AS Credit_Debit_Notices_Paid,
        CAST(scd.scd_c_purchases + scd.scd_c_cashOut + scd.scd_c_salesReturn AS DECIMAL(18,2)) AS Total_Cash_Drained_From_Register,
        COALESCE(emp.e_Name, CAST(scd.scd_insert_uid AS VARCHAR(50))) AS Cashier_Name
    FROM Shift_Changes_d scd WITH (NOLOCK)
    JOIN Store st WITH (NOLOCK) ON st.sto_id = scd.scd_sto_id
    JOIN F_Cash_Store fcs WITH (NOLOCK) ON fcs.fcs_id = scd.scd_fcs_id
    LEFT JOIN Employee emp WITH (NOLOCK) ON emp.e_id = TRY_CAST(scd.scd_insert_uid AS DECIMAL(18,0))
    WHERE (scd.scd_c_purchases + scd.scd_c_cashOut + scd.scd_c_salesReturn) > 0
      AND (scd.scd_start IS NULL OR YEAR(scd.scd_start) <> 2024)
    ORDER BY scd.scd_sc_id DESC;
    """

    # 6. Cashier Drawer Authorizations Matrix
    q_drawer_authorizations = """
    SELECT 
        ecs.ecs_e_id AS Employee_ID,
        emp.e_Name AS Staff_Name,
        st.sto_name_ar AS Branch_Assigned,
        fcs.fcs_name_ar AS Authorized_POS_Drawer,
        fcs.fcs_computer_name AS Dedicated_Workstation,
        CASE emp.e_select_cashier WHEN 1 THEN 'Yes' ELSE 'No' END AS Can_Switch_Drawers,
        ecs.ecs_insert_date AS Assignment_Date
    FROM Emp_cash_store ecs WITH (NOLOCK)
    JOIN Employee emp WITH (NOLOCK) ON emp.e_id = ecs.ecs_e_id
    JOIN Store st WITH (NOLOCK) ON st.sto_id = ecs.ecs_sto_id
    JOIN F_Cash_Store fcs WITH (NOLOCK) ON fcs.fcs_id = ecs.ecs_fcs_id
    ORDER BY st.sto_id, emp.e_id;
    """

    # 7. Staff Store Security Permissions
    q_emp_stores = """
    SELECT 
        es.es_e_id AS Employee_ID,
        emp.e_Name AS Staff_Name,
        st.sto_name_ar AS Authorized_Store,
        CASE es.es_readOnly WHEN 1 THEN 'Read-Only (عرض فقط)' ELSE 'Full Read/Write (كامل)' END AS Permission_Level,
        es.sec_update_date AS Authorization_Timestamp
    FROM Emp_Stores es WITH (NOLOCK)
    JOIN Employee emp WITH (NOLOCK) ON emp.e_id = es.es_e_id
    JOIN Store st WITH (NOLOCK) ON st.sto_id = es.es_sto_id
    ORDER BY emp.e_id, st.sto_id;
    """

    # 8. Employee Base Salary & Contract Terms
    q_emp_salary_terms = """
    SELECT 
        es.es_id AS Contract_ID,
        emp.e_Name AS Staff_Name,
        j.j_name_ar AS Role_Title,
        es.es_Type AS Salary_Periodicity,
        es.es_value AS Base_Contracted_Salary,
        es.es_over_h AS Overtime_Hourly_Rate,
        es.es_insurance_disc AS Contracted_Insurance_Deduction,
        es.es_tax_disc AS Contracted_Tax_Deduction,
        es.es_date AS Contract_Effective_Date
    FROM Emp_Salary es WITH (NOLOCK)
    JOIN Employee emp WITH (NOLOCK) ON emp.e_id = es.es_e_id
    LEFT JOIN Jobs j WITH (NOLOCK) ON j.j_id = emp.e_job_id
    ORDER BY emp.e_id;
    """

    # 9. Duplicate Financial Voucher Interceptions Log
    q_duplicate_vouchers = """
    SELECT 
        fh_id AS Intercepted_Record_ID,
        fh_code AS Voucher_Code,
        fh_value AS Voucher_Value,
        fh_notes AS Operational_Narrative,
        fh_actual_date AS Voucher_Date,
        fh_computer AS Workstation_Hostname,
        sec_insert_date AS Purged_Timestamp
    FROM F_Transaction_Header_SaveDeleteRecords WITH (NOLOCK)
    WHERE YEAR(fh_actual_date) <> 2024
    ORDER BY fh_actual_date DESC;
    """

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        read_sql_and_enrich(q_shifts, ENGINE).to_excel(
            writer, sheet_name="Shift_Reconciliation", index=False
        )
        read_sql_and_enrich(q_shifts_d, ENGINE).to_excel(
            writer, sheet_name="Shift_Operations_Breakdown", index=False
        )
        read_sql_and_enrich(q_sessions, ENGINE).to_excel(
            writer, sheet_name="Station_Login_Sessions", index=False
        )
        read_sql_and_enrich(q_staff, ENGINE).to_excel(
            writer, sheet_name="Cashier_Staff_Roster", index=False
        )
        read_sql_and_enrich(q_drawer_payouts, ENGINE).to_excel(
            writer, sheet_name="Drawer_Expense_Payouts", index=False
        )
        read_sql_and_enrich(q_drawer_authorizations, ENGINE).to_excel(
            writer, sheet_name="Cashier_Drawer_Authorizations", index=False
        )
        read_sql_and_enrich(q_emp_stores, ENGINE).to_excel(
            writer, sheet_name="Staff_Store_Permissions", index=False
        )
        read_sql_and_enrich(q_emp_salary_terms, ENGINE).to_excel(
            writer, sheet_name="Employee_Contract_Terms", index=False
        )
        read_sql_and_enrich(q_duplicate_vouchers, ENGINE).to_excel(
            writer, sheet_name="Duplicate_Voucher_Interceptions", index=False
        )

    wb = openpyxl.load_workbook(filename)
    table_names = set()
    for ws_name in wb.sheetnames:
        style_worksheet(
            wb[ws_name],
            header_fill_hex="1A5276",
            is_kpi=False,
            existing_table_names=table_names,
        )
    wb.save(filename)
    logger.info("Module 5 Complete.")


# ==================================================================================================
# MODULE 6: EXECUTIVE BUSINESS INTELLIGENCE & ANALYTICAL COCKPIT
# ==================================================================================================
def generate_executive_bi_workbook(output_dir):
    filename = os.path.join(output_dir, "06_Genius_Executive_Intelligence_Cockpit.xlsx")
    logger.info(f"Generating Module 6: {filename}...")

    # 1. Executive Master KPI Scorecard (Accurate Signs & Header Discount Factoring)
    q_kpi_scorecard = """
    SELECT '1. Total Gross Sales Revenue' AS Metric, 
           CAST(ISNULL(SUM(total_bill), 0) AS DECIMAL(18,2)) AS Value_L_E, 
           'Total customer sales before discounts (Excl. 2024)' AS Notes 
    FROM sales_trans_h WITH (NOLOCK) WHERE sth_flag='C' AND YEAR(sec_insert_date) <> 2024

    UNION ALL
    SELECT '2. Total Customer Discounts Conceded', 
           CAST(ISNULL(SUM(total_bill - total_bill_after_disc), 0) AS DECIMAL(18,2)), 
           'Customer discounts granted at invoice level (Excl. 2024)' 
    FROM sales_trans_h WITH (NOLOCK) WHERE sth_flag='C' AND YEAR(sec_insert_date) <> 2024

    UNION ALL
    SELECT '3. Total Net Sales Realized', 
           CAST(ISNULL(SUM(total_bill_after_disc), 0) AS DECIMAL(18,2)), 
           'Actual top-line revenue generated (Excl. 2024)' 
    FROM sales_trans_h WITH (NOLOCK) WHERE sth_flag='C' AND YEAR(sec_insert_date) <> 2024

    UNION ALL
    SELECT '4. Total Cost of Goods Sold (COGS)', 
           CAST(ISNULL(SUM((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_cost), 0) AS DECIMAL(18,2)), 
           'Direct pharmacy acquisition cost of net dispensed merchandise (Excl. 2024)' 
    FROM sales_trans_d d WITH (NOLOCK) 
    JOIN sales_trans_h h WITH (NOLOCK) ON h.sth_id=d.sth_id 
    WHERE h.sth_flag='C' AND ISNULL(d.itm_nexist, 0) <> 1 AND YEAR(h.sec_insert_date) <> 2024

    UNION ALL
    SELECT '5. Cumulative Realized Gross Profit', 
           CAST(ISNULL(
               SUM(
                   (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell * (1.0 - (ISNULL(d.itm_dis_per, 0) / 100.0))) - ISNULL(d.itm_dis_mon, 0))
                   * (1.0 - (ISNULL(h.total_dis_per, 0) / 100.0))
                   - ISNULL(ISNULL(h.total_des_mon, 0) / NULLIF(h.no_of_items, 0), 0)
               ) - SUM((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_cost), 0) AS DECIMAL(18,2)), 
           'Net Revenue minus Total Product Cost (Excl. 2024)' 
    FROM sales_trans_d d WITH (NOLOCK) 
    JOIN sales_trans_h h WITH (NOLOCK) ON h.sth_id=d.sth_id 
    WHERE h.sth_flag='C' AND ISNULL(d.itm_nexist, 0) <> 1 AND YEAR(h.sec_insert_date) <> 2024

    UNION ALL
    SELECT '6. Total Retail Dispensary Stock (At Cost)', 
           CAST(ISNULL(SUM(itm_qty * pharm_price), 0) AS DECIMAL(18,2)), 
           'Current capital tied up in Main Pharmacy Dispensary (Store 47)' 
    FROM Item_Class_Store WITH (NOLOCK) WHERE itm_qty > 0 AND sto_id = 47

    UNION ALL
    SELECT '7. Total Consignment / Quarantine Stock (At Cost)', 
           CAST(ISNULL(SUM(itm_qty * pharm_price), 0) AS DECIMAL(18,2)), 
           'Current stock under quarantine or consignment (Store 48: تحت التصريف)' 
    FROM Item_Class_Store WITH (NOLOCK) WHERE itm_qty > 0 AND sto_id = 48

    UNION ALL
    SELECT '8. Total Combined Stock (At Retail)', 
           CAST(ISNULL(SUM(itm_qty * sell_price), 0) AS DECIMAL(18,2)), 
           'Potential retail revenue of all on-shelf stock across all branches' 
    FROM Item_Class_Store WITH (NOLOCK) WHERE itm_qty > 0

    UNION ALL
    SELECT '9. Total Customer Outstanding Debt (Accounts Receivable)', 
           CAST(ISNULL(SUM(cust_current_credit), 0) AS DECIMAL(18,2)), 
           'Total money owed to pharmacy by credit customers (Receivables)' 
    FROM Customer WITH (NOLOCK) WHERE cust_current_credit > 0 AND cust_active = '1'

    UNION ALL
    SELECT '10. Total Vendor Liabilities (Accounts Payable)', 
           CAST(ISNULL(SUM(ven_current_credit), 0) AS DECIMAL(18,2)), 
           'Total unpaid invoices owed to pharmaceutical distributors (We owe them)' 
    FROM Vendor WITH (NOLOCK) WHERE ven_current_credit > 0 AND ven_active = '1'

    UNION ALL
    SELECT '11. Total Vendor Advances & Return Balances', 
           CAST(ISNULL(ABS(SUM(ven_current_credit)), 0) AS DECIMAL(18,2)), 
           'Total distributor credit balances due back to pharmacy (They owe us)' 
    FROM Vendor WITH (NOLOCK) WHERE ven_current_credit < 0 AND ven_active = '1'

    UNION ALL
    SELECT '12. Total Live Liquidity in Safes & Banks', 
           CAST(ISNULL(SUM(fcs_current_balance), 0) AS DECIMAL(18,2)), 
           'Total live cash available across all drawers, vaults and bank accounts' 
    FROM F_Cash_Store WITH (NOLOCK) WHERE fcs_active = 1

    UNION ALL
    SELECT '13. Net POS Cash Drawer Discrepancies (Shortage/Surplus)', 
           CAST(ISNULL(SUM(sc_act_cash - sc_curr_cash), 0) AS DECIMAL(18,2)), 
           'Cumulative historical cashier cash drawer variance (Excl. 2024)' 
    FROM Shift_Changes WITH (NOLOCK) WHERE YEAR(sec_insert_date) <> 2024;
    """

    # 2. Executive Monthly Financial Income Statement (P&L)
    q_monthly_pnl = """
    WITH MonthlySales AS (
        SELECT 
            YEAR(h.sec_insert_date) AS Fin_Year,
            MONTH(h.sec_insert_date) AS Fin_Month,
            CAST(SUM(h.total_bill) AS DECIMAL(18,2)) AS Gross_Sales,
            CAST(SUM(h.total_bill - h.total_bill_after_disc) AS DECIMAL(18,2)) AS Discounts_Granted,
            CAST(SUM(h.total_bill_after_disc) AS DECIMAL(18,2)) AS Net_Revenue,
            CAST(SUM((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_cost) AS DECIMAL(18,2)) AS COGS
        FROM sales_trans_h h WITH (NOLOCK)
        JOIN sales_trans_d d WITH (NOLOCK) ON d.sth_id = h.sth_id
        WHERE h.sth_flag = 'C' AND ISNULL(d.itm_nexist, 0) <> 1 AND YEAR(h.sec_insert_date) <> 2024
        GROUP BY YEAR(h.sec_insert_date), MONTH(h.sec_insert_date)
    ),
    MonthlyDrawerExpenses AS (
        SELECT 
            YEAR(scd_start) AS Fin_Year,
            MONTH(scd_start) AS Fin_Month,
            CAST(SUM(scd_c_cashOut) AS DECIMAL(18,2)) AS Drawer_Petty_Expenses
        FROM Shift_Changes_d WITH (NOLOCK)
        WHERE scd_start IS NOT NULL AND YEAR(scd_start) <> 2024
        GROUP BY YEAR(scd_start), MONTH(scd_start)
    ),
    MonthlyPayroll AS (
        SELECT 
            sal_year AS Fin_Year,
            sal_mon AS Fin_Month,
            CAST(SUM(sal_amount) AS DECIMAL(18,2)) AS Total_Payroll_Disbursed
        FROM mon_sal WITH (NOLOCK)
        WHERE sal_year <> 2024
        GROUP BY sal_year, sal_mon
    )
    SELECT 
        s.Fin_Year,
        s.Fin_Month,
        s.Gross_Sales,
        s.Discounts_Granted,
        s.Net_Revenue,
        s.COGS,
        CAST(s.Net_Revenue - s.COGS AS DECIMAL(18,2)) AS Gross_Profit,
        CAST((s.Net_Revenue - s.COGS) / NULLIF(s.Net_Revenue, 0) AS DECIMAL(18,4)) AS Gross_Margin_Perc,
        ISNULL(e.Drawer_Petty_Expenses, 0.00) AS POS_Operating_Expenses,
        ISNULL(p.Total_Payroll_Disbursed, 0.00) AS Payroll_Personnel_Expenses,
        CAST((s.Net_Revenue - s.COGS) - ISNULL(e.Drawer_Petty_Expenses, 0) - ISNULL(p.Total_Payroll_Disbursed, 0) AS DECIMAL(18,2)) AS Estimated_Net_Operating_Profit
    FROM MonthlySales s
    LEFT JOIN MonthlyDrawerExpenses e ON e.Fin_Year = s.Fin_Year AND e.Fin_Month = s.Fin_Month
    LEFT JOIN MonthlyPayroll p ON p.Fin_Year = s.Fin_Year AND p.Fin_Month = s.Fin_Month
    ORDER BY s.Fin_Year DESC, s.Fin_Month DESC;
    """

    # 3. ABC Pareto Analysis (With Mathematical Multi-Unit Pack Normalization)
    q_abc_analysis = """
    WITH ItemSales AS (
        SELECT 
            ic.itm_id,
            ic.itm_code AS Barcode,
            ic.itm_name_ar AS Item_Name,
            c.com_name_ar AS Manufacturer,
            CAST(SUM((d.qnty - ISNULL(d.itm_back, 0)) * (
                CASE d.itm_unit 
                    WHEN 3 THEN (1.0 / NULLIF(ic.itm_unit1_unit3, 0)) 
                    WHEN 2 THEN (1.0 / NULLIF(ic.itm_unit1_unit2, 0)) 
                    ELSE 1.0 
                END
            )) AS DECIMAL(18,2)) AS Standard_Packs_Sold,
            CAST(SUM(
                (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell * (1.0 - (ISNULL(d.itm_dis_per, 0) / 100.0))) - ISNULL(d.itm_dis_mon, 0))
                * (1.0 - (ISNULL(h.total_dis_per, 0) / 100.0))
                - ISNULL(ISNULL(h.total_des_mon, 0) / NULLIF(h.no_of_items, 0), 0)
            ) AS DECIMAL(18,2)) AS Net_Revenue,
            CAST(SUM((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_cost) AS DECIMAL(18,2)) AS Total_COGS,
            CAST(SUM(
                (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell * (1.0 - (ISNULL(d.itm_dis_per, 0) / 100.0))) - ISNULL(d.itm_dis_mon, 0))
                * (1.0 - (ISNULL(h.total_dis_per, 0) / 100.0))
                - ISNULL(ISNULL(h.total_des_mon, 0) / NULLIF(h.no_of_items, 0), 0)
            ) - SUM((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_cost) AS DECIMAL(18,2)) AS Gross_Profit
        FROM sales_trans_d d WITH (NOLOCK)
        JOIN sales_trans_h h WITH (NOLOCK) ON h.sth_id = d.sth_id
        JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = d.itm_id
        LEFT JOIN Company c WITH (NOLOCK) ON c.com_id = ic.com_id
        WHERE h.sth_flag = 'C'
          AND ISNULL(ic.itm_srvc, 0) = 0
          AND ISNULL(d.itm_nexist, 0) <> 1
          AND YEAR(h.sec_insert_date) <> 2024
        GROUP BY ic.itm_id, ic.itm_code, ic.itm_name_ar, c.com_name_ar
        HAVING SUM(d.qnty - ISNULL(d.itm_back, 0)) > 0
    ),
    RankedSales AS (
        SELECT *,
            SUM(Net_Revenue) OVER () AS Total_Market_Revenue,
            SUM(Net_Revenue) OVER (ORDER BY Net_Revenue DESC ROWS BETWEEN UNBOUNDED PRECEDING AND CURRENT ROW) AS Running_Revenue
        FROM ItemSales
    )
    SELECT 
        itm_id AS Item_ID,
        Barcode,
        Item_Name,
        Manufacturer,
        Standard_Packs_Sold,
        Net_Revenue,
        Total_COGS,
        Gross_Profit,
        CAST((Gross_Profit / NULLIF(Net_Revenue, 0)) AS DECIMAL(18,4)) AS Gross_Margin_Perc,
        CAST((Running_Revenue / NULLIF(Total_Market_Revenue, 0)) AS DECIMAL(18,4)) AS Cumulative_Revenue_Share,
        CASE 
            WHEN (Running_Revenue / NULLIF(Total_Market_Revenue, 0)) <= 0.80 THEN 'Class A (Top 80% Rev)'
            WHEN (Running_Revenue / NULLIF(Total_Market_Revenue, 0)) <= 0.95 THEN 'Class B (Next 15% Rev)'
            ELSE 'Class C (Bottom 5% Rev)'
        END AS ABC_Classification
    FROM RankedSales
    ORDER BY Net_Revenue DESC;
    """

    # 4. Expiry Risk & Dead Stock Radar
    q_dead_stock_radar = """
    WITH LastItemSale AS (
        SELECT d.itm_id, MAX(h.sec_insert_date) AS Last_Sold_Date
        FROM sales_trans_d d WITH (NOLOCK)
        JOIN sales_trans_h h WITH (NOLOCK) ON h.sth_id = d.sth_id
        WHERE h.sth_flag = 'C' AND YEAR(h.sec_insert_date) <> 2024
        GROUP BY d.itm_id
    )
    SELECT 
        ics.sto_id AS Store_ID,
        st.sto_name_ar AS Store_Name,
        ic.itm_id AS Item_ID,
        ic.itm_code AS Barcode,
        ic.itm_name_ar AS Item_Name,
        ics.c_id AS Batch_ID,
        ics.itm_qty AS Stock_On_Hand,
        ics.pharm_price AS Unit_Cost,
        CAST(ics.itm_qty * ics.pharm_price AS DECIMAL(18,2)) AS Capital_At_Risk,
        CONVERT(VARCHAR(10), ics.itm_expiry_date, 120) AS Expiry_Date,
        DATEDIFF(DAY, GETDATE(), ics.itm_expiry_date) AS Days_Until_Expiration,
        CASE 
            WHEN ics.itm_expiry_date < GETDATE() THEN 'EXPIRED (منتهي الصلاحية)'
            WHEN DATEDIFF(DAY, GETDATE(), ics.itm_expiry_date) <= 90 THEN 'CRITICAL (< 90 Days)'
            WHEN DATEDIFF(DAY, GETDATE(), ics.itm_expiry_date) <= 180 THEN 'WARNING (< 180 Days)'
            ELSE 'STABLE'
        END AS Expiry_Risk_Category,
        ISNULL(lis.Last_Sold_Date, '2000-01-01') AS Last_Sold_Date
    FROM Item_Class_Store ics WITH (NOLOCK)
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = ics.itm_id
    JOIN Store st WITH (NOLOCK) ON st.sto_id = ics.sto_id
    LEFT JOIN LastItemSale lis ON lis.itm_id = ics.itm_id
    WHERE ics.itm_qty > 0 
      AND (ics.itm_expiry_date <= DATEADD(MONTH, 6, GETDATE()) OR ics.itm_expiry_date IS NULL)
    ORDER BY Days_Until_Expiration ASC;
    """

    # 5. Hourly Sales Distribution & Traffic Velocity
    q_hourly_traffic = """
    SELECT 
        DATEPART(HOUR, sec_insert_date) AS Hour_Of_Day_24H,
        COUNT(sth_id) AS Total_Bills_Count,
        CAST(SUM(total_bill_after_disc) AS DECIMAL(18,2)) AS Total_Revenue_Generated,
        CAST(AVG(total_bill_after_disc) AS DECIMAL(18,2)) AS Average_Basket_Size,
        CAST(SUM(total_bill_after_disc) / NULLIF((SELECT SUM(total_bill_after_disc) FROM sales_trans_h WHERE sth_flag='C' AND YEAR(sec_insert_date) <> 2024), 0) AS DECIMAL(18,4)) AS Share_Of_Total_Sales
    FROM sales_trans_h WITH (NOLOCK)
    WHERE sth_flag = 'C' AND YEAR(sec_insert_date) <> 2024
    GROUP BY DATEPART(HOUR, sec_insert_date)
    ORDER BY Hour_Of_Day_24H ASC;
    """

    # 6. Distributor / Vendor Supply Performance Matrix
    q_vendor_performance = """
    SELECT 
        v.ven_id AS Vendor_ID,
        v.ven_name_ar AS Vendor_Name,
        COUNT(DISTINCT pth.pth_id) AS Total_Purchase_Bills,
        CAST(SUM(pth.total_bill) AS DECIMAL(18,2)) AS Total_Purchased_Gross,
        CAST(SUM(pth.total_after_back) AS DECIMAL(18,2)) AS Total_Purchased_Net,
        CAST(AVG(pth.total_dis_per) AS DECIMAL(18,2)) AS Avg_Negotiated_Discount_Perc,
        ISNULL((
            SELECT CAST(SUM(ibh_total_value) AS DECIMAL(18,2)) 
            FROM Item_back_h WITH (NOLOCK) 
            WHERE ibh_ven_id = v.ven_id AND YEAR(sec_insert_date) <> 2024
        ), 0.00) AS Total_Returned_Value,
        v.ven_current_credit AS Current_Outstanding_Debt
    FROM Vendor v WITH (NOLOCK)
    JOIN pur_trans_h pth WITH (NOLOCK) ON pth.ven_id = v.ven_id
    WHERE pth.bill_status <> 1 AND YEAR(pth.sec_insert_date) <> 2024
    GROUP BY v.ven_id, v.ven_name_ar, v.ven_current_credit
    ORDER BY Total_Purchased_Net DESC;
    """

    # 7. Branch Operational Benchmark (Store 47 Retail vs Store 48 Consignment)
    q_store_matrix = """
    SELECT 
        st.sto_id AS Store_ID,
        st.sto_name_ar AS Store_Name,
        CASE WHEN st.sto_id = 47 THEN 'Dispensary Retail' ELSE 'Consignment / Quarantine' END AS Store_Role,
        ISNULL((SELECT COUNT(*) FROM Item_Class_Store WHERE sto_id=st.sto_id AND itm_qty > 0), 0) AS Active_SKUs_In_Stock,
        ISNULL((SELECT CAST(SUM(itm_qty * pharm_price) AS DECIMAL(18,2)) FROM Item_Class_Store WHERE sto_id=st.sto_id AND itm_qty > 0), 0) AS Total_Stock_Cost_Valuation,
        ISNULL((SELECT CAST(SUM(itm_qty * sell_price) AS DECIMAL(18,2)) FROM Item_Class_Store WHERE sto_id=st.sto_id AND itm_qty > 0), 0) AS Total_Stock_Retail_Valuation,
        ISNULL((SELECT COUNT(*) FROM sales_trans_h WHERE sto_id=st.sto_id AND sth_flag='C' AND YEAR(sec_insert_date) <> 2024), 0) AS Total_Sales_Transactions,
        ISNULL((SELECT CAST(SUM(total_bill_after_disc) AS DECIMAL(18,2)) FROM sales_trans_h WHERE sto_id=st.sto_id AND sth_flag='C' AND YEAR(sec_insert_date) <> 2024), 0) AS Total_Sales_Revenue,
        ISNULL((
            SELECT CAST(SUM(sc.sc_act_cash - sc.sc_curr_cash) AS DECIMAL(18,2)) 
            FROM Shift_Changes sc WITH (NOLOCK) 
            JOIN Shift_Changes_d scd WITH (NOLOCK) ON scd.scd_sc_id=sc.sc_id 
            WHERE scd.scd_sto_id=st.sto_id AND YEAR(sc.sec_insert_date) <> 2024
        ), 0.00) AS Cumulative_Shift_Variance
    FROM Store st WITH (NOLOCK)
    WHERE st.activated = 1;
    """

    # 8. Monthly Inventory Valuation Trend
    q_monthly_stock_trend = """
    SELECT 
        YEAR(imc.sec_insert_date) AS Close_Year,
        MONTH(imc.sec_insert_date) AS Close_Month,
        st.sto_name_ar AS Store_Name,
        COUNT(DISTINCT imc.itm_id) AS Distinct_SKUs_Stored,
        CAST(SUM(imc.itm_qty) AS DECIMAL(18,2)) AS Total_Stock_Units,
        CAST(SUM(imc.itm_qty * imc.pharm_price) AS DECIMAL(18,2)) AS Total_Inventory_Cost_Valuation,
        CAST(SUM(imc.itm_qty * imc.sell_price) AS DECIMAL(18,2)) AS Total_Inventory_Retail_Valuation,
        CAST(SUM(imc.itm_qty * (imc.sell_price - imc.pharm_price)) AS DECIMAL(18,2)) AS Unrealized_Gross_Margin,
        CAST(SUM(imc.itm_qty * (imc.sell_price - imc.pharm_price)) / NULLIF(SUM(imc.itm_qty * imc.sell_price), 0) AS DECIMAL(18,4)) AS Potential_Profit_Margin_Perc
    FROM ICS_Month_Close imc WITH (NOLOCK)
    JOIN Store st WITH (NOLOCK) ON st.sto_id = imc.sto_id
    WHERE YEAR(imc.sec_insert_date) <> 2024
    GROUP BY YEAR(imc.sec_insert_date), MONTH(imc.sec_insert_date), st.sto_name_ar
    ORDER BY Close_Year DESC, Close_Month DESC;
    """

    # 9. Therapeutic Pharmacological Class Profitability (Pack-Normalized)
    q_therapeutic_profitability = """
    SELECT 
        ISNULL(sg.sg_name_ar, 'General Pharmaceuticals (غير مصنف)') AS Therapeutic_Class,
        COUNT(DISTINCT d.itm_id) AS Active_Medicines_Count,
        CAST(SUM((d.qnty - ISNULL(d.itm_back, 0)) * (
            CASE d.itm_unit 
                WHEN 3 THEN (1.0 / NULLIF(ic.itm_unit1_unit3, 0)) 
                WHEN 2 THEN (1.0 / NULLIF(ic.itm_unit1_unit2, 0)) 
                ELSE 1.0 
            END
        )) AS DECIMAL(18,2)) AS Standard_Packs_Sold,
        CAST(SUM(
            (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell * (1.0 - (ISNULL(d.itm_dis_per, 0)/100.0))) - ISNULL(d.itm_dis_mon, 0))
            * (1.0 - (ISNULL(h.total_dis_per, 0)/100.0))
            - ISNULL(ISNULL(h.total_des_mon, 0) / NULLIF(h.no_of_items, 0), 0)
        ) AS DECIMAL(18,2)) AS Net_Revenue,
        CAST(SUM((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_cost) AS DECIMAL(18,2)) AS Total_COGS,
        CAST(SUM(
            (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell * (1.0 - (ISNULL(d.itm_dis_per, 0)/100.0))) - ISNULL(d.itm_dis_mon, 0))
            * (1.0 - (ISNULL(h.total_dis_per, 0)/100.0))
            - ISNULL(ISNULL(h.total_des_mon, 0) / NULLIF(h.no_of_items, 0), 0)
        ) - SUM((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_cost) AS DECIMAL(18,2)) AS Gross_Profit,
        CAST((SUM(
            (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell * (1.0 - (ISNULL(d.itm_dis_per, 0)/100.0))) - ISNULL(d.itm_dis_mon, 0))
            * (1.0 - (ISNULL(h.total_dis_per, 0)/100.0))
            - ISNULL(ISNULL(h.total_des_mon, 0) / NULLIF(h.no_of_items, 0), 0)
        ) - SUM((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_cost)) 
             / NULLIF(SUM(
                 (((d.qnty - ISNULL(d.itm_back, 0)) * d.itm_sell * (1.0 - (ISNULL(d.itm_dis_per, 0)/100.0))) - ISNULL(d.itm_dis_mon, 0))
                 * (1.0 - (ISNULL(h.total_dis_per, 0)/100.0))
                 - ISNULL(ISNULL(h.total_des_mon, 0) / NULLIF(h.no_of_items, 0), 0)
             ), 0) AS DECIMAL(18,4)) AS Gross_Margin_Rate
    FROM sales_trans_d d WITH (NOLOCK)
    JOIN sales_trans_h h WITH (NOLOCK) ON h.sth_id = d.sth_id
    JOIN Item_Catalog ic WITH (NOLOCK) ON ic.itm_id = d.itm_id
    LEFT JOIN Scientific_Groups sg WITH (NOLOCK) ON sg.sg_id = ic.itm_scientific_group_id
    WHERE h.sth_flag = 'C'
      AND ISNULL(ic.itm_srvc, 0) = 0
      AND ISNULL(d.itm_nexist, 0) <> 1
      AND YEAR(h.sec_insert_date) <> 2024
    GROUP BY sg.sg_name_ar
    HAVING SUM(d.qnty - ISNULL(d.itm_back, 0)) > 0
    ORDER BY Net_Revenue DESC;
    """

    # 10. Customer Accounts Receivable Aging Schedule
    q_receivables_aging = """
    SELECT 
        c.cust_id AS Customer_ID,
        c.cust_code AS Customer_Code,
        c.cust_name_ar AS Customer_Name,
        c.cust_mobile AS Phone_Number,
        c.cust_current_credit AS Total_Outstanding_Balance,
        DATEDIFF(DAY, c.sec_update_date, GETDATE()) AS Days_Since_Last_Activity,
        CASE 
            WHEN DATEDIFF(DAY, c.sec_update_date, GETDATE()) <= 30 THEN 'Current (0-30 Days)'
            WHEN DATEDIFF(DAY, c.sec_update_date, GETDATE()) <= 60 THEN 'Aging (31-60 Days)'
            WHEN DATEDIFF(DAY, c.sec_update_date, GETDATE()) <= 90 THEN 'Overdue (61-90 Days)'
            ELSE 'Critical Default (> 90 Days)'
        END AS Aging_Bucket,
        c.cust_max_credit AS Credit_Ceiling_Limit,
        CAST(c.cust_current_credit / NULLIF(c.cust_max_credit, 0) AS DECIMAL(18,2)) AS Credit_Limit_Utilization,
        c.sec_update_date AS Last_Transaction_Date
    FROM Customer c WITH (NOLOCK)
    WHERE c.cust_active = '1' AND c.cust_current_credit > 0
    ORDER BY c.cust_current_credit DESC;
    """

    # 11. Prescription Basket Size Distribution
    q_basket_size_distribution = """
    SELECT 
        CASE 
            WHEN total_bill_after_disc < 50 THEN '1. Micro (< 50 L.E)'
            WHEN total_bill_after_disc <= 150 THEN '2. Small (50 - 150 L.E)'
            WHEN total_bill_after_disc <= 500 THEN '3. Medium (151 - 500 L.E)'
            WHEN total_bill_after_disc <= 1500 THEN '4. Large (501 - 1500 L.E)'
            ELSE '5. VIP (> 1500 L.E)'
        END AS Basket_Tier,
        COUNT(sth_id) AS Transaction_Count,
        CAST(SUM(total_bill_after_disc) AS DECIMAL(18,2)) AS Total_Revenue,
        CAST(AVG(total_bill_after_disc) AS DECIMAL(18,2)) AS Average_Basket_Value,
        CAST(SUM(total_bill_after_disc) / NULLIF((SELECT SUM(total_bill_after_disc) FROM sales_trans_h WHERE sth_flag='C' AND YEAR(sec_insert_date) <> 2024), 0) AS DECIMAL(18,4)) AS Share_Of_Total_Sales
    FROM sales_trans_h WITH (NOLOCK)
    WHERE sth_flag = 'C' AND YEAR(sec_insert_date) <> 2024
    GROUP BY 
        CASE 
            WHEN total_bill_after_disc < 50 THEN '1. Micro (< 50 L.E)'
            WHEN total_bill_after_disc <= 150 THEN '2. Small (50 - 150 L.E)'
            WHEN total_bill_after_disc <= 500 THEN '3. Medium (151 - 500 L.E)'
            WHEN total_bill_after_disc <= 1500 THEN '4. Large (501 - 1500 L.E)'
            ELSE '5. VIP (> 1500 L.E)'
        END
    ORDER BY Basket_Tier;
    """

    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        read_sql_and_enrich(q_kpi_scorecard, ENGINE).to_excel(
            writer, sheet_name="Executive_KPI_Summary", index=False
        )
        read_sql_and_enrich(q_monthly_pnl, ENGINE).to_excel(
            writer, sheet_name="Executive_Monthly_P&L", index=False
        )
        read_sql_and_enrich(q_abc_analysis, ENGINE).to_excel(
            writer, sheet_name="ABC_Revenue_Pareto", index=False
        )
        read_sql_and_enrich(q_dead_stock_radar, ENGINE).to_excel(
            writer, sheet_name="Expiry_&_DeadStock_Radar", index=False
        )
        read_sql_and_enrich(q_hourly_traffic, ENGINE).to_excel(
            writer, sheet_name="Hourly_Traffic_Patterns", index=False
        )
        read_sql_and_enrich(q_vendor_performance, ENGINE).to_excel(
            writer, sheet_name="Vendor_Supply_Matrix", index=False
        )
        read_sql_and_enrich(q_store_matrix, ENGINE).to_excel(
            writer, sheet_name="MultiBranch_Benchmark", index=False
        )
        read_sql_and_enrich(q_monthly_stock_trend, ENGINE).to_excel(
            writer, sheet_name="Monthly_Stock_Valuation_Trend", index=False
        )
        read_sql_and_enrich(q_therapeutic_profitability, ENGINE).to_excel(
            writer, sheet_name="Therapeutic_Class_Profitability", index=False
        )
        read_sql_and_enrich(q_receivables_aging, ENGINE).to_excel(
            writer, sheet_name="Customer_Receivables_Aging", index=False
        )
        read_sql_and_enrich(q_basket_size_distribution, ENGINE).to_excel(
            writer, sheet_name="Basket_Size_Distribution", index=False
        )

    wb = openpyxl.load_workbook(filename)
    table_names = set()
    for ws_name in wb.sheetnames:
        style_worksheet(
            wb[ws_name],
            header_fill_hex="0B132B",
            is_kpi=(ws_name == "Executive_KPI_Summary"),
            existing_table_names=table_names,
        )
    wb.save(filename)
    logger.info("Module 6 Complete.")


# ==================================================================================================
# MASTER WORKFLOW ORCHESTRATOR
# ==================================================================================================
def main():
    start_time = time.time()
    logger.info(
        "=========================================================================="
    )
    logger.info(
        "STARTING GENIUS ERP COMPREHENSIVE DATA CONSOLIDATION PIPELINE (MASTER)"
    )
    logger.info(
        "Configuration: Year 2024 Universal Exclusion | Clean Tables (No Totals Row)"
    )
    logger.info(
        "=========================================================================="
    )

    output_dir = os.path.join(os.getcwd(), "Genius_Consolidated_Reports")
    os.makedirs(output_dir, exist_ok=True)
    logger.info(f"Target Output Directory: {output_dir}")

    # Verify Database Connectivity
    try:
        with ENGINE.connect() as conn:
            conn.execute(text("SELECT 1"))
        logger.info("Database Connection Test: SUCCESSFUL.")
    except Exception as e:
        logger.error(f"Failed to connect to Microsoft SQL Server: {e}")
        sys.exit(1)

    target_mod = sys.argv[1].strip() if len(sys.argv) > 1 else "all"

    try:
        if target_mod in ("all", "1"):
            generate_inventory_workbook(output_dir)
        if target_mod in ("all", "2"):
            generate_sales_workbook(output_dir)
        if target_mod in ("all", "3"):
            generate_purchases_workbook(output_dir)
        if target_mod in ("all", "4"):
            generate_financial_workbook(output_dir)
        if target_mod in ("all", "5"):
            generate_shifts_workbook(output_dir)
        if target_mod in ("all", "6"):
            generate_executive_bi_workbook(output_dir)
    except Exception as e:
        logger.error(
            f"Extraction halted due to an unexpected exception: {e}", exc_info=True
        )
        sys.exit(1)

    elapsed = round(time.time() - start_time, 2)
    logger.info(
        "=========================================================================="
    )
    logger.info(
        f"ALL 6 MASTER WORKBOOKS (67 SHEETS TOTAL) SUCCESSFULLY GENERATED IN {elapsed}s."
    )
    logger.info(f"Output Location: {output_dir}")
    logger.info(
        "=========================================================================="
    )


if __name__ == "__main__":
    main()
